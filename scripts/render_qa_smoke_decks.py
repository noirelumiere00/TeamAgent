"""描画 QA の CI 用の見本デッキを 2 つ作る（収まる版・はみ出す版）。

render-qa のイメージの中（python-pptx あり・ネットワークなし）で動かす想定:
    docker run --rm --network none -v "$OUT:/out" --entrypoint python teamagent-render-qa:ci /smoke.py /out
本物の資料は使わない（社外秘）。
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


def _deck(text: str, height_in: float) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(height_in))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.text = text
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(18)
            run.font.name = "Yu Gothic UI"
            # 本番の 83 枚のテンプレと同じく、和文（ea）も Yu Gothic UI を指定する。
            rpr = run._r.get_or_add_rPr()
            ea = rpr.makeelement(qn("a:ea"), {"typeface": "Yu Gothic UI"})
            rpr.append(ea)
    return prs


def main() -> None:
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    # 1 行で収まる短い文と、枠の高さの数倍になる長い文（どちらも同じ 4in 幅の枠）。
    _deck("冬のごほうびに一粒。", 1.0).save(out / "clean.pptx")
    long_text = "冬限定という希少性が、一粒の重みを増す。" * 12
    _deck(long_text, 0.6).save(out / "overflow.pptx")


if __name__ == "__main__":
    main()
