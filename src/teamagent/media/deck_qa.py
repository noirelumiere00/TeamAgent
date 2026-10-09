"""描画・外部アクセスを行わない、media runtime 専用の PPTX 検査。"""

from __future__ import annotations

import math
import posixpath
import re
import unicodedata
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Literal

# EMU 丸めや影・裁ち落としの微差を許す。縦横それぞれスライド寸法の 0.5%。
BOUNDS_TOLERANCE = 0.005
# リサイズの丸めを許すが、視認できる歪みは指摘する。
ASPECT_TOLERANCE = 0.03
# 枠候補は画像とほぼ同じ中心・寸法に限定（背景や隣接図形を除外）。
FRAME_CENTER_TOLERANCE = 0.10
FRAME_SIZE_TOLERANCE = 0.25
# 小さい交差は文字の余白とみなし、小さい方の面積の 20% を超えたら指摘。
OVERLAP_TOLERANCE = 0.20
# 字形・改行の推定誤差を許す。行高は標準的なフォントサイズの 1.2 倍。
TEXT_HEIGHT_TOLERANCE = 1.10
LINE_HEIGHT = 1.2
DEFAULT_FONT_PT = 18.0
# 空の小装飾・タイトル余白は除外し、スライド面積の 2% 以上を対象とする。
EMPTY_TEXT_AREA = 0.02
PLACEHOLDER_MARKERS = ("{{", "}}", "<<", "TODO", "xxx", "○○", "要確認（データ未検出）")
# Win/Mac 標準・Office 同梱の一般的な書体。環境差があるため未知は warn のみ。
STANDARD_FONTS = frozenset(
    {
        "Arial",
        "Arial Unicode MS",
        "Calibri",
        "Calibri Light",
        "Aptos",
        "Aptos Display",
        "Times New Roman",
        "Verdana",
        "Tahoma",
        "Georgia",
        "Trebuchet MS",
        "Courier New",
        "Meiryo",
        "メイリオ",
        "Yu Gothic",
        "Yu Gothic UI",
        "游ゴシック",
        "游ゴシック体",
        "Yu Mincho",
        "游明朝",
        "游明朝体",
        "MS Gothic",
        "ＭＳ ゴシック",
        "MS PGothic",
        "ＭＳ Ｐゴシック",
        "MS Mincho",
        "ＭＳ 明朝",
        "MS PMincho",
        "ＭＳ Ｐ明朝",
        "Hiragino Sans",
        "ヒラギノ角ゴシック",
        "ヒラギノ角ゴ ProN W3",
        "ヒラギノ角ゴ ProN W6",
        "Hiragino Kaku Gothic ProN",
        "Hiragino Mincho ProN",
        "ヒラギノ明朝 ProN W3",
        "Helvetica",
        "Helvetica Neue",
        "Symbol",
        "Wingdings",
    }
)
IMAGE_TYPES = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "gif": "image/gif",
    "bmp": "image/bmp",
    "tif": "image/tiff",
    "tiff": "image/tiff",
    "webp": "image/webp",
    "avif": "image/avif",
    "heic": "image/heic",
    "heif": "image/heif",
    "svg": "image/svg+xml",
    "emf": "image/x-emf",
    "wmf": "image/x-wmf",
}
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


@dataclass
class DeckQaFinding:
    slide: int
    shape: str
    kind: str
    severity: Literal["error", "warn"]
    details: dict[str, Any]


@dataclass
class DeckQaReport:
    slide_count: int = 0
    findings: list[DeckQaFinding] = field(default_factory=list)
    fonts: list[str] = field(default_factory=list)

    def add(
        self, slide: int, shape: str, kind: str, severity: Literal["error", "warn"], **details: Any
    ) -> None:
        finding = DeckQaFinding(slide, shape, kind, severity, details)
        # 同一 slide / name / 比率の指摘は一件にまとめる。
        if kind == "image_aspect" and finding in self.findings:
            return
        self.findings.append(finding)

    @property
    def error_count(self) -> int:
        return sum(f.severity == "error" for f in self.findings)

    @property
    def warn_count(self) -> int:
        return sum(f.severity == "warn" for f in self.findings)

    def to_dict(self) -> dict[str, Any]:
        return {**asdict(self), "error_count": self.error_count, "warn_count": self.warn_count}

    def summary(self) -> dict[str, Any]:
        ranked = sorted(self.findings, key=lambda f: f.severity != "error")
        return {
            "error_count": self.error_count,
            "warn_count": self.warn_count,
            "top_findings": [asdict(f) for f in ranked[:5]],
        }


def _signature(data: bytes) -> str | None:
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif"
    if data.startswith(b"BM"):
        return "image/bmp"
    if data.startswith((b"II*\0", b"MM\0*")):
        return "image/tiff"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data[4:8] == b"ftyp" and (b"avif" in data[8:64] or b"avis" in data[8:64]):
        return "image/avif"
    if data[4:8] == b"ftyp" and any(b in data[8:64] for b in (b"heic", b"heix", b"hevc", b"mif1")):
        return "image/heic"
    return None  # SVG/EMF/WMF は bitmap のシグネチャ比較の対象外。


def _xml(data: bytes) -> Any:
    """python-pptx が Presentation() で同じ部品を読むときと同じ parser（entity を展開しない）。"""
    from pptx.oxml import parse_xml

    return parse_xml(data)


def _tagged(node: Any) -> list[Any]:
    """lxml はコメント・処理命令も子に含めるので、要素だけを返す。"""
    return [child for child in node if isinstance(child.tag, str)]


def _package(path: str | Path, report: DeckQaReport) -> None:
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        types = _xml(package.read("[Content_Types].xml"))
        defaults = {
            e.attrib["Extension"].lower(): e.attrib["ContentType"]
            for e in _tagged(types)
            if e.tag.endswith("Default")
        }
        overrides = {
            e.attrib["PartName"].lstrip("/"): e.attrib["ContentType"]
            for e in _tagged(types)
            if e.tag.endswith("Override")
        }
        contexts: dict[str, list[tuple[int, str]]] = {}
        # slideN.xml の番号は削除・並び替え後の表示順とは限らない。
        slide_order: dict[str, int] = {}
        if "ppt/presentation.xml" in names and "ppt/_rels/presentation.xml.rels" in names:
            rel_tree = _xml(package.read("ppt/_rels/presentation.xml.rels"))
            targets = {
                r.get("Id"): posixpath.normpath(posixpath.join("ppt", r.get("Target", ""))).lstrip(
                    "/"
                )
                for r in _tagged(rel_tree)
            }
            tree = _xml(package.read("ppt/presentation.xml"))
            for ordinal, node in enumerate(tree.findall("./p:sldIdLst/p:sldId", NS), 1):
                rid = node.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                )
                slide_order[targets.get(rid, "")] = ordinal
        for name in sorted(names):
            if not name.endswith(".rels"):
                continue
            owner = posixpath.join(
                posixpath.dirname(posixpath.dirname(name)), posixpath.basename(name)[:-5]
            )
            rels = _xml(package.read(name))
            match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", owner)
            slide = slide_order.get(owner, int(match[1]) if match else 1)
            shape_names: dict[str, str] = {}
            if match and owner in names:
                tree = _xml(package.read(owner))
                for pic in tree.findall(".//p:pic", NS):
                    nv = pic.find(".//p:cNvPr", NS)
                    blip = pic.find(".//a:blip", NS)
                    if blip is not None and nv is not None:
                        for key, value in blip.attrib.items():
                            if key.endswith("}embed"):
                                shape_names[value] = nv.get("name", "image")
            rel_ids = {rel.get("Id") for rel in _tagged(rels)}
            for rid, shape in shape_names.items():
                if rid not in rel_ids:
                    report.add(slide, shape, "image_missing", "error", relationship=rid)
            for rel in _tagged(rels):
                if rel.get("TargetMode") == "External" or not rel.get("Type", "").endswith(
                    "/image"
                ):
                    continue
                target = rel.get("Target", "")
                resolved = posixpath.normpath(
                    posixpath.join(posixpath.dirname(owner), target)
                ).lstrip("/")
                shape = shape_names.get(rel.get("Id", ""), resolved)
                contexts.setdefault(resolved, []).append((slide, shape))
                if resolved not in names:
                    report.add(slide, shape, "image_missing", "error", target=resolved)
        for name in sorted(names):
            ext = name.rsplit(".", 1)[-1].lower()
            if not name.startswith("ppt/media/") or ext not in IMAGE_TYPES:
                continue
            content_type = overrides.get(name, defaults.get(ext))
            if content_type == "image/jpg":
                content_type = "image/jpeg"
            with package.open(name) as stream:
                head = stream.read(64)
            actual = _signature(head)
            problems: list[str] = []
            if package.getinfo(name).file_size == 0:
                problems.append("empty")
            if content_type is None:
                problems.append("missing_content_type")
            elif content_type != IMAGE_TYPES[ext]:
                problems.append("extension_content_type_mismatch")
            if actual and (actual != IMAGE_TYPES[ext] or actual != content_type):
                problems.append("signature_mismatch")
            if (
                actual in {"image/webp", "image/avif", "image/heic"}
                or ext in {"webp", "avif", "heic", "heif"}
                or content_type in {"image/webp", "image/avif", "image/heic", "image/heif"}
            ):
                problems.append("unsupported_powerpoint_format")
            if (
                actual is None
                and ext in {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"}
                and head
            ):
                problems.append("invalid_signature")
            for slide, shape in dict.fromkeys(contexts.get(name, [(1, name)])):
                for problem in problems:
                    report.add(
                        slide,
                        shape,
                        "image_registration",
                        "error"
                        if problem in {"missing_content_type", "unsupported_powerpoint_format"}
                        else "warn",
                        problem=problem,
                        bytes=package.getinfo(name).file_size,
                        extension=ext,
                        content_type=content_type,
                        actual=actual,
                    )


# affine = a,b,c,d,e,f ; x'=a*x+c*y+e, y'=b*x+d*y+f
Affine = tuple[float, float, float, float, float, float]
Rect = tuple[float, float, float, float]


def _compose(m: Affine, n: Affine) -> Affine:
    a, b, c, d, e, f = m
    g, h, i, j, k, offset_y = n
    return (
        a * g + c * h,
        b * g + d * h,
        a * i + c * j,
        b * i + d * j,
        a * k + c * offset_y + e,
        b * k + d * offset_y + f,
    )


def _shape_sources(shape: Any, slide: Any) -> list[Any]:
    """slide → idx 対応 layout → type 対応 master の入力欄。"""
    sources = [shape._element]
    if not shape.is_placeholder:
        return sources
    idx = shape.placeholder_format.idx
    layout = next(
        (p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == idx), None
    )
    if layout is not None:
        sources.append(layout._element)
    ph_type = int((layout or shape).placeholder_format.type)
    # OOXML の master placeholder は layout の idx ではなく type で対応する。
    master_type = 1 if ph_type in {1, 3, 5} else ph_type if ph_type in {13, 14, 15, 16} else 2
    master = next(
        (
            p
            for p in slide.slide_layout.slide_master.placeholders
            if int(p.placeholder_format.type) == master_type
        ),
        None,
    )
    if master is not None:
        sources.append(master._element)
    return sources


def _dimensions(shape: Any, slide: Any) -> tuple[float, float, float, float, float]:
    values: list[float] = []
    sources = _shape_sources(shape, slide) if slide is not None else [shape._element]
    for query, attr, fallback in (
        ("./p:spPr/a:xfrm/a:off | ./p:grpSpPr/a:xfrm/a:off | ./p:xfrm/a:off", "x", 0),
        ("./p:spPr/a:xfrm/a:off | ./p:grpSpPr/a:xfrm/a:off | ./p:xfrm/a:off", "y", 0),
        ("./p:spPr/a:xfrm/a:ext | ./p:grpSpPr/a:xfrm/a:ext | ./p:xfrm/a:ext", "cx", 0),
        ("./p:spPr/a:xfrm/a:ext | ./p:grpSpPr/a:xfrm/a:ext | ./p:xfrm/a:ext", "cy", 0),
        ("./p:spPr/a:xfrm | ./p:grpSpPr/a:xfrm | ./p:xfrm", "rot", 0),
    ):
        value = next(
            (
                node.get(attr)
                for source in sources
                for node in source.xpath(query)
                if node.get(attr) is not None
            ),
            fallback,
        )
        values.append(float(value))
    return values[0], values[1], values[2], values[3], values[4] / 60000


def _walk(
    shapes: Any, parent: Affine = (1, 0, 0, 1, 0, 0), slide: Any = None
) -> list[tuple[Any, Rect, Affine]]:
    result = []
    for shape in shapes:
        x, y, w, h, rotation = _dimensions(shape, slide)
        angle = math.radians(rotation)
        co, si = math.cos(angle), math.sin(angle)
        m = _compose(
            parent,
            (
                co,
                si,
                -si,
                co,
                x + w / 2 - co * w / 2 + si * h / 2,
                y + h / 2 - si * w / 2 - co * h / 2,
            ),
        )
        points = [
            (m[0] * u + m[2] * v + m[4], m[1] * u + m[3] * v + m[5])
            for u, v in ((0, 0), (w, 0), (0, h), (w, h))
        ]
        xs, ys = zip(*points, strict=True)
        rect = (min(xs), min(ys), max(xs), max(ys))
        if hasattr(shape, "shapes"):
            xf = shape._element.grpSpPr.xfrm
            if xf is not None and xf.chExt is not None and xf.chOff is not None:
                if xf.get("flipH") in {"1", "true"}:
                    m = _compose(m, (-1, 0, 0, 1, w, 0))
                if xf.get("flipV") in {"1", "true"}:
                    m = _compose(m, (1, 0, 0, -1, 0, h))
                sx, sy = w / (xf.chExt.cx or 1), h / (xf.chExt.cy or 1)
                child = _compose(m, (sx, 0, 0, sy, -xf.chOff.x * sx, -xf.chOff.y * sy))
                result.extend(_walk(shape.shapes, child, slide))
        else:
            result.append((shape, rect, m))
    return result


def _is_picture(shape: Any) -> bool:
    # hasattr(image) は property を評価し、切れた rels で例外になる。
    return bool(shape._element.xpath(".//a:blip"))


def _intersection(a: Rect, b: Rect) -> float:
    return max(0, min(a[2], b[2]) - max(a[0], b[0])) * max(0, min(a[3], b[3]) - max(a[1], b[1]))


def _area(a: Rect) -> float:
    return max(0, a[2] - a[0]) * max(0, a[3] - a[1])


def _outside(rect: Rect, width: float, height: float) -> bool:
    return (
        rect[0] < -width * BOUNDS_TOLERANCE
        or rect[1] < -height * BOUNDS_TOLERANCE
        or rect[2] > width * (1 + BOUNDS_TOLERANCE)
        or rect[3] > height * (1 + BOUNDS_TOLERANCE)
    )


def _inherited_style(shape: Any, slide: Any) -> tuple[float, str]:
    """プレースホルダ・マスター・テーマの既定書式を辿る（描画はしない）。"""
    sources = _shape_sources(shape, slide)
    title = shape.is_placeholder and int(shape.placeholder_format.type) in {1, 3}
    style = "titleStyle" if title else "bodyStyle" if shape.is_placeholder else "otherStyle"
    sources.extend(slide.slide_layout.slide_master._element.xpath(f"./p:txStyles/p:{style}"))
    sources.extend(slide.part.package.presentation_part._element.xpath("./p:defaultTextStyle"))
    size, font = DEFAULT_FONT_PT * 12700, ""
    found_size = False
    for source in sources:
        nodes = source.findall(".//a:lvl1pPr/a:defRPr", NS) or source.findall(".//a:defRPr", NS)
        for node in nodes[:1]:
            if not found_size and node.get("sz"):
                size, found_size = int(node.get("sz")) * 127, True
            if not font:
                faces = node.findall("./a:ea", NS) + node.findall("./a:latin", NS)
                font = next((n.get("typeface") for n in faces if n.get("typeface")), "")
    if not font:
        refs = shape._element.xpath("./p:style/a:fontRef")
        font = "+mj-lt" if (refs and refs[0].get("idx") == "major") or title else "+mn-lt"
    if font.startswith("+"):
        master = slide.slide_layout.slide_master.part
        for rel in master.rels.values():
            if rel.reltype.endswith("/theme"):
                theme = _xml(rel.target_part.blob)
                family = "majorFont" if font.startswith("+mj") else "minorFont"
                node = (
                    theme.find(f".//a:{family}/a:font[@script='Jpan']", NS)
                    if font.endswith("-ea")
                    else None
                )
                if node is None:
                    node = theme.find(f".//a:{family}/a:latin", NS)
                font = node.get("typeface", "") if node is not None else ""
                break
    return size, font


def _body_properties(frame: Any, shape: Any, slide: Any) -> tuple[dict[str, str], str, float]:
    attrs: dict[str, str] = {}
    autofit = ""
    scale = 1.0
    bodies = [frame._txBody.bodyPr]
    if shape.has_text_frame and frame._txBody is shape.text_frame._txBody:
        bodies.extend(
            source.xpath("./p:txBody/a:bodyPr")[0]
            for source in _shape_sources(shape, slide)[1:]
            if source.xpath("./p:txBody/a:bodyPr")
        )
    for body in bodies:
        for key, value in body.attrib.items():
            attrs.setdefault(key, value)
        if not autofit:
            nodes = body.xpath("./a:noAutofit | ./a:normAutofit | ./a:spAutoFit")
            if nodes:
                autofit = nodes[0].tag.rsplit("}", 1)[-1]
                scale = int(nodes[0].get("fontScale", "100000")) / 100000
    return attrs, autofit, scale


def _paragraph_properties(paragraph: Any, shape: Any, slide: Any) -> list[Any]:
    level = paragraph.level + 1
    nodes = list(paragraph._p.xpath("./a:pPr"))
    for source in _shape_sources(shape, slide):
        # paragraph properties on an inherited placeholder precede its list style.
        if source is not shape._element:
            nodes.extend(source.xpath(f"./p:txBody/a:p/a:pPr[@lvl='{level - 1}']"))
            if level == 1:
                nodes.extend(source.xpath("./p:txBody/a:p/a:pPr[not(@lvl)]"))
        nodes.extend(source.xpath(f"./p:txBody/a:lstStyle/a:lvl{level}pPr"))
    title = shape.is_placeholder and int(shape.placeholder_format.type) in {1, 3}
    style = "titleStyle" if title else "bodyStyle" if shape.is_placeholder else "otherStyle"
    nodes.extend(
        slide.slide_layout.slide_master._element.xpath(f"./p:txStyles/p:{style}/a:lvl{level}pPr")
    )
    nodes.extend(
        slide.part.package.presentation_part._element.xpath(f"./p:defaultTextStyle/a:lvl{level}pPr")
    )
    return nodes


def _text_height(
    frame: Any,
    width: float,
    inherited_size: float,
    shape: Any,
    slide: Any,
    attrs: dict[str, str],
    scale: float,
) -> float:
    usable = max(1.0, width - int(attrs.get("lIns", "91440")) - int(attrs.get("rIns", "91440")))
    height = float(int(attrs.get("tIns", "45720")) + int(attrs.get("bIns", "45720")))
    for paragraph in frame.paragraphs:
        properties = _paragraph_properties(paragraph, shape, slide)
        sizes = [n.get("sz") for p in properties for n in p.findall("./a:defRPr[@sz]", NS)]
        size = (int(sizes[0]) * 127 if sizes else inherited_size) * scale
        lines, used, max_size = 1, 0.0, 0.0
        chunks = []
        for node in paragraph._p:
            if node.tag.endswith("}br"):
                chunks.append(("\v", size))
            elif node.tag.endswith(("}r", "}fld")):
                txt = node.find("a:t", NS)
                props = node.find("a:rPr", NS)
                run_size = (
                    int(props.get("sz")) * 127 * scale
                    if props is not None and props.get("sz")
                    else size
                )
                chunks.append((txt.text or "" if txt is not None else "", float(run_size)))
        for chunk, run_size in chunks:
            max_size = max(max_size, run_size)
            for char in chunk:
                if char in "\n\v":
                    lines, used = lines + 1, 0.0
                    continue
                advance = run_size * (1 if unicodedata.east_asian_width(char) in "WF" else 0.55)
                if attrs.get("wrap", "square") != "none" and used and used + advance > usable:
                    lines, used = lines + 1, 0.0
                used += advance
        if not chunks:
            end = paragraph._p.find("a:endParaRPr", NS)
            max_size = (
                int(end.get("sz")) * 127 * scale if end is not None and end.get("sz") else size
            )

        def spacing(
            tag: str,
            default: float,
            props_list: list[Any] = properties,
            font_size: float = max_size,
        ) -> float:
            for props in props_list:
                nodes = props.findall(f"./a:{tag}/*", NS)
                if nodes:
                    node = nodes[0]
                    value = int(node.get("val", "0"))
                    return (
                        font_size * LINE_HEIGHT * value / 100000
                        if node.tag.endswith("}spcPct")
                        else value * 127
                    )
            return default

        height += lines * spacing("lnSpc", max_size * LINE_HEIGHT)
        height += spacing("spcBef", 0) + spacing("spcAft", 0)
    return height


def _picture_size(shape: Any) -> tuple[int, int] | None:
    """リンク・ベクターは bitmap decoder に渡さない。検査例外は破損の証拠ではない。"""
    from io import BytesIO

    from PIL import Image, UnidentifiedImageError

    blips = shape._element.xpath("./p:blipFill/a:blip")
    if not blips:
        return None
    rid = blips[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    if not rid:
        return None
    try:
        part = shape.part.related_part(rid)
    except KeyError:
        return None  # package 検査が参照切れを報告する。
    if part.content_type in {
        "image/svg+xml",
        "image/x-emf",
        "image/x-wmf",
        "image/emf",
        "image/wmf",
    }:
        return None
    try:
        with Image.open(BytesIO(part.blob)) as bitmap:
            return bitmap.size
    except (UnidentifiedImageError, OSError, ValueError):
        return None  # 未対応 decoder と破損を混同しない。


def inspect_pptx(
    path: str | Path, *, expected_slides: int | None = None, strict: bool = False
) -> DeckQaReport:
    """ファイルを変更せず検査。文字・配置は描画無しの保守的な推定。

    strict=True は意図的な伸縮も含め、画像の縦横比の差を error にする。
    """
    from pptx import Presentation

    report = DeckQaReport()
    _package(path, report)
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        report.add(1, "package", "package_unreadable", "error", exception=type(exc).__name__)
        return report
    report.slide_count = len(presentation.slides)
    if expected_slides is not None and expected_slides != report.slide_count:
        report.add(
            1,
            "presentation",
            "slide_count",
            "error",
            expected=expected_slides,
            actual=report.slide_count,
        )
    width, height = float(presentation.slide_width or 0), float(presentation.slide_height or 0)
    fonts: set[str] = set()
    for number, slide in enumerate(presentation.slides, 1):
        entries = _walk(slide.shapes, slide=slide)
        texts: list[tuple[str, Rect]] = []
        for shape, rect, m in entries:
            name = shape.name
            if _outside(rect, width, height):
                report.add(
                    number,
                    name,
                    "out_of_bounds",
                    "warn",  # 裁ち落としや枠の余白は描画無しでは判別できない。
                    rect=rect,
                    slide_width=width,
                    slide_height=height,
                )
            if shape._element.tag.endswith("}pic"):
                if not any(
                    source.xpath("./p:spPr/a:prstGeom | ./p:spPr/a:custGeom")
                    for source in _shape_sources(shape, slide)
                ):
                    report.add(number, name, "picture_no_geometry", "error")
                dimensions = _picture_size(shape)
                if dimensions is not None:
                    iw, ih = dimensions
                    crop = shape._element.xpath("./p:blipFill/a:srcRect")
                    crop_w = (
                        1 - sum(int(crop[0].get(k, "0")) / 100000 for k in ("l", "r"))
                        if crop
                        else 1
                    )
                    crop_h = (
                        1 - sum(int(crop[0].get(k, "0")) / 100000 for k in ("t", "b"))
                        if crop
                        else 1
                    )
                    natural = iw * crop_w / (ih * crop_h) if ih * crop_h > 0 else 0
                    _, _, sw, sh, _ = _dimensions(shape, slide)
                    displayed = math.hypot(m[0], m[1]) * sw / max(1, math.hypot(m[2], m[3]) * sh)
                    difference = abs(displayed / natural - 1) if natural > 0 else 1
                    if difference > ASPECT_TOLERANCE:
                        report.add(
                            number,
                            name,
                            "image_aspect",
                            "error" if strict else "warn",
                            difference=difference,
                            displayed=displayed,
                            natural=natural,
                        )
                for frame, box, _ in entries:
                    if frame is shape or _is_picture(frame) or frame.has_table or frame.has_chart:
                        continue
                    if getattr(frame, "has_text_frame", False) and frame.text.strip():
                        continue
                    rw, rh, bw, bh = (
                        rect[2] - rect[0],
                        rect[3] - rect[1],
                        box[2] - box[0],
                        box[3] - box[1],
                    )
                    if min(rw, rh, bw, bh) <= 0:
                        continue
                    if (
                        abs(rw / bw - 1) <= FRAME_SIZE_TOLERANCE
                        and abs(rh / bh - 1) <= FRAME_SIZE_TOLERANCE
                        and abs((rect[0] + rect[2] - box[0] - box[2]) / 2)
                        <= bw * FRAME_CENTER_TOLERANCE
                        and abs((rect[1] + rect[3] - box[1] - box[3]) / 2)
                        <= bh * FRAME_CENTER_TOLERANCE
                        and _intersection(rect, box) < _area(rect) * (1 - BOUNDS_TOLERANCE)
                    ):
                        report.add(
                            number,
                            name,
                            "image_frame_overflow",
                            "warn",
                            frame=frame.name,
                            rect=rect,
                            frame_rect=box,
                        )
            frames = []
            if shape.has_text_frame:
                _, _, sw, sh, _ = _dimensions(shape, slide)
                frames.append((shape.text_frame, sw, sh, rect))
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if not cell.is_spanned:
                            frames.append(
                                (
                                    cell.text_frame,
                                    float(shape.width) / len(row.cells),
                                    float(row.height),
                                    rect,
                                )
                            )
            for frame, fw, fh, text_rect in frames:
                text = frame.text
                inherited_size, inherited_font = _inherited_style(shape, slide)
                explicit_fonts = {
                    n.get("typeface") for n in frame._txBody.xpath(".//a:latin | .//a:ea | .//a:cs")
                }
                if text.strip() and inherited_font and not explicit_fonts:
                    fonts.add(inherited_font)
                    if inherited_font not in STANDARD_FONTS:
                        report.add(number, name, "font_unavailable", "warn", font=inherited_font)
                for font in sorted(f for f in explicit_fonts if f):
                    if font and not font.startswith("+"):
                        fonts.add(font)
                        if font not in STANDARD_FONTS:
                            report.add(number, name, "font_unavailable", "warn", font=font)
                for marker in PLACEHOLDER_MARKERS:
                    count = (
                        len(re.findall(r"(?<!\w)" + marker + r"(?!\w)", text))
                        if marker in {"TODO", "xxx"}
                        else text.count(marker)
                    )
                    if count:
                        report.add(
                            number,
                            name,
                            "placeholder",
                            "warn" if marker == "○○" else "error",
                            marker=marker,
                            count=count,
                        )
                if "要確認" in text:
                    report.add(number, name, "review_required", "warn", count=text.count("要確認"))
                if not text.strip():
                    if (
                        _area(text_rect) >= width * height * EMPTY_TEXT_AREA
                        and not _is_picture(shape)
                        and (
                            shape.is_placeholder
                            or (
                                shape._element.tag.endswith("}sp")
                                and shape._element.xpath("./p:nvSpPr/p:cNvSpPr[@txBox='1']")
                            )
                        )
                    ):
                        report.add(
                            number,
                            name,
                            "empty_text",
                            "warn",
                            area_ratio=_area(text_rect) / (width * height),
                        )
                    continue
                attrs, autofit, scale = _body_properties(frame, shape, slide)
                needed = _text_height(frame, fw, inherited_size, shape, slide, attrs, scale)
                if needed > fh * TEXT_HEIGHT_TOLERANCE:
                    if autofit == "spAutoFit":
                        grown = (
                            text_rect[0],
                            text_rect[1],
                            text_rect[2],
                            text_rect[3] + (needed - fh) * math.hypot(m[2], m[3]),
                        )
                        hits = [
                            other.name
                            for other, box, _ in entries
                            if other is not shape
                            and _intersection(text_rect, box)
                            < _area(text_rect) * (1 - BOUNDS_TOLERANCE)
                            and _intersection(grown, box) > _intersection(text_rect, box) + 1
                        ]
                        if _outside(grown, width, height) or hits:
                            report.add(
                                number,
                                name,
                                "text_autofit_collision",
                                "warn",
                                needed_height=needed,
                                available_height=fh,
                                grown_rect=grown,
                                collisions=hits,
                            )
                    else:
                        # 過大な insets / anchor だけではクリッピングを断定できない。
                        # error は余白を除いても 2 倍以上の文字量がある場合に限る。
                        text_needed = (
                            needed
                            - int(attrs.get("tIns", "45720"))
                            - int(attrs.get("bIns", "45720"))
                        )
                        severity: Literal["error", "warn"] = (
                            "error"
                            if autofit == "noAutofit"
                            and text_needed >= 2 * fh
                            and int(attrs.get("tIns", "45720")) + int(attrs.get("bIns", "45720"))
                            < fh
                            and int(attrs.get("lIns", "91440")) + int(attrs.get("rIns", "91440"))
                            < fw
                            else "warn"
                        )
                        report.add(
                            number,
                            name,
                            "text_overflow",
                            severity,
                            needed_height=needed,
                            available_height=fh,
                        )
                if shape.has_text_frame:
                    texts.append((name, text_rect))
        for index, (name, rect) in enumerate(texts):
            for other, box in texts[index + 1 :]:
                area = min(_area(rect), _area(box))
                ratio = _intersection(rect, box) / area if area else 0
                if ratio > OVERLAP_TOLERANCE:
                    report.add(
                        number, name, "text_overlap", "warn", other=other, overlap_ratio=ratio
                    )
    report.fonts = sorted(fonts)
    return report
