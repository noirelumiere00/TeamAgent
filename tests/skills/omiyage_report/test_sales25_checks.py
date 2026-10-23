"""Final-byte machine evidence, using real L1 inspection and fake delivery only."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any
from unittest.mock import Mock

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pydantic import ValidationError

from teamagent.adapters.proposal_job_store import ProposalJobStore
from teamagent.media.deck_qa import inspect_pptx
from teamagent.skills._shared.long_jobs import ORIGIN_KEY
from teamagent.skills.omiyage_report import sales25
from tests.skills.omiyage_report.test_sales25_parent import (
    FakeSlack,
    Pipeline,
    build,
    ctx,
    request,
    result,
)
from tests.skills.omiyage_report.test_sales25_parent import flags as flags

SHA = "a" * 64
CHECKS = ("check_fmt", "check_top5_analysis", "deck_qa_l1", "text_fit")


def machine_checks() -> sales25.MachineChecks:
    return sales25.MachineChecks(
        check_fmt=sales25.CheckResult(status="passed", checked_sha256=SHA),
        check_top5_analysis=sales25.CheckResult(status="passed", checked_sha256=SHA),
        deck_qa_l1=sales25.DeckQaCheck(
            status="passed", checked_sha256=SHA, error_count=0, review_pages=[]
        ),
        text_fit=sales25.CheckResult(status="passed", checked_sha256=SHA),
    )


def test_real_l1_retains_all_warning_pages_and_errors(tmp_path: Path) -> None:
    presentation = Presentation()
    for number in range(1, 26):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        paragraph = text.text_frame.paragraphs[0]
        paragraph.text = f"要確認: {number}"
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(16)
    path = tmp_path / "warning-pages.pptx"
    presentation.save(str(path))
    sha = hashlib.sha256(path.read_bytes()).hexdigest()

    report = inspect_pptx(path, expected_slides=25)
    assert report.error_count == 0
    assert len(report.summary()["top_findings"]) == 5
    check = sales25.DeckQaCheck.from_report(report, sha)
    assert check.status == "passed" and check.checked_sha256 == sha
    assert check.review_pages == list(range(1, 26))

    wrong_slide_count = inspect_pptx(path, expected_slides=24)
    failed = sales25.DeckQaCheck.from_report(wrong_slide_count, sha)
    assert failed.status == "failed"
    assert failed.error_count == wrong_slide_count.error_count > 0
    assert failed.review_pages == list(range(1, 26))


def test_all_four_checks_accept_only_their_final_sha() -> None:
    checks = machine_checks()
    assert checks.accepts(SHA)
    assert not checks.accepts("b" * 64)


@pytest.mark.parametrize("name", CHECKS)
@pytest.mark.parametrize("status", ["failed", "incomplete"])
def test_each_machine_failure_blocks_delivery(name: str, status: str) -> None:
    data = machine_checks().model_dump()
    data[name]["status"] = status
    assert not sales25.MachineChecks.model_validate(data).accepts(SHA)


@pytest.mark.parametrize("name", CHECKS)
def test_each_machine_check_must_match_final_sha(name: str) -> None:
    data = machine_checks().model_dump()
    data[name]["checked_sha256"] = "b" * 64
    assert not sales25.MachineChecks.model_validate(data).accepts(SHA)


@pytest.mark.parametrize("name", CHECKS)
def test_machine_evidence_cannot_be_missing(name: str) -> None:
    data = machine_checks().model_dump()
    del data[name]
    with pytest.raises(ValidationError):
        sales25.MachineChecks.model_validate(data)


def test_l1_error_count_blocks_even_when_status_claims_passed() -> None:
    data = machine_checks().model_dump()
    data["deck_qa_l1"]["error_count"] = 1
    assert not sales25.MachineChecks.model_validate(data).accepts(SHA)


@pytest.mark.parametrize("reason", ["audio_unverified", "shortfall", None])
def test_top5_incomplete_requires_an_explicit_deliverable_shortage(reason: str | None) -> None:
    data = machine_checks().model_dump()
    data["check_top5_analysis"].update(status="incomplete", reason=reason)
    assert sales25.MachineChecks.model_validate(data).accepts(SHA) is (reason is not None)


def test_saved_pending_without_machine_evidence_cannot_restore_attachment() -> None:
    memory: dict[str, dict[str, Any]] = {}
    store = ProposalJobStore(table_name="", memory=memory)
    pipeline = Pipeline()
    accepted = build(store, pipeline).run(request(), ctx())
    saved = result(store, accepted.job_id)
    assert saved.delivery_status == "pending"
    old_result = saved.model_dump()
    del old_result["checkpoints"]["machine_check"]["checks"]
    memory[accepted.job_id]["result_json"] = json.dumps(old_result)

    reader = Mock(return_value=pipeline.data)
    slack = FakeSlack(store, accepted.job_id)
    parent = sales25.Parent(store, reader=reader, slack=slack)
    row = store.get_job(accepted.job_id)
    assert row is not None
    context = ctx()
    output = parent.status(accepted.job_id, row, context)
    assert output.delivery_status == "failed"
    assert not context.metadata[ORIGIN_KEY].pending
    reader.assert_not_called()
    assert not slack.uploads and not slack.dm_users
