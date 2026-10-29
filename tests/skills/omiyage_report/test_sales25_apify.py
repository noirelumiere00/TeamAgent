"""Fake worker → existing Apify actor → sales25 intake/authoring; no external services."""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime
from types import SimpleNamespace
from typing import Any
from unittest.mock import Mock

import httpx
import pytest
from pptx import Presentation

from teamagent.adapters import tiktok_search_fallback
from teamagent.adapters.apify_client import ACTOR_TIKTOK_VIDEO, ApifyClient, ApifyRunResult
from teamagent.adapters.cost_guard import CostLimitExceededError
from teamagent.adapters.media_job import MediaMarkerExistsError
from teamagent.adapters.proposal_job_store import ProposalJobStore
from teamagent.adapters.tiktok_search_fallback import apply_sales25_fallback
from teamagent.mcp_gateway import async_job_notify as notify
from teamagent.media.contracts import MediaArtifact, MediaJobResult, S3ObjectRef
from teamagent.skills._shared.long_jobs import ORIGIN_KEY
from teamagent.skills.first_visit import sales25_contract as contract
from teamagent.skills.omiyage_report import sales25, sales25_wire
from teamagent.skills.omiyage_report.sales25_wire import Wire, production_stages
from teamagent.skills.omiyage_report.schema import OmiyageReportStatusInput
from teamagent.skills.omiyage_report.skill import OmiyageReportStatusSkill
from tests.adapters.test_apify_client import _FakeApify, _ReservationLedger
from tests.skills.first_visit.sales25_fixture import inputs
from tests.skills.omiyage_report.test_sales25_parent import FakeSlack
from tests.skills.omiyage_report.test_sales25_wire import (
    Transport,
    build,
    ctx,
    flags,
    model,
    request,
    result,
)

__all__ = ["flags"]

QUERIES = ["ブランドA", "ブランドB", "ブランドC", "朝のケア", "乾燥対策", "時短ケア"]


class Worker(Transport):
    def __init__(self, failures: dict[int, str]) -> None:
        super().__init__()
        self.failures = failures

    def stage_marker(self, **kwargs: Any) -> S3ObjectRef:
        if self.find_staged(**kwargs):
            raise MediaMarkerExistsError("MEDIA_MARKER_EXISTS")
        return self.stage_bytes(**kwargs)

    def run_sync(self, request: Any, **kwargs: Any) -> tuple[dict[str, bytes], dict[str, Any]]:
        if request.operation.kind != "tiktok_acquire":
            return super().run_sync(request, **kwargs)
        files, metadata = super().run_sync(request)
        if request.operation.kind == "tiktok_acquire":
            code = self.failures.get(QUERIES.index(request.operation.keywords[0]))
            if code:
                raw = json.loads(files["sales25-raw.json"])
                status = json.loads(files["acquisition-status.json"])
                raw.update(
                    ok=False,
                    status="failed",
                    error_code=code,
                    stop_reason="refused" if code == "TIKTOK_BOT_WALL" else "api_error",
                    diagnostics={
                        "pagesFetched": 1,
                        "scrollCount": 1,
                        "captchaDetected": code == "TIKTOK_BOT_WALL",
                    },
                )
                status.update(
                    accepted=False,
                    status="failed",
                    stop_reason=raw["stop_reason"],
                    diagnostics=raw["diagnostics"],
                )
                files.update(
                    {
                        "sales25-raw.json": json.dumps(raw).encode(),
                        "acquisition-status.json": json.dumps(status).encode(),
                    }
                )
        if kwargs.get("capture_result"):
            kwargs["capture_result"](
                MediaJobResult(
                    job_id=request.job_id,
                    status="done",
                    artifacts=tuple(
                        MediaArtifact(
                            name=name,
                            object=self.stage_bytes(
                                job_id=request.job_id,
                                name=name,
                                body=body,
                                content_type="application/json",
                            ),
                        )
                        for name, body in files.items()
                    ),
                )
            )
        return files, metadata


def actor_items(index: int, count: int = 12) -> list[dict[str, Any]]:
    posts = inputs()[1][f"q{index}"]["videos"][:count]
    return [
        {
            "id": post["id"],
            "webVideoUrl": post["url"].replace("/photo/", "/video/"),
            "text": post["desc"],
            "authorMeta": {"name": "synthetic", "fans": None},
            **post["stats"],
            "isAd": post["isAd"],
            "hashtags": [{"name": tag} for tag in post["hashtags"]],
            "isSlideshow": True,
            "slideshowImageLinks": ["synthetic"] * 3,
        }
        for post in reversed(posts)
    ]


@pytest.fixture
def apify(monkeypatch: pytest.MonkeyPatch) -> Mock:
    monkeypatch.setenv("USE_TIKTOK_APIFY_FALLBACK", "1")
    fake = Mock()

    def search(query: str, **kwargs: Any) -> ApifyRunResult:
        return ApifyRunResult(
            items=actor_items(QUERIES.index(query)),
            actor_id=ACTOR_TIKTOK_VIDEO,
            run_id="synthetic-run",
            status="SUCCEEDED",
            estimated_cost_usd=0.048,
        )

    fake.tiktok_search.side_effect = search
    ledger = object()
    monkeypatch.setattr(
        "teamagent.adapters.tiktok_search_fallback.CostGuard.from_env", lambda: ledger
    )

    def factory(**kwargs: Any) -> Mock:
        assert kwargs["ledger"] is ledger
        return fake

    monkeypatch.setattr(ApifyClient, "from_env", factory)
    return fake


@pytest.mark.parametrize("failed", [list(range(6)), [1, 4]])
def test_intake_recovers_only_failed_queries_and_renders_provenance(
    apify: Mock, failed: list[int]
) -> None:
    failed = list(range(min(failed), 6))
    client = Worker({n: "TIKTOK_BOT_WALL" if n < 4 else "TIKTOK_API_ERROR" for n in failed})
    store = ProposalJobStore(table_name="", memory={})
    wire = Wire(client, model, store=store)
    accepted = build(store, sales25_stages=production_stages(wire)).run(request(), ctx())
    saved = result(store, accepted.job_id)
    assert saved.delivery_status == "pending", saved.model_dump_json(indent=2)
    assert saved.final_pptx_ref is not None
    assert len(saved.checkpoints) == 12
    assert {call.args[0] for call in apify.tiktok_search.call_args_list} == {
        QUERIES[n] for n in failed
    }
    assert apify.tiktok_search.call_count == 2 * len(failed)
    assert all(call.kwargs["max_items"] == 100 for call in apify.tiktok_search.call_args_list)
    evidence = []
    for n in range(6):
        checkpoint = saved.checkpoints[f"search_{n + 1}"]
        assert checkpoint.actual == 12
        payload = wire._json(
            {
                "owner_id": checkpoint.refs[0]["owner_id"],
                "job_id": accepted.job_id,
                "started_at": store.get_job(accepted.job_id)["created_at"],
            },
            checkpoint.refs[0],
        )
        evidence.append(payload)
        expected = "apify" if n in failed else "worker"
        assert payload["raw"]["acquired_via"] == payload["acquired_via"] == expected
        assert payload["acquisition_status"]["actual"] == 12
        assert payload["acquisition_status"]["target"] == 100
        if n in failed:
            assert [p["id"] for p in payload["raw"]["videos"]] == [p["id"] for p in actor_items(n)]
            assert [p["rank_display"] for p in payload["raw"]["videos"]] == list(range(1, 13))
            assert payload["worker_evidence"]["raw"]["status"] == "failed"
            assert payload["raw"]["order_provider"] == "apify_dataset_order"
    assert (
        len([r for r in client.requests if r.operation.kind == "tiktok_acquire"]) == min(failed) + 1
    )
    deck = Presentation(io.BytesIO(client.download(saved.final_pptx_ref)))
    text = "\n".join(
        node.text or ""
        for node in deck.slides[24]._element.iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
        )
    )
    assert "取得経路: " + ("Apify" if len(failed) == 6 else "ブラウザ／Apify") in text
    assert "100本目標→12本取得" in text
    notes = deck.slides[24].notes_slide.notes_text_frame.text
    assert all(query in notes for query in QUERIES)
    assert "Apifyは返却順を表示順として記録" in notes
    # Recovery evidence survives an adapter replay without another paid run.
    original = evidence[failed[0]]["worker_evidence"]
    replay = apply_sales25_fallback(
        original, media_client=client, deadline_s=300, request_id="replay"
    )
    assert replay["raw"] == evidence[failed[0]]["raw"]
    assert apify.tiktok_search.call_count == 2 * len(failed)


@pytest.mark.parametrize("enabled", [True, False])
def test_double_failure_remains_incomplete_and_off_never_calls_apify(
    apify: Mock, monkeypatch: pytest.MonkeyPatch, enabled: bool
) -> None:
    monkeypatch.setenv("USE_TIKTOK_APIFY_FALLBACK", "1" if enabled else "0")
    apify.tiktok_search.side_effect = RuntimeError("APIFY_TIMEOUT: secret-should-not-leak")
    client = Worker(dict.fromkeys(range(6), "TIKTOK_BOT_WALL"))
    store = ProposalJobStore(table_name="", memory={})
    wire = Wire(client, model, store=store)
    accepted = build(store, sales25_stages=production_stages(wire)).run(request(), ctx())
    saved = result(store, accepted.job_id)
    assert saved.status == "incomplete"
    assert len(saved.checkpoints) == 6
    assert saved.final_pptx_ref is None and saved.delivery_status == "not_ready"
    assert apify.tiktok_search.call_count == (6 if enabled else 0)
    assert not any(r.operation.kind == "sales25_pptx" for r in client.requests)


def test_page_two_captcha_sends_remaining_queries_directly_to_apify_and_is_job_local(
    apify: Mock, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("OMIYAGE_SALES25_SEARCH_DEPTH", "50")
    client = Worker({0: "TIKTOK_BOT_WALL"})
    wire = Wire(client, model, store=ProposalJobStore(table_name="", memory={}))
    base: dict[str, Any] = {
        "owner_id": "synthetic",
        "job_id": "gentle",
        "started_at": datetime.now(UTC).isoformat(),
        "_search_session": {},
    }
    outcomes = [
        wire.search({**base, "query": query}, {}, f"gentle:search_{i}")
        for i, query in enumerate(QUERIES, 1)
    ]
    assert all(outcome.status == "partial" and outcome.actual == 12 for outcome in outcomes)
    assert [r.operation.keywords for r in client.requests] == [(QUERIES[0],)]
    assert [call.args[0] for call in apify.tiktok_search.call_args_list] == [
        query for query in QUERIES for _ in range(2)
    ]
    assert all(call.kwargs["max_items"] == 50 for call in apify.tiktok_search.call_args_list)
    first = wire._json(base, outcomes[0].refs[0])["worker_evidence"]
    assert len(first["raw"]["videos"]) == 12
    assert first["raw"]["diagnostics"]["pagesFetched"] == 1
    assert first["raw"]["diagnostics"]["captchaDetected"] is True
    for outcome in outcomes[1:]:
        assert outcome.child_job_ids == []
        skipped = wire._json(base, outcome.refs[0])
        assert skipped["worker_evidence"]["browser_skipped"] is True
    wire.search(
        {**base, "job_id": "other", "query": QUERIES[1], "_search_session": {}},
        {},
        "other:search_1",
    )
    assert len(client.requests) == 2


def test_successful_browser_queries_pause_thirty_seconds_with_fake_time(
    apify: Mock, monkeypatch: pytest.MonkeyPatch
) -> None:
    elapsed = [0.0]
    pauses = []

    def sleep(seconds: float) -> None:
        pauses.append(seconds)
        elapsed[0] += seconds

    monkeypatch.setattr(
        sales25_wire,
        "time",
        SimpleNamespace(
            time=lambda: datetime.now(UTC).timestamp(),
            monotonic=lambda: elapsed[0],
            sleep=sleep,
        ),
    )
    client = Worker({})
    wire = Wire(client, model, store=ProposalJobStore(table_name="", memory={}))
    base: dict[str, Any] = {
        "owner_id": "synthetic",
        "job_id": "gentle",
        "started_at": datetime.now(UTC).isoformat(),
        "_search_session": {},
    }
    for i, query in enumerate(QUERIES[:2], 1):
        assert wire.search({**base, "query": query}, {}, f"gentle:search_{i}").actual == 12
    assert pauses == [30]
    assert len(client.requests) == 2
    apify.tiktok_search.assert_not_called()


@pytest.mark.parametrize("parent_deadline,maximum", [(35, 35), (300, 50)])
def test_empty_search_preserves_missing_and_observed_zero_and_deadline(
    apify: Mock, monkeypatch: pytest.MonkeyPatch, parent_deadline: int, maximum: int
) -> None:
    client = Worker({})
    original = client.search_tiktok_sales25(
        QUERIES[0], owner_id="synthetic", request_fingerprint="empty"
    )
    original["raw"].update(videos=[], status="observed_zero", error_code="TIKTOK_TRULY_EMPTY")
    original["acquisition_status"].update(actual=0, shortfall=100)
    items = actor_items(0, 3)
    items[0].pop("playCount")
    items[1]["playCount"] = 0
    items[2]["playCount"] = 42
    apify.tiktok_search.side_effect = None
    apify.tiktok_search.return_value = ApifyRunResult(
        items, ACTOR_TIKTOK_VIDEO, "run", "SUCCEEDED", 0.012
    )
    monkeypatch.setenv("TIKTOK_APIFY_FALLBACK_DEADLINE_S", "50")
    recovered = apply_sales25_fallback(
        original, media_client=client, deadline_s=parent_deadline, request_id="empty"
    )
    assert recovered["acquired_via"] == "apify"
    assert [p["stats"]["playCount"] for p in recovered["raw"]["videos"]] == [None, 0, 42]
    assert "stats.playCount" in recovered["raw"]["videos"][0]["missingFields"]
    assert "stats.playCount" not in recovered["raw"]["videos"][1]["missingFields"]
    assert recovered["acquisition_status"]["actual"] == 3
    assert recovered["acquisition_status"]["shortfall"] == 97
    assert 0 < apify.tiktok_search.call_args.kwargs["deadline_s"] <= maximum
    assert original["raw"]["videos"] == []


@pytest.mark.parametrize("code", ["TIKTOK_BOT_WALL", "TIKTOK_CDN_DENIED", "TIKTOK_API_ERROR"])
def test_full_target_and_unrelated_failure(apify: Mock, code: str) -> None:
    client = Worker({0: code})
    original = client.search_tiktok_sales25(
        QUERIES[0], owner_id="synthetic", request_fingerprint="full"
    )
    original["raw"]["stop_reason"] = "exception"
    apify.tiktok_search.side_effect = None
    apify.tiktok_search.return_value = ApifyRunResult(
        actor_items(0, 100), ACTOR_TIKTOK_VIDEO, "run", "SUCCEEDED", 0.4
    )
    recovered = apply_sales25_fallback(
        original, media_client=client, deadline_s=300, request_id="full"
    )
    assert recovered["acquisition_status"]["actual"] == 100
    assert recovered["acquisition_status"]["status"] == "complete"
    assert recovered["acquisition_status"]["shortfall"] == 0
    original["raw"].update(error_code="TIKTOK_EXCEPTION", stop_reason="exception")
    unrelated = apply_sales25_fallback(
        original, media_client=client, deadline_s=300, request_id="unrelated"
    )
    assert unrelated["acquisition_status"]["accepted"] is False
    assert apify.tiktok_search.call_count == 1


@pytest.mark.parametrize("failure", ["empty", "invalid_url", "over_target", "cost", "deadline"])
def test_invalid_or_unavailable_fallback_preserves_failure_once(
    apify: Mock, monkeypatch: pytest.MonkeyPatch, failure: str
) -> None:
    client = Worker({0: "TIKTOK_BOT_WALL"})
    original = client.search_tiktok_sales25(
        QUERIES[0], owner_id="synthetic", request_fingerprint=failure
    )
    items = actor_items(0, 3)
    if failure == "empty":
        items = []
    elif failure == "invalid_url":
        items[0]["webVideoUrl"] = "https://example.invalid/private"
    elif failure == "over_target":
        items *= 34
    apify.tiktok_search.side_effect = None
    apify.tiktok_search.return_value = ApifyRunResult(
        items, ACTOR_TIKTOK_VIDEO, "run", "SUCCEEDED", 0.012
    )
    if failure == "cost":
        apify.tiktok_search.side_effect = CostLimitExceededError("budget")
    for _ in range(2):
        recovered = apply_sales25_fallback(
            original,
            media_client=client,
            deadline_s=0 if failure == "deadline" else 300,
            request_id="failure",
        )
        assert recovered["acquisition_status"]["accepted"] is False
        assert recovered["raw"]["error_code"] == "TIKTOK_BOT_WALL"
        assert "APIFY_FALLBACK_FAILED:" in recovered["apify_fallback_warning"]
    assert apify.tiktok_search.call_count == (0 if failure == "deadline" else 1)


@pytest.mark.parametrize("retry_count", [13, 50])
def test_one_succeeded_actor_stops_at_thirteen_and_delivers_after_one_reserved_retry(
    monkeypatch: pytest.MonkeyPatch, retry_count: int
) -> None:
    """r81: five 50/50 queries, q2 SUCCEEDED + Scraped 0/1 + 13/50 after 221s."""
    monkeypatch.setenv("OMIYAGE_SALES25_SEARCH_DEPTH", "50")
    monkeypatch.setenv("USE_TIKTOK_APIFY_FALLBACK", "1")
    monkeypatch.setenv("TIKTOK_APIFY_FALLBACK_DEADLINE_S", "240")
    elapsed = [0.0]
    monkeypatch.setattr(
        tiktok_search_fallback,
        "time",
        SimpleNamespace(time=lambda: datetime.now(UTC).timestamp(), monotonic=lambda: elapsed[0]),
    )
    fake = _FakeApify()
    ledger = _ReservationLedger()
    ledger.token = SimpleNamespace(global_reserved=True)
    reservations = Mock(wraps=ledger.reserve)
    monkeypatch.setattr(ledger, "reserve", reservations)
    monkeypatch.setattr(tiktok_search_fallback.CostGuard, "from_env", lambda: ledger)
    queries: list[str] = []

    def handler(http_request: httpx.Request) -> httpx.Response:
        if http_request.method == "POST" and "/acts/" in http_request.url.path:
            body = json.loads(http_request.content)
            query = body["searchQueries"][0]
            queries.append(query)
            count = 50
            if query == QUERIES[2]:
                count = 13 if queries.count(query) == 1 else retry_count
                if queries.count(query) == 1:
                    elapsed[0] += 221
            fake.items_by_actor[ACTOR_TIKTOK_VIDEO] = actor_items(QUERIES.index(query), count)
            fake.status_message = f"Scraped {0 if count == 13 else 1}/1 search queries"
        return fake.handler(http_request)

    actor = ApifyClient(
        "synthetic",
        ledger=ledger,
        http=httpx.Client(transport=httpx.MockTransport(handler)),
        poll_interval_s=0,
    )
    search = Mock(wraps=actor.tiktok_search)
    monkeypatch.setattr(actor, "tiktok_search", search)
    monkeypatch.setattr(ApifyClient, "from_env", lambda **kwargs: actor)
    client = Worker({0: "TIKTOK_BOT_WALL"})
    store = ProposalJobStore(table_name="", memory={})
    approved = model()
    wire = Wire(client, lambda: approved, store=store)
    accepted = build(store, sales25_stages=production_stages(wire)).run(request(), ctx())
    saved = result(store, accepted.job_id)
    assert saved.delivery_status == "pending", saved.model_dump_json(indent=2)
    assert saved.checkpoints["relevance"].status == "complete"
    third = saved.checkpoints["search_3"]
    assert third.actual == retry_count
    assert third.status == ("partial" if retry_count == 13 else "complete")
    assert third.reason == ("shortfall" if retry_count == 13 else None)
    assert [saved.checkpoints[f"search_{i}"].actual for i in (1, 2, 4, 5, 6)] == [50] * 5
    case = {
        "owner_id": third.refs[0]["owner_id"],
        "job_id": accepted.job_id,
        "started_at": store.get_job(accepted.job_id)["created_at"],
    }
    payload = wire._json(case, third.refs[0])
    status, raw = payload["acquisition_status"], payload["raw"]
    assert status["accepted"] is True
    assert status["target"] == raw["requested"] == 50
    assert status["actual"] == raw["actual"] == len(raw["videos"]) == retry_count
    assert status["shortfall"] == 50 - retry_count
    assert len(payload["apify_attempts"]) == 2
    assert queries == [*QUERIES[:3], QUERIES[2], *QUERIES[3:]]
    assert ledger.reserved == [0.2] * 7
    assert search.call_args_list[3].kwargs["deadline_s"] == 19
    assert len(ledger.settled) == 7
    reservation_ids = [call.kwargs["request_id"] for call in reservations.call_args_list]
    assert len(set(reservation_ids)) == 7 and reservation_ids[3].endswith(":retry")
    # The final relevance stage consumes the short source too; model evidence is cached by input.
    batch_sizes = [
        len(json.loads(call.args[0].split("Input:\n", 1)[1])["posts"])
        for call in approved.generate_text.call_args_list
        if len(call.args) > 1 and call.args[1].startswith("sales25-relevance-q2-")
    ]
    assert (13 in batch_sizes) if retry_count == 13 else (batch_sizes == [25, 25])
    # Replay uses the selected immutable result and never reserves a third run.
    replay = apply_sales25_fallback(
        payload["worker_evidence"], media_client=client, deadline_s=240, request_id="replay"
    )
    assert replay["raw"] == raw and len(ledger.reserved) == 7
    slack = FakeSlack(store, accepted.job_id)
    parent = sales25.Parent(store, reader=client.download, slack=slack)
    context = ctx()
    output = OmiyageReportStatusSkill(store=store, sales25_parent=parent).run(
        OmiyageReportStatusInput(job_id=accepted.job_id), context
    )
    monkeypatch.setattr(notify, "ProposalJobStore", lambda: store)
    assert notify.publish_notice(
        output.message,
        origin=context.metadata[ORIGIN_KEY],
        request_id="r81",
        job_id=accepted.job_id,
    )
    assert result(store, accepted.job_id).delivery_status == "delivered"
    if retry_count == 13:
        assert saved.status == "partial"
        assert "一部の数字が足りません" in slack.uploads[0]["initial_comment"]
        assert "50本目標→13本取得" in slack.uploads[0]["initial_comment"]


@pytest.mark.parametrize("failure", ["cost", "timeout", "invalid", "smaller", "deadline", "cache"])
def test_shortfall_retains_first_valid_session_when_retry_cannot_improve(
    apify: Mock, monkeypatch: pytest.MonkeyPatch, failure: str
) -> None:
    client = Worker({0: "TIKTOK_BOT_WALL"})
    original = client.search_tiktok_sales25(
        QUERIES[0], owner_id="synthetic", request_fingerprint=failure
    )
    elapsed = [0.0]
    monkeypatch.setattr(
        tiktok_search_fallback,
        "time",
        SimpleNamespace(time=lambda: 1000, monotonic=lambda: elapsed[0]),
    )

    def search(*args: Any, **kwargs: Any) -> ApifyRunResult:
        if apify.tiktok_search.call_count == 1:
            if failure == "deadline":
                elapsed[0] = 300
            items = actor_items(0, 13)
        elif failure == "cost":
            raise CostLimitExceededError("budget")
        elif failure == "timeout":
            raise TimeoutError("private URL")
        else:
            items = actor_items(0, 3)
            if failure == "invalid":
                items[0]["webVideoUrl"] = "https://example.invalid/private"
        return ApifyRunResult(items, ACTOR_TIKTOK_VIDEO, "run", "SUCCEEDED", 0.052)

    apify.tiktok_search.side_effect = search
    if failure == "cache":
        stage = client.stage_bytes

        def stage_bytes(**kwargs: Any) -> S3ObjectRef:
            if kwargs["name"] == "sales25-apify.json":
                raise TimeoutError("cache expired")
            return stage(**kwargs)

        monkeypatch.setattr(client, "stage_bytes", stage_bytes)
    recovered = apply_sales25_fallback(
        original, media_client=client, deadline_s=240, request_id=failure
    )
    assert recovered["acquired_via"] == "apify"
    assert recovered["acquisition_status"]["accepted"] is True
    assert recovered["acquisition_status"]["actual"] == len(recovered["raw"]["videos"]) == 13
    assert apify.tiktok_search.call_count == (1 if failure == "deadline" else 2)
    if failure in {"cost", "timeout", "invalid"}:
        assert "APIFY_RETRY_FAILED:" in recovered["apify_retry_warning"]


@pytest.mark.parametrize(
    "updates",
    [{}, {"requested": 12}, {"requested": True}, {"actual": 50}, {"actual": True}],
)
def test_partial_apify_contract_requires_consistent_counts(
    apify: Mock, updates: dict[str, Any]
) -> None:
    client = Worker({0: "TIKTOK_BOT_WALL"})
    original = client.search_tiktok_sales25(
        QUERIES[0], owner_id="synthetic", request_fingerprint="contract"
    )
    apify.tiktok_search.side_effect = None
    apify.tiktok_search.return_value = ApifyRunResult(
        actor_items(0, 13), ACTOR_TIKTOK_VIDEO, "run", "SUCCEEDED", 0.052
    )
    raw = apply_sales25_fallback(
        original, media_client=client, deadline_s=240, request_id="contract"
    )["raw"]
    raw.update(updates)
    decisions = [dict(id=p["id"], eligible=True, reason="合成") for p in raw["videos"]]
    if updates:
        with pytest.raises(ValueError, match="Apify requested/actual"):
            contract._source({"query": QUERIES[0]}, raw, decisions, "synthetic")
    else:
        source = contract._source({"query": QUERIES[0]}, raw, decisions, "synthetic")
        assert len(source["posts"]) == 13
