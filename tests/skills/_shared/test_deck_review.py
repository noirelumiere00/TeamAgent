"""合成 PPTX で shape 境界・テンプレ差分・表示順を固定する。"""

from io import BytesIO

from pptx import Presentation
from pptx.util import Inches

from teamagent.skills._shared.deck_review import count_review, review_slides, warning_lines


def _deck(texts: list[list[str]]) -> bytes:
    prs = Presentation()
    for shapes in texts:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for text in shapes:
            slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = text
    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()


def test_shape_boundary_split_runs_and_joined_paragraph_text() -> None:
    prs = Presentation(BytesIO(_deck([["重要", "確認"], [], ["■概要要確認（データ未検出）"]])))
    shape = prs.slides[1].shapes.add_textbox(0, 0, Inches(4), Inches(1))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.add_run().text = "要"
    paragraph.add_run().text = "確認（出典URL未取得）"
    stream = BytesIO()
    prs.save(stream)
    assert count_review(stream.getvalue()) == {1: 0, 2: 1, 3: 1}


def test_template_counts_are_subtracted_per_slide() -> None:
    template = _deck([["要確認"], ["要確認"], []])
    output = _deck([["要確認"], ["要確認 要確認"], ["要確認"]])
    assert review_slides(output, template) == [2, 3]


def test_notes_excluded_hidden_slide_and_display_order_counted() -> None:
    prs = Presentation(BytesIO(_deck([[], ["要確認"], ["要確認"]])))
    prs.slides[0].notes_slide.notes_text_frame.text = "要確認"
    prs.slides[1]._element.set("show", "0")
    ids = prs.slides._sldIdLst
    ids.insert(0, ids[-1])
    stream = BytesIO()
    prs.save(stream)
    assert count_review(stream.getvalue()) == {1: 1, 2: 0, 3: 1}


def test_warning_limit_empty_and_unknown() -> None:
    assert warning_lines([]) == []
    assert warning_lines(None) == ["要確認の位置は自動で数えられませんでした"]
    assert warning_lines([63, 7, 7]) == [
        "PowerPoint の左の一覧で 7・63枚目に『要確認』が残っています（確かめてから使ってください）"
    ]
    assert warning_lines(list(range(12, 0, -1)))[0] == (
        "PowerPoint の左の一覧で 1・2・3・4・5・6・7・8・9・10枚目ほか2枚に『要確認』が残っています"
        "（確かめてから使ってください）"
    )


def test_tables_and_group_shapes_are_counted() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, 0, 0, Inches(4), Inches(1)).table
    table.cell(0, 0).text = "要確認"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "要確認"
    stream = BytesIO()
    prs.save(stream)
    assert count_review(stream.getvalue()) == {1: 2}
