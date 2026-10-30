"""外部接続なしで、MCP受付→自動調査/JSON→Builder→両レンダラの契約を通す。"""

import json
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Inches

from teamagent.adapters.proposal_job_store import ProposalJobStore
from teamagent.media.render_child import _proposal
from teamagent.skills.base import SkillContext
from teamagent.skills.proposal_builder.schema import (
    ProposalBuilderStatusInput,
    ProposalBuilderSubmitInput,
)
from teamagent.skills.proposal_builder.skill import (
    ProposalBuilderSkill,
    ProposalBuilderStatusSkill,
    ProposalBuilderSubmitSkill,
)
from teamagent.skills.proposal_deck.contract import VALID_IDS
from teamagent.skills.proposal_deck.renderer import render_deck
from teamagent.skills.proposal_deck.schema import ProposalDeckInput, ProposalDeckOutput
from teamagent.skills.proposal_deck.skill import ProposalDeckSkill
from tests.skills.proposal_builder.test_auto_research import _ctx, _immediate, _result
from tests.skills.proposal_builder.test_tiktok_enrichment import (
    _CapturingDeck,
    _configure_builder_test,
)
from tests.skills.proposal_deck.test_bu2_shape import _pin_template, _template


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("automatic", [False, True])
@pytest.mark.parametrize(
    "date_input", [{}, {"posting_start_date": None}, {"posting_start_date": "2027-01-10"}]
)
def test_submit_to_generation_preserves_optional_date(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    backend: str,
    automatic: bool,
    date_input: dict[str, Any],
) -> None:
    _configure_builder_test(monkeypatch, media_configured=False)
    monkeypatch.setenv("PROPOSAL_BUILDER_ALLOWED_EMAILS", "*")
    monkeypatch.setenv("PROPOSAL_RESEARCH_AUTO", "1")
    template = tmp_path / "template.pptx"
    _template(template)
    prs = Presentation(str(template))
    prs.slides[4].shapes.add_textbox(
        0, 0, Inches(9), Inches(1)
    ).text = "配信期間: {{PB-DATE:-56:%Y年%m月%d日}}～{{PB-DATE:21:%Y年%m月%d日}}"
    prs.slides[4].shapes.add_textbox(
        0, Inches(1), Inches(9), Inches(1)
    ).text = "投稿月: {{PB-MONTH}}"
    prs.save(str(template))
    _pin_template(monkeypatch, template)
    monkeypatch.setenv("PROPOSAL_BUILDER_TEMPLATE_PATH", str(template))
    output = tmp_path / "output.pptx"

    class Research:
        def run(self, input: Any, ctx: SkillContext) -> Any:
            return _result()

    class RenderingDeck(_CapturingDeck):
        def run(self, input: ProposalDeckInput, ctx: SkillContext) -> ProposalDeckOutput:
            # 調査本文の生成は対象外。95枠の不足契約を通し、日付メタデータは実経路で渡す。
            data = {
                "placeholders": {
                    1: "投稿開始 2050/12/31",
                    47: "「朝の余白」— 朝の気持ちに寄り添う。\n「夜のご褒美」— 夜の習慣に結びつく。\n「手軽な習慣」— 続ける気持ちを後押しする。",
                },
                "skipped_placeholders": [
                    {"id": pid, "reason": "要確認（データ未検出）"}
                    for pid in sorted(VALID_IDS - {1, 47})
                ],
                "posting_start_date": "未定" if backend == "local" else "",
                "auxiliary_placeholders": {"PB-MONTH": "2050年12月"},
            }
            composer = ProposalDeckSkill._finalize_composer(data, input)
            if backend == "local":
                render_deck(composer, template, output)
            else:
                (tmp_path / "composer.json").write_text(composer.model_dump_json())
                _proposal(
                    tmp_path,
                    {
                        "kind": "proposal_pptx",
                        "template": "template.pptx",
                        "composer": "composer.json",
                        "evidence": [],
                        "output": "output.pptx",
                        "fail_if_missing": True,
                    },
                )
            return super().run(input, ctx).model_copy(update={"pptx_path": str(output)})

    deck = RenderingDeck()
    builder = ProposalBuilderSkill(
        search=object(),
        deck=deck,  # type: ignore[arg-type]
        account_db_path="unused.xlsx",
    )
    memory: dict[str, dict[str, Any]] = {}
    store = ProposalJobStore(table_name="", memory=memory)
    request = ProposalBuilderSubmitInput.model_validate(
        {
            **(
                {"research_brief": {"product_name": "ACME"}}
                if automatic
                else {"gemini_json": _result().research_json}
            ),
            "proposal_date": "2026-10-22",
            "plans": [{"name": "Trial", "period": "4週間", "assumption": "依頼文: 4週間"}],
            **date_input,
        }
    )
    ctx = _ctx()
    ctx.metadata["request_text"] = "期間は4週間"
    accepted = ProposalBuilderSubmitSkill(
        builder_factory=lambda: builder,
        research_factory=Research,
        store=store,
        thread_launcher=_immediate,
        heartbeat_seconds=0,
    ).run(request, ctx)
    state = ProposalBuilderStatusSkill(store=store).run(
        ProposalBuilderStatusInput(job_id=accepted.job_id), _ctx()
    )
    assert accepted.status == "queued"
    assert state.status == "done", state.message
    assert output.is_file()
    assert deck.inputs[0].posting_start_date == request.posting_start_date
    assert json.loads(memory[accepted.job_id]["request_summary"])[
        "posting_start_date"
    ] == date_input.get("posting_start_date")
    assert deck.inputs[0].proposal_date == date(2026, 10, 22)
    generated = Presentation(str(output))
    assert len(generated.slides) == 68
    texts = [
        shape.text for slide in generated.slides for shape in slide.shapes if shape.has_text_frame
    ]
    period = next(text for text in texts if text.startswith("配信期間:"))
    month = next(text for text in texts if text.startswith("投稿月:"))
    if request.posting_start_date is None:
        assert period == "配信期間: 未定～未定"
        assert month == "投稿月: 未定"
        assert "未定" in deck.inputs[0].deadline
        assert "2026年" not in deck.inputs[0].deadline
    else:
        start = request.posting_start_date
        assert (
            period
            == f"配信期間: {(start - timedelta(days=56)):%Y年%m月%d日}～{(start + timedelta(days=21)):%Y年%m月%d日}"
        )
        assert month == "投稿月: 2027年01月"
    combined = "\n".join(texts)
    if request.posting_start_date is None:
        assert "2050" not in combined
        assert "投稿開始 未定" in combined
    assert "2026年10月22日" in combined
    assert "4週間" in combined
    assert "{{PB-" not in combined
