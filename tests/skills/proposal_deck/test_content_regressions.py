"""10-30実機で見つかった本文崩れの回帰。外部接続なし。"""

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from teamagent.media.proposal_quality import (
    proposal_field_results,
    proposal_field_spec,
    update_proposal_image_results,
)
from tests.skills.proposal_deck.test_template_cleanup import _render


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize(
    "suffix", ["PR×ショート動画施策ご提案書", "ご提案書 PR × ショート動画施策　\n"]
)
def test_cover_strips_repeated_suffix(tmp_path: Path, backend: str, suffix: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(
        0, 0, Inches(8), Inches(3)
    ).text = "｛1：タイトル｝ご提案書\nPR×ショート動画施策"
    prs.save(str(tmp_path / "template.pptx"))
    result: Any = _render(backend, tmp_path, {1: f"合成ブランド {suffix}"})
    assert "".join(result.slides[0].shapes[0].text.splitlines()) == (
        "合成ブランドご提案書PR×ショート動画施策"
    )


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_fixed_schedule_and_current_box_keep_body(tmp_path: Path, backend: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    texts = [
        "構成案 提出 → FIX（要確認）",
        "マスター動画初稿・投稿文 提出（要確認）",
        "一般ユーザー・インフルエンサーによる\nリアリティのある“体感コメント”\n「○○おすすめ5選」「正直レビュー」といった\nUGC型コンテンツ\nSNS上でのUGCが興味喚起～検討初期の材料として強く機能",
    ]
    for index, text in enumerate(texts):
        slide.shapes.add_textbox(0, Inches(index), Inches(8), Inches(1)).text = text
    prs.save(str(tmp_path / "template.pptx"))
    result: Any = _render(backend, tmp_path, {})
    assert [shape.text for shape in result.slides[0].shapes][:2] == [
        text.replace("（要確認）", "（クライアント確認）") for text in texts[:2]
    ]
    current = result.slides[0].shapes[2].text
    assert "UGC" in current and "○○" not in current


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("body_length,warn", [(500, False), (501, True)])
def test_500_limit_counts_only_authored_body(
    tmp_path: Path, backend: str, body_length: int, warn: bool
) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(8), Inches(3)).text = "固定文" * 200 + "｛96｝"
    slide.notes_slide.notes_text_frame.text = "既存メモ" * 200
    prs.save(str(tmp_path / "template.pptx"))
    result: Any = _render(backend, tmp_path, {96: "本" * body_length})
    notes = result.slides[0].notes_slide.notes_text_frame.text
    assert ("警告: 本文500字超過" in notes) is warn


def test_pink_frames_outside_image_slots_warn_without_blocking() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    for parent in (slide, group):
        shape = parent.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 51, 153)
        shape.text = "※該当記事の引用（画像/テキスト）"
    spec = proposal_field_spec(prs)
    assert spec["images"] == []
    metadata = proposal_field_results(prs, spec, {}, {})
    update_proposal_image_results(metadata, requested_images=0)
    assert metadata["field_issues"] == []
    assert metadata["field_warnings"] == ["要対応: 画像の仮枠が残っています 2 枠（1枚目）"]


def test_integrated_pr_candidates_require_three_and_retry() -> None:
    import json
    from unittest.mock import MagicMock

    from teamagent.media.proposal_template import BASE_AUXILIARY_KEYS
    from teamagent.skills.base import SkillContext
    from teamagent.skills.proposal_deck.schema import ProposalDeckInput
    from teamagent.skills.proposal_deck.skill import ProposalDeckSkill
    from tests.skills.proposal_deck.test_proposal_deck_skill import _full_composer_json, _resp

    data = json.loads(_full_composer_json())
    data["placeholders"].pop("47")
    data["skipped_placeholders"] = [{"id": 47, "reason": "要確認（データ未検出）"}]
    repaired = {
        **data,
        "placeholders": {
            **data["placeholders"],
            "47": "「朝の余白」— 朝の気持ちに寄り添う。\n「夜のご褒美」— 夜の習慣に結びつく。\n「手軽な習慣」— 続ける気持ちを後押しする。",
        },
        "skipped_placeholders": [],
    }
    bedrock = MagicMock(model_id="jp.anthropic.claude-sonnet-4-6")
    bedrock.converse.side_effect = [_resp(json.dumps(data)), _resp(json.dumps(repaired))]
    output, _, _ = ProposalDeckSkill(bedrock=bedrock)._compose(
        ProposalDeckInput(
            product_name="合成商材",
            goal="認知",
            target_persona="成人",
            template_profile="proposal-builder-v1",
            auxiliary_placeholders={key: "設定値" for key in BASE_AUXILIARY_KEYS},
            max_repair=1,
        ),
        SkillContext(),
    )
    assert bedrock.converse.call_count == 2
    assert output.placeholders[47] == repaired["placeholders"]["47"]


def test_account_description_does_not_redact_idiom_but_redacts_metrics() -> None:
    from teamagent.skills.proposal_builder.research import redact_unverified_quantities

    text, changed = redact_unverified_quantities(
        "何気ない毎日をもう一度見る。二度と見逃さない。月間10万回再生。"
    )
    assert "もう一度見る" in text and "二度と見逃さない" in text
    assert changed and "10万回" not in text


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_narrow_fields_do_not_keep_model_line_breaks(tmp_path: Path, backend: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for index, pid in enumerate((7, 9, 11, 47)):
        slide.shapes.add_textbox(0, Inches(index), Inches(8), Inches(1)).text = f"｛{pid}｝"
    prs.save(str(tmp_path / "template.pptx"))
    rendered: Any = _render(
        backend,
        tmp_path,
        {
            7: "朝の\n余白",
            9: "夜の\nご褒美",
            11: "手軽な\n習慣",
            47: "「朝の余白」— 朝の気持ちに寄り添う。\n「夜のご褒美」— 夜の習慣に結びつく。\n「手軽な習慣」— 続ける気持ちを後押しする。",
        },
    )
    assert [shape.text.replace("\v", "\n") for shape in rendered.slides[0].shapes] == [
        "朝の 余白",
        "夜の ご褒美",
        "手軽な 習慣",
        "「朝の余白」— 朝の気持ちに寄り添う。\n「夜のご褒美」— 夜の習慣に結びつく。\n「手軽な習慣」— 続ける気持ちを後押しする。",
    ]


@pytest.mark.parametrize(
    "text", ["一度の温度差", "二度の温度差", "十一度の温度差", "二十一度の室温", "3度の温度差"]
)
def test_actual_temperatures_are_still_redacted(text: str) -> None:
    from teamagent.skills.proposal_builder.research import redact_unverified_quantities

    actual, changed = redact_unverified_quantities(text)
    assert changed and "要確認（出典URL未取得）" in actual


def test_length_fallback_keeps_quoted_names_instead_of_cutting_mid_sentence() -> None:
    from teamagent.skills.proposal_deck.skill import _fit_label

    # 10-30 実機: 見出し・他候補が説明文ごと来て、上限で切ると文の途中で終わった。
    assert _fit_label("「実家のような安心感」——戻ってきた人が語る", 20) == "「実家のような安心感」"
    assert _fit_label("①「毎日できたて」——日常語と組み合わせ", 19) == "「毎日できたて」"
    assert _fit_label("案名だけの長い見出しテキストです", 5) == "案名だけの"
    assert (
        _fit_label("「とても長い案名がここに入っていて収まらない」", 10) == "「とても長い案名がこ"
    )
