"""Text filling must not alter preset/custom decorations or group transforms."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches

from teamagent.media.render_child import _proposal
from teamagent.skills.proposal_deck.renderer import render_pptx


def _template(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text = "{1:本文}"
    star = slide.shapes.add_shape(
        MSO_SHAPE.STAR_5_POINT, Inches(1), Inches(1), Inches(0.4), Inches(0.6)
    )
    star.rotation = 27
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(MSO_SHAPE.MATH_MULTIPLY, Inches(2), Inches(1), Inches(0.3), Inches(0.4))
    nested = group.shapes.add_group_shape()
    arrow = nested.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(3), Inches(1), Inches(0.6), Inches(0.2)
    )
    arrow.rotation = 315
    custom = nested.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(4), Inches(1), Inches(0.4), Inches(0.5)
    )
    geometry = custom._element.spPr.prstGeom
    geometry.getparent().replace(
        geometry,
        parse_xml(
            f"<a:custGeom {nsdecls('a')}><a:avLst/><a:gdLst/><a:ahLst/>"
            '<a:cxnLst/><a:rect l="l" t="t" r="r" b="b"/>'
            '<a:pathLst><a:path w="100" h="100"><a:moveTo><a:pt x="0" y="0"/>'
            '</a:moveTo><a:lnTo><a:pt x="100" y="50"/></a:lnTo><a:lnTo>'
            '<a:pt x="0" y="100"/></a:lnTo><a:close/></a:path></a:pathLst></a:custGeom>'
        ),
    )
    # Non-identity child coordinates expose accidental group-bound recalculation.
    transform = group._element.grpSpPr.xfrm
    transform.set("flipH", "1")
    transform.set("rot", "900000")
    transform.chOff.x, transform.chOff.y = 1000, 2000
    transform.chExt.cx, transform.chExt.cy = 5000000, 2000000
    prs.save(path)


def _geometry(path: Path) -> dict[tuple[int, ...], tuple[Any, ...]]:
    result = {}

    def walk(shapes: Any, parent: tuple[int, ...]) -> None:
        for shape in shapes:
            key = (*parent, shape.shape_id)
            properties = shape._element.xpath("./p:spPr | ./p:grpSpPr")
            result[key] = (
                shape.left,
                shape.top,
                shape.width,
                shape.height,
                shape.rotation,
                tuple(etree.tostring(node, method="c14n") for node in properties),
            )
            if hasattr(shape, "shapes"):
                walk(shape.shapes, key)

    for number, slide in enumerate(Presentation(path).slides, 1):
        walk(slide.shapes, (number,))
    return result


@pytest.mark.parametrize("renderer", ["local", "worker"])
def test_render_preserves_decoration_geometry(tmp_path: Path, renderer: str) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "rendered.pptx"
    _template(template)
    original = _geometry(template)
    value = "生成された本文です。" * 10
    if renderer == "local":
        render_pptx(template, {1: value}, output)
    else:
        (tmp_path / "composer.json").write_text(
            json.dumps({"placeholders": {"1": value}}), encoding="utf-8"
        )
        _proposal(
            tmp_path,
            {
                "kind": "proposal_pptx",
                "template": template.name,
                "composer": "composer.json",
                "evidence": [],
                "output": output.name,
                "fail_if_missing": True,
            },
        )
    assert _geometry(output) == original
    assert Presentation(output).slides[0].shapes[0].text == value


@pytest.mark.parametrize("renderer", ["local", "worker"])
def test_render_clamps_only_roundrect_adjustment(tmp_path: Path, renderer: str) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "rendered.pptx"
    _template(template)
    prs = Presentation(template)
    for index, adjustment in enumerate((239808, 1204822, -100, 0, 16667, 50000)):
        shape = prs.slides[0].shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(index), Inches(3), Inches(0.4), Inches(0.3)
        )
        shape.adjustments[0] = adjustment / 100000
    prs.save(template)
    original = _geometry(template)
    if renderer == "local":
        render_pptx(template, {1: "本文"}, output)
    else:
        (tmp_path / "composer.json").write_text(
            json.dumps({"placeholders": {"1": "本文"}}), encoding="utf-8"
        )
        _proposal(
            tmp_path,
            {
                "kind": "proposal_pptx",
                "template": template.name,
                "composer": "composer.json",
                "evidence": [],
                "output": output.name,
                "fail_if_missing": True,
            },
        )
    after = _geometry(output)
    assert {key: data[:5] for key, data in after.items()} == {
        key: data[:5] for key, data in original.items()
    }
    values = [
        shape._element.xpath("./p:spPr/a:prstGeom/a:avLst/a:gd")[0].get("fmla")
        for shape in list(Presentation(output).slides[0].shapes)[-6:]
    ]
    assert values == ["val 50000", "val 50000", "val 0", "val 0", "val 16667", "val 50000"]
    assert {key: value for key, value in after.items() if len(key) > 2} == {
        key: value for key, value in original.items() if len(key) > 2
    }
