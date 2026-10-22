"""The 86-page extension is SHA-pinned and shares selection/empty-value semantics."""

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from teamagent.adapters.proposal_assets import ProposalAssetProvisionError, _validate_template
from teamagent.media import proposal_template
from teamagent.media.operations import MediaOperationError
from teamagent.media.proposal_bu2 import BU2Content
from teamagent.media.render_child import _proposal
from teamagent.skills.proposal_deck.contract import ComposerOutput
from teamagent.skills.proposal_deck.renderer import UnfilledPlaceholderError, render_deck
from tests.skills.proposal_deck.test_bu2_content import bu2_payload
from tests.skills.proposal_deck.test_bu2_shape import _pin_template, _render, _template
from tests.skills.proposal_deck.test_proposal_deck_skill import _full_composer_json


def _extension(path: Path) -> None:
    _template(path)
    prs = Presentation(path)
    for prefix in ("PB-BU2-PLAN-", "PB-BU2-SIM-", "PB-BU2-HEARING"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        keys = {key for key in proposal_template.BU2_AUXILIARY_KEYS if key.startswith(prefix)}
        if prefix == "PB-BU2-PLAN-":
            keys.add("PB-BU2-ESTIMATE-NOTES")
        for index, key in enumerate(sorted(keys)):
            slide.shapes.add_textbox(
                0, Inches(index / 4), Inches(9), Inches(0.25)
            ).text = f"{{{{{key}}}}}"
    prs.save(path)


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("unit,expected", [("BU1", 83), ("BU2", 72)])
def test_extension_pages_are_bu2_only_and_missing_values_go_to_notes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, backend: str, unit: str, expected: int
) -> None:
    path = tmp_path / "template.pptx"
    _extension(path)
    _pin_template(monkeypatch, path, 86)
    before = path.read_bytes()
    output = _render(tmp_path, backend, business_unit=unit)
    prs = Presentation(output)
    assert len(prs.slides) == expected
    if unit == "BU2":
        for slide in list(prs.slides)[-3:]:
            assert all(
                shape.text.strip() in {"", "※税・管理費別"}
                for shape in slide.shapes
                if shape.has_text_frame
            )
            assert "要確認" in slide.notes_slide.notes_text_frame.text
            assert "PB-BU2-" in slide.notes_slide.notes_text_frame.text
    assert path.read_bytes() == before


@pytest.mark.parametrize("backend", ["local", "worker", "assets"])
@pytest.mark.parametrize("slide_count", [83, 86])
def test_unknown_sha_cannot_enable_template(tmp_path: Path, backend: str, slide_count: int) -> None:
    path = tmp_path / "template.pptx"
    (_extension if slide_count == 86 else _template)(path)
    with pytest.raises(
        (UnfilledPlaceholderError, MediaOperationError, ProposalAssetProvisionError)
    ):
        _validate_template(path) if backend == "assets" else _render(tmp_path, backend)
    assert not (tmp_path / "output.pptx").exists()


@pytest.mark.parametrize("backend", ["local", "worker", "assets"])
def test_edited_extension_bytes_do_not_reuse_a_valid_sha(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, backend: str
) -> None:
    path = tmp_path / "template.pptx"
    _extension(path)
    _pin_template(monkeypatch, path, 86)
    with path.open("ab") as output:
        output.write(b"changed after review")
    with pytest.raises(
        (UnfilledPlaceholderError, MediaOperationError, ProposalAssetProvisionError)
    ):
        _validate_template(path) if backend == "assets" else _render(tmp_path, backend)
    assert not (tmp_path / "output.pptx").exists()


def test_asset_validator_accepts_old_and_pinned_new_inventory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "template.pptx"
    _template(path)
    _pin_template(monkeypatch, path)
    _validate_template(path)
    _extension(path)
    _pin_template(monkeypatch, path, 86)
    _validate_template(path)


@pytest.mark.parametrize("backend", ["local", "worker", "assets"])
@pytest.mark.parametrize("damage", ["auxiliary", "numeric", "page"])
def test_pinned_extension_still_requires_its_exact_inventory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, backend: str, damage: str
) -> None:
    path = tmp_path / "template.pptx"
    _extension(path)
    prs = Presentation(path)
    if damage == "auxiliary":
        prs.slides[-1].shapes[0].text = "{{PB-UNKNOWN}}"
    elif damage == "numeric":
        prs.slides[2].shapes[0].text += " {104}"
    else:
        slide = prs.slides._sldIdLst[-1]
        prs.part.drop_rel(slide.rId)
        prs.slides._sldIdLst.remove(slide)
    prs.save(path)
    _pin_template(monkeypatch, path, 86)
    with pytest.raises(
        (UnfilledPlaceholderError, MediaOperationError, ProposalAssetProvisionError)
    ):
        _validate_template(path) if backend == "assets" else _render(tmp_path, backend)


@pytest.mark.parametrize("partial_terms", [False, True])
@pytest.mark.parametrize("hearing_count", [4, 5, 6])
def test_structured_values_and_source_notes_match_in_both_renderers(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, partial_terms: bool, hearing_count: int
) -> None:
    path = tmp_path / "template.pptx"
    _extension(path)
    _pin_template(monkeypatch, path, 86)
    payload = json.loads(_full_composer_json())
    payload.update(
        **bu2_payload(),
        template_profile="proposal-builder-v1",
        posting_start_date="2026-12-20",
        proposal_date="2026-10-22",
        product_sector="美容",
        auxiliary_placeholders={
            **{key: "設定値" for key in proposal_template.BASE_AUXILIARY_KEYS},
            "PB-BU2-PLAN-TRIAL-AMOUNT": "入力を無視した金額",
        },
    )
    payload["hearing_items"].extend(
        ["素材について伺いたい", "承認者について伺いたい"][: hearing_count - 4]
    )
    if partial_terms:
        payload["estimate_terms"] = {"net_pricing": True}
    else:
        payload["plans"].extend({**payload["plans"][0], "name": name} for name in ("A", "B", "C"))
    render_deck(ComposerOutput.model_validate(payload), path, tmp_path / "local.pptx")
    (tmp_path / "composer.json").write_text(json.dumps(payload), encoding="utf-8")
    _proposal(
        tmp_path,
        {
            "kind": "proposal_pptx",
            "template": "template.pptx",
            "composer": "composer.json",
            "evidence": [],
            "output": "worker.pptx",
            "fail_if_missing": True,
        },
    )
    contents = []
    for backend in ("local", "worker"):
        added = list(Presentation(tmp_path / f"{backend}.pptx").slides)[-3:]
        text = "\n".join(shape.text for slide in added for shape in slide.shapes)
        notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in added)
        contents.append((text.replace("\v", "\n"), notes))
        assert "250,003 円" in text
        assert "25,000人（推定）" in text
        assert "15,000人（試算）" in text
        assert "1. 目的とKPIを確認したい" in text
        hearing_texts = [shape.text.strip() for shape in added[-1].shapes]
        assert len(hearing_texts) == hearing_count
        if hearing_count < 6:
            assert f"入力{hearing_count}件のため6番枠未使用" in notes
        assert "入力を無視した金額" not in text
        assert "※ネット表記" in text
        assert "※税・管理費別" in text
        assert ("※管理費 10% 別途" in text) is not partial_terms
        assert "合成の料金条件" in notes
        assert "入力条件・補足（自動生成）" in notes
        assert ("要確認" in notes) is partial_terms
        assert "要確認" not in text
    assert contents[0] == contents[1]


@pytest.mark.parametrize(
    "request_text",
    [
        "",
        "本数３本、再生１２万回、単価８円/回、期間６週間、管理費１５%、接触７回、理解率８０%",
    ],
    ids=["missing-request", "unmatched-values"],
)
def test_ungrounded_bu2_values_are_blank_with_review_notes_in_both_renderers(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, request_text: str
) -> None:
    path = tmp_path / "template.pptx"
    _extension(path)
    _pin_template(monkeypatch, path, 86)
    supplied = bu2_payload()
    supplied["plans"][0]["assumption"] = "料金は900000円、単価2.5円で計算"
    grounded = BU2Content.model_validate(supplied).grounded_in_request(request_text)
    payload = json.loads(_full_composer_json())
    payload.update(
        **grounded.model_dump(mode="json"),
        template_profile="proposal-builder-v1",
        posting_start_date="2026-12-20",
        proposal_date="2026-10-22",
        product_sector="美容",
        auxiliary_placeholders={key: "設定値" for key in proposal_template.BASE_AUXILIARY_KEYS},
    )
    render_deck(ComposerOutput.model_validate(payload), path, tmp_path / "local.pptx")
    (tmp_path / "composer.json").write_text(json.dumps(payload), encoding="utf-8")
    _proposal(
        tmp_path,
        {
            "kind": "proposal_pptx",
            "template": "template.pptx",
            "composer": "composer.json",
            "evidence": [],
            "output": "worker.pptx",
            "fail_if_missing": True,
        },
    )
    contents = []
    for backend in ("local", "worker"):
        added = list(Presentation(tmp_path / f"{backend}.pptx").slides)[-3:]
        text = "\n".join(shape.text for slide in added[:2] for shape in slide.shapes)
        notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in added[:2])
        contents.append((text.replace("\v", "\n"), notes))
        assert all(not shape.text.strip() for shape in added[1].shapes)
        assert all(
            not shape.text.strip() or shape.text.startswith("※ネット表記")
            for shape in added[0].shapes
        )
        for slide in added[:2]:
            review = slide.notes_slide.notes_text_frame.text
            assert "要確認一覧（自動生成）" in review
            assert "依頼に記載が無いため未記入" in review
        assert "要確認" not in text
        for value in ("100,001", "250,003", "900000", "2.5円", "4週間", "4回接触", "60%", "10%"):
            assert value not in text
            assert value not in notes
    assert contents[0] == contents[1]


def test_worker_rejects_invalid_simulation_before_output(tmp_path: Path) -> None:
    _extension(tmp_path / "template.pptx")
    with pytest.raises(MediaOperationError, match="composer JSON is invalid"):
        _render(
            tmp_path,
            "worker",
            simulation={
                "views": 100,
                "contacts": 0,
                "understanding_rate": 0.5,
                "assumption": "仮定",
            },
        )
    assert not (tmp_path / "output.pptx").exists()


@pytest.mark.parametrize("reason", [None, "入力5件のため6番枠未使用"])
def test_unused_hearing_shape_is_removed_after_its_note_is_saved(reason: str | None) -> None:
    from teamagent.media.proposal_cleanup import prepare_proposal_content

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    kept = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1))
    kept.text = "既存の本文"
    before = kept._element.xml
    frame = slide.shapes.add_textbox(0, Inches(1), Inches(4), Inches(1)).text_frame
    frame.text = "{{PB-BU2-HEARING-6}}"
    notes = {"PB-BU2-HEARING-6": reason} if reason else {}
    prepare_proposal_content(prs, {}, {"PB-BU2-HEARING-6": ""}, notes)
    assert len(slide.shapes) == 1
    assert slide.shapes[0]._element.xml == before
    assert ("要確認" in slide.notes_slide.notes_text_frame.text) is (reason is None)
    assert (reason or "PB-BU2-HEARING-6") in slide.notes_slide.notes_text_frame.text
