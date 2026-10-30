"""SHA固定の実FMTで画像欄の検出・挿入を確認。資材はローカル限定・登録しない。"""

import hashlib
import json
import shutil
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from teamagent.media.operations import _iter_proposal_image_slots
from teamagent.media.proposal_quality import proposal_field_spec, proposal_readiness_issues
from teamagent.media.render_child import _proposal
from teamagent.skills.proposal_deck.contract import ComposerOutput
from teamagent.skills.proposal_deck.renderer import render_deck
from tests.skills.proposal_deck.test_renderer_images import _TINY_JPEG

_ARTIFACTS = Path.home() / "Documents/Claude/Artifacts"
_TEMPLATE = _ARTIFACTS / "handoff-20261030/fmt86_07ea4170.pptx"
_SHA256 = "07ea41700933d380295c5b8e713d803c8276c46c94b18008dbd8bf230539f883"
# ラベルで確認した画像欄。元FMTの表示順とshape_id（座標には依存しない）。
_IMAGE_FIELDS = {38: [2], 40: [7, 8, 10], 41: [12, 15, 17], 72: [5]}


@pytest.fixture
def real_template() -> Path:
    if not _TEMPLATE.is_file():
        pytest.skip("SHA固定の実FMTローカルコピーがない")
    assert hashlib.sha256(_TEMPLATE.read_bytes()).hexdigest() == _SHA256
    return _TEMPLATE


def test_real_template_image_spec_matches_worker_slots(real_template: Path) -> None:
    prs = Presentation(str(real_template))
    assert len(prs.slides) == 86
    expected = {
        (prs.slides[page - 1].slide_id, shape_id)
        for page, ids in _IMAGE_FIELDS.items()
        for shape_id in ids
    }
    spec = proposal_field_spec(prs)
    assert {(field["slide_id"], field["shape_id"]) for field in spec["images"]} == expected
    assert {
        (slide.slide_id, shape.shape_id) for slide, shape in _iter_proposal_image_slots(prs)
    } == expected
    assert all(field["expected"] == 1 for field in spec["images"])


@pytest.mark.parametrize(
    "scenario", ["none", "single", "valid", "broken", "overflow", "no-conditions"]
)
@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("unit,sector,expected_images", [("BU2", "食品", 1), ("BU1", "美容", 8)])
def test_real_template_filled_images_have_no_image_issue(
    real_template: Path,
    tmp_path: Path,
    backend: str,
    unit: str,
    sector: str,
    expected_images: int,
    scenario: str,
) -> None:
    composer_path = _ARTIFACTS / f"20261022_BU2FMT/composer_{unit}_{sector}.json"
    if not composer_path.is_file():
        pytest.skip("合成composerのローカルコピーがない")
    data = json.loads(composer_path.read_text())
    assert data["business_unit"] == unit and data["product_sector"] == sector
    if scenario == "no-conditions":
        data.update(plans=[], simulation=None, hearing_items=[], estimate_terms={})
    image_count = {
        "none": 0,
        "single": 1,
        "valid": 8,
        "broken": 1,
        "overflow": 9,
        "no-conditions": 0,
    }[scenario]
    # 合成画像だけを追加。原本FMT・composerは変更しない。
    image = tmp_path / "synthetic.jpg"
    image.write_bytes(b"broken" if scenario == "broken" else _TINY_JPEG)
    data["evidence_images"] = {
        58: [
            {"placeholder_id": 58, "rank": rank, "keyword": "合成画像", "image_path": str(image)}
            for rank in range(1, image_count + 1)
        ]
    }
    if image_count == 0:
        data["evidence_images"] = {}
    output = tmp_path / "output.pptx"
    metadata: dict[str, Any] = {}
    if backend == "local":
        render_deck(
            ComposerOutput.model_validate(data),
            real_template,
            output,
            enable_images=True,
            metadata=metadata,
        )
    else:
        shutil.copyfile(real_template, tmp_path / "template.pptx")
        (tmp_path / "composer.json").write_text(json.dumps(data))
        metadata = _proposal(
            tmp_path,
            {
                "kind": "proposal_pptx",
                "template": "template.pptx",
                "composer": "composer.json",
                "output": "output.pptx",
                "fail_if_missing": True,
                "evidence": [
                    {"placeholder_id": 58, "rank": rank, "path": image.name}
                    for rank in range(1, image_count + 1)
                ],
            },
        )
    issues = proposal_readiness_issues(metadata)
    if scenario == "broken":
        assert "image_insertion_failed:1" in issues
    else:
        assert issues == []
    assert "required_image_spec_missing" not in metadata["field_issues"]
    assert len(metadata["image_slots"]) == expected_images
    inserted = sum(field["inserted"] for field in metadata["image_slots"])
    assert metadata["evidence_images"] == inserted
    if scenario in {"single", "valid", "overflow"}:
        assert inserted == min(image_count, expected_images)
        if unit == "BU2":
            assert metadata["image_slots"][0]["slide"] == 54
            slide = Presentation(str(output)).slides[53]
            assert any(shape.shape_type == 13 for shape in slide.shapes)
    if image_count > expected_images:
        assert (
            f"要対応: 画像 {image_count - expected_images} 件は貼り先なし"
            in metadata["field_warnings"]
        )
    if scenario in {"none", "no-conditions"}:
        assert metadata["evidence_images"] == 0
        assert all(field["inserted"] == 0 for field in metadata["image_slots"])
        assert any(
            f"要対応: 画像未挿入 {expected_images} 枠（" in warning
            for warning in metadata["field_warnings"]
        )
    if scenario == "no-conditions" and unit == "BU2":
        hearing = next(slide for slide in Presentation(str(output)).slides if slide.slide_id == 341)
        assert 20 in {shape.shape_id for shape in hearing.shapes}
        assert not {21, 22, 23, 24, 25} & {shape.shape_id for shape in hearing.shapes}
        assert (
            sum(
                shape.text.count("条件未指定・要相談")
                for shape in hearing.shapes
                if shape.has_text_frame
            )
            == 1
        )
