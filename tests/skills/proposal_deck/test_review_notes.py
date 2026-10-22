"""不足枠は本文を空にし、ノートと完成通知へ回す。"""

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from teamagent.media.render_child import _proposal
from teamagent.skills._shared.deck_review import review_slides
from teamagent.skills.proposal_deck.renderer import render_pptx


@pytest.mark.parametrize("worker", [False, True])
def test_missing_content_moves_to_notes_and_keeps_notification(
    tmp_path: Path, worker: bool
) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "発表者のメモ"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "{1:商品}"
    slide.shapes.add_textbox(0, Inches(1), Inches(3), Inches(1)).text = "{2:不足}"
    slide.shapes.add_table(1, 1, 0, Inches(2), Inches(3), Inches(1)).table.cell(
        0, 0
    ).text = "{{PB-CASES}}"
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "{3:未検出}"
    template = tmp_path / "template.pptx"
    prs.save(template)
    output = tmp_path / "out.pptx"
    placeholders = {1: "ごちそうズボラ飯", 2: "要確認（文字数調整）"}
    auxiliary = {"PB-CASES": "要確認（出典付き実績候補未検出）"}
    if worker:
        (tmp_path / "composer.json").write_text(
            json.dumps({"placeholders": placeholders, "auxiliary_placeholders": auxiliary}),
            encoding="utf-8",
        )
        _proposal(
            tmp_path,
            {
                "kind": "proposal_pptx",
                "template": template.name,
                "composer": "composer.json",
                "evidence": [],
                "output": output.name,
                "fail_if_missing": True,
            },
        )
    else:
        render_pptx(
            template,
            {**placeholders, 3: "要確認（データ未検出）"},
            output,
            auxiliary_placeholders=auxiliary,
        )

    reopened = Presentation(output)
    assert len(reopened.slides) == 2
    first, second = reopened.slides
    assert first.shapes[0].shapes[0].text == "ごちそうズボラ飯"
    assert first.shapes[1].text == ""
    assert first.shapes[2].table.cell(0, 0).text == ""
    assert second.shapes[0].text == ""
    notes = first.notes_slide.notes_text_frame.text
    assert "発表者のメモ" in notes
    assert "要確認一覧" in notes and "文字数調整" in notes and "PB-CASES" in notes
    assert "出典付き実績候補未検出" in notes
    assert "データ未検出" in second.notes_slide.notes_text_frame.text
    assert review_slides(output.read_bytes(), template.read_bytes()) == [1, 2]


def test_presenter_notes_do_not_create_review_warning(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "発表前に要確認"
    template = tmp_path / "t.pptx"
    prs.save(template)
    assert review_slides(template.read_bytes(), template.read_bytes()) == []
