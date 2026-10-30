"""出荷前ゲートは local/worker で同じ本文・ノートを作る。"""

import json
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Inches

from teamagent.media.render_child import _proposal
from teamagent.skills._shared.deck_review import review_slides
from teamagent.skills.proposal_deck.contract import EvidenceImage
from teamagent.skills.proposal_deck.renderer import render_pptx
from tests.skills.proposal_deck.test_renderer_images import _TINY_JPEG


def _render(tmp_path: Path, backend: str, **context: Any) -> tuple[Any, Path]:
    template, output = tmp_path / "template.pptx", tmp_path / "output.pptx"
    values = {1: "〇〇構文", 2: "商品の実際の訴求"}
    if backend == "local":
        render_pptx(template, values, output, **context)
    else:
        (tmp_path / "composer.json").write_text(
            json.dumps({"placeholders": values, **context}, default=str), encoding="utf-8"
        )
        _proposal(
            tmp_path,
            {
                "kind": "proposal_pptx",
                "template": template.name,
                "composer": "composer.json",
                "evidence": [],
                "output": output.name,
                "fail_if_missing": True,
            },
        )
    return Presentation(str(output)), output


def _template(tmp_path: Path, pages: list[list[str]]) -> None:
    prs = Presentation()
    for texts in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for index, text in enumerate(texts):
            slide.shapes.add_textbox(0, Inches(index), Inches(8), Inches(1)).text = text
    prs.save(str(tmp_path / "template.pptx"))


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_instruction_frames_move_to_notes(tmp_path: Path, backend: str) -> None:
    artifacts = [
        "{1}",
        "#○○",
        "××",
        "{商材}",
        "・・・",
        "SNS投稿イメージなど貼り付け",
        "※該当記事の引用（画像/テキスト）",
        "ここに画像",
        "商品・サービス画像",
        "Xでの\nリアルなインサイトに関する\n投稿画像",
        "TBD",
    ]
    _template(tmp_path, [[*artifacts, "{2}", "※オリエン時にご提案", "施策A×施策B"]])
    rendered, output = _render(tmp_path, backend)
    texts = [shape.text for shape in rendered.slides[0].shapes]
    assert texts[: len(artifacts)] == [""] * len(artifacts)
    assert texts[-3:] == ["商品の実際の訴求", "※オリエン時にご提案", "施策A×施策B"]
    notes = rendered.slides[0].notes_slide.notes_text_frame.text
    assert "要確認一覧（自動生成）" in notes and "〇〇構文" in notes and "投稿画像" in notes
    assert review_slides(output.read_bytes(), (tmp_path / "template.pptx").read_bytes()) == [1]


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_unchanged_sample_warns_but_static_company_copy_does_not(
    tmp_path: Path, backend: str
) -> None:
    sample = "【商品紹介】ヤクルトが作っているスキンケア「ラクティフル」を紹介します。"
    _template(tmp_path, [["｜EMMEコンテンツイメージ", sample], ["株式会社ベクトルの実績紹介"]])
    rendered, _ = _render(tmp_path, backend, product_sector="美容")
    assert rendered.slides[0].shapes[1].text == sample
    assert "警告: FMT見本未上書き" in rendered.slides[0].notes_slide.notes_text_frame.text
    assert not rendered.slides[1].has_notes_slide


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("sector,count", [("食品", 2), ("化粧品", 6), ("ヘルスケア", 6), ("", 6)])
def test_emme_selection_preserves_other_pages(
    tmp_path: Path, backend: str, sector: str, count: int
) -> None:
    titles = [
        "｜美容メディア：EMME（アカウント群）",
        "｜美容メディア：EMME（編集部の特徴・実績）",
        "｜ライフスタイルメディア：EMME",
        "｜EMMEコンテンツイメージ",
    ]
    _template(tmp_path, [["提案"], *[[title] for title in titles], ["締め"]])
    rendered, _ = _render(tmp_path, backend, product_sector=sector)
    assert len(rendered.slides) == count
    assert rendered.slides[0].shapes[0].text == "提案"
    assert rendered.slides[-1].shapes[0].text == "締め"
    if not sector:
        assert "業種未指定" in rendered.slides[1].notes_slide.notes_text_frame.text


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_emme_mention_on_other_page_does_not_delete_it(tmp_path: Path, backend: str) -> None:
    _template(tmp_path, [["{2}"], ["今回の比較対象にEMMEを含む"], ["｜EMMEコンテンツイメージ"]])
    rendered, _ = _render(
        tmp_path,
        backend,
        product_sector="食品",
        auxiliary_placeholders={"PB-CASES": "EMMEとの比較"},
    )
    assert len(rendered.slides) == 2
    assert rendered.slides[1].shapes[0].text == "今回の比較対象にEMMEを含む"


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("agency", [False, True])
def test_estimate_margin_only_for_agency(tmp_path: Path, backend: str, agency: bool) -> None:
    original = "※税別・ネットになります。貴社マージンを乗せてご提出ください"
    _template(tmp_path, [[original]])
    rendered, _ = _render(tmp_path, backend, agency_route=agency)
    assert rendered.slides[0].shapes[0].text == (original if agency else "※税別")


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_long_body_warns_without_counting_notes(tmp_path: Path, backend: str) -> None:
    _template(tmp_path, [["{96}"], ["短" * 500]])
    source = Presentation(str(tmp_path / "template.pptx"))
    source.slides[1].notes_slide.notes_text_frame.text = "ノート" * 500
    source.save(str(tmp_path / "template.pptx"))
    # The helper's two model values are unrelated to this long-body case.
    from tests.skills.proposal_deck.test_template_cleanup import _render as render

    rendered: Any = render(backend, tmp_path, {96: "長" * 501})
    assert rendered.slides[0].shapes[0].text == "長" * 501
    assert "警告: 本文500字超過" in rendered.slides[0].notes_slide.notes_text_frame.text
    assert rendered.slides[1].notes_slide.notes_text_frame.text == "ノート" * 500


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_injected_image_is_not_reported_as_empty(tmp_path: Path, backend: str) -> None:
    from pptx.util import Emu

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Emu(914400), Emu(5454720), Emu(609905), Emu(609905)).text = "商品画像"
    template, output = tmp_path / "template.pptx", tmp_path / "output.pptx"
    prs.save(str(template))
    image = tmp_path / "evidence.jpg"
    image.write_bytes(_TINY_JPEG)
    if backend == "local":
        render_pptx(
            template,
            {},
            output,
            evidence_images={
                1: [EvidenceImage(placeholder_id=1, rank=1, keyword="商品", image_path=str(image))]
            },
        )
    else:
        (tmp_path / "composer.json").write_text('{"placeholders": {}}', encoding="utf-8")
        _proposal(
            tmp_path,
            {
                "kind": "proposal_pptx",
                "template": template.name,
                "composer": "composer.json",
                "output": output.name,
                "fail_if_missing": True,
                "evidence": [{"placeholder_id": 1, "rank": 1, "path": image.name}],
            },
        )
    rendered = Presentation(str(output))
    assert rendered.slides[0].shapes[0].text == ""
    assert sum(shape.shape_type == 13 for shape in rendered.slides[0].shapes) == 1
    assert review_slides(output.read_bytes(), template.read_bytes()) == []
