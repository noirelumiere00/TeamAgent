"""PRワード頁の見本写真と、FMTの書体で欠ける候補番号の回帰テスト。"""

from __future__ import annotations

import base64
import io
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches

from teamagent.media.render_child import _proposal
from teamagent.skills.proposal_deck.renderer import render_pptx

_JPEG = base64.b64decode(
    "/9j/4AAQSkZJRgABAQEAYABgAAD/2wBDAAMCAgICAgMCAgIDAwMDBAYEBAQEBAgGBgUGCQgKCgkICQkKDA8MCgsO"
    "CwkJDRENDg8QEBEQCgwSExIQEw8QEBD/wAALCAABAAEBAREA/8QAFAABAAAAAAAAAAAAAAAAAAAACP/EABQQAQAA"
    "AAAAAAAAAAAAAAAAAAD/2gAIAQEAAD8AfwD/2Q=="
)


def _render(
    backend: str,
    root: Path,
    placeholders: dict[int, str],
    auxiliary: dict[str, str] | None = None,
) -> Presentation:
    template = root / "template.pptx"
    output = root / "output.pptx"
    if backend == "local":
        render_pptx(template, placeholders, output, auxiliary_placeholders=auxiliary)
    else:
        (root / "composer.json").write_text(
            json.dumps(
                {"placeholders": placeholders, "auxiliary_placeholders": auxiliary or {}},
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        _proposal(
            root,
            {
                "kind": "proposal_pptx",
                "template": "template.pptx",
                "composer": "composer.json",
                "evidence": [],
                "output": "output.pptx",
                "fail_if_missing": True,
            },
        )
    return Presentation(str(output))


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_pr_sample_photo_removed_without_changing_logo_or_other_slides(
    tmp_path: Path, backend: str
) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    sample = slide.shapes.add_picture(
        io.BytesIO(_JPEG), 0, 0, presentation.slide_width, presentation.slide_height
    )
    sample.name = "Picture 2"
    sample._element.nvPicPr.cNvPr.set("descr", "SNSで話題“洗濯機クリーニング”の見本写真")
    overlay = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height
    )
    overlay.name = "見本写真の上の装飾"
    overlay_xml = overlay._element.xml
    logo = slide.shapes.add_picture(io.BytesIO(_JPEG), 0, 0, Inches(0.5), Inches(0.5))
    logo.name = "ブランドロゴ"
    logo_xml = logo._element.xml
    real_photo = slide.shapes.add_picture(
        io.BytesIO(_JPEG), 0, 0, presentation.slide_width, presentation.slide_height
    )
    real_photo.name = "差し替えた商品写真"
    real_photo._element.nvPicPr.cNvPr.set("descr", "公式から受領した商品画像")
    real_photo_xml = real_photo._element.xml
    for index, pid in enumerate((46, 47, 56)):
        slide.shapes.add_textbox(0, Inches(index + 1), Inches(5), Inches(0.5)).text = f"｛{pid}｝"
    other = presentation.slides.add_slide(presentation.slide_layouts[6])
    case_image = other.shapes.add_picture(
        io.BytesIO(_JPEG), 0, 0, presentation.slide_width, presentation.slide_height
    )
    case_image.name = "既存の施策事例画像"
    case_xml = case_image._element.xml
    presentation.save(str(tmp_path / "template.pptx"))

    rendered = _render(backend, tmp_path, {46: "ごちそうズボラ飯", 47: "候補文", 56: "背景"})

    assert [
        shape.name
        for shape in rendered.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ] == ["ブランドロゴ", "差し替えた商品写真"]
    shapes = {shape.name: shape for shape in rendered.slides[0].shapes}
    assert shapes["ブランドロゴ"]._element.xml == logo_xml
    assert shapes["差し替えた商品写真"]._element.xml == real_photo_xml
    assert shapes["見本写真の上の装飾"]._element.xml == overlay_xml
    assert rendered.slides[1].shapes[0]._element.xml == case_xml
    assert Presentation(str(tmp_path / "template.pptx")).slides[0].shapes[0].name == "Picture 2"


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_pr_candidate_numbers_and_missing_glyph_marker_use_safe_characters(
    tmp_path: Path, backend: str
) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(0, 0, Inches(8), Inches(1)).text_frame
    run = frame.paragraphs[0].add_run()
    run.text = "【PRワード 他候補】｛47｝"
    run.font.name = "Yu Gothic UI"
    slide.shapes.add_textbox(0, Inches(2), Inches(8), Inches(1)).text = "{{PB-KEY-MESSAGE}}"
    presentation.save(str(tmp_path / "template.pptx"))

    rendered = _render(
        backend,
        tmp_path,
        {47: "❶時短 ❷満足 ❾ご褒美 ☒ USP×本音・背景 ✅"},
        {"PB-KEY-MESSAGE": "❸毎日のごちそう"},
    )

    assert rendered.slides[0].shapes[0].text == (
        "【PRワード 他候補】\v1時短 2満足 9ご褒美 x USP×本音・背景 ✅"
    )
    assert rendered.slides[0].shapes[1].text == "3毎日のごちそう"
    assert rendered.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.name == "Yu Gothic UI"
