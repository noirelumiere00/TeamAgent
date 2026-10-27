"""BU2の頁選択は元FMTで確定し、残す頁の装飾・座標と入力検査を保つ。"""

import hashlib
import json
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Inches

from teamagent.media.operations import MediaOperationError
from teamagent.media.proposal_template import BU2_AUXILIARY_KEYS, KNOWN_TEMPLATE_SLIDES
from teamagent.media.render_child import _proposal
from teamagent.skills.proposal_deck.contract import (
    PROPOSAL_BUILDER_REQUIRED_AUXILIARY,
    VALID_IDS,
    ComposerOutput,
)
from teamagent.skills.proposal_deck.renderer import (
    UnfilledPlaceholderError,
    render_deck,
    render_pptx,
)
from tests.skills.proposal_deck.test_decoration_geometry import _geometry
from tests.skills.proposal_deck.test_decoration_geometry import _template as _decorated_template


def _legacy_template(path: Path) -> None:
    _decorated_template(path)
    prs = Presentation(path)
    titles = {
        12: "｜ショート動画を成果につなげる5つのポイント",
        13: "｜5つのポイントが連動して、売上エンジンになる",
        34: "02 |PRワードの開発",
        46: "まとめ\n｛43｝\n｛44｝\n｛45｝",
        47: "PRワード案\n｛46｝\n｛47｝\n｛56｝",
        48: "03 |コミュニケーション戦略",
        67: "｜美容メディア：EMME（アカウント群）",
        68: "｜美容メディア：EMME（編集部の特徴・実績）",
        69: "｜ライフスタイルメディア：EMME",
        70: "｜EMMEコンテンツイメージ",
    }
    for number in range(2, 84):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.3), Inches(0.5), Inches(8), Inches(1)).text = titles.get(
            number, f"食品メーカー向けの施策 {number}"
        )
    prs.slides[2].shapes.add_textbox(0, 0, Inches(8), Inches(5)).text = " ".join(
        f"{{{pid}}}" for pid in sorted(VALID_IDS)
    )
    prs.slides[3].shapes.add_textbox(0, 0, Inches(8), Inches(5)).text = " ".join(
        [f"{{{{{key}}}}}" for key in sorted(PROPOSAL_BUILDER_REQUIRED_AUXILIARY)]
        + [f"{{{{PB-DATE:{offset}:%m/%d}}}}" for offset in range(-56, 22, 7)]
        + ["{{PB-TEMPLATE:proposal-builder-v1}}"]
    )
    prs.save(path)


def _template(path: Path) -> None:
    _legacy_template(path)
    prs = Presentation(path)
    for prefix in ("PB-BU2-PLAN-", "PB-BU2-SIM-", "PB-BU2-HEARING"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        keys = {key for key in BU2_AUXILIARY_KEYS if key.startswith(prefix)}
        if prefix == "PB-BU2-PLAN-":
            keys.add("PB-BU2-ESTIMATE-NOTES")
        for index, key in enumerate(sorted(keys)):
            slide.shapes.add_textbox(
                0, Inches(index / 4), Inches(9), Inches(0.25)
            ).text = f"{{{{{key}}}}}"
    prs.save(path)


def _render(root: Path, backend: str, **context: Any) -> Path:
    values = {pid: "本文" for pid in VALID_IDS}
    # 生成文に除外対象の見出しが現れても、元FMTの頁選択に影響させない。
    values[1] = "｜ショート動画を成果につなげる5つのポイント"
    payload = {
        "placeholders": values,
        "template_profile": "proposal-builder-v1",
        "auxiliary_placeholders": {key: "設定値" for key in PROPOSAL_BUILDER_REQUIRED_AUXILIARY},
        "posting_start_date": date(2026, 12, 20),
        "proposal_date": date(2026, 10, 22),
        "product_sector": "美容",
        **context,
    }
    output = root / "output.pptx"
    if backend == "local":
        render_deck(ComposerOutput.model_construct(**payload), root / "template.pptx", output)
    else:
        (root / "composer.json").write_text(json.dumps(payload, default=str), encoding="utf-8")
        _proposal(
            root,
            {
                "kind": "proposal_pptx",
                "template": "template.pptx",
                "composer": "composer.json",
                "evidence": [],
                "output": "output.pptx",
                "fail_if_missing": True,
            },
        )
    return output


def _pin_template(monkeypatch: pytest.MonkeyPatch, path: Path, slide_count: int = 86) -> None:
    monkeypatch.setitem(
        KNOWN_TEMPLATE_SLIDES, hashlib.sha256(path.read_bytes()).hexdigest(), slide_count
    )


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("unit", [None, "BU2", "BU1"])
@pytest.mark.parametrize("sector", ["食品", "美容"])
def test_bu2_selection_preserves_remaining_geometry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, backend: str, unit: str | None, sector: str
) -> None:
    template = tmp_path / "template.pptx"
    _template(template)
    _pin_template(monkeypatch, template)
    before = template.read_bytes()
    source = Presentation(template)
    excluded = ({12, 13} | set(range(34, 46))) if unit != "BU1" else {84, 85, 86}
    if sector == "食品":
        excluded.update(range(67, 71))
    kept = [number for number in range(1, 87) if number not in excluded]
    context = {"business_unit": unit} if unit else {}
    context["hearing_items"] = [f"確認事項{number}" for number in range(1, 7)]
    output = _render(tmp_path, backend, product_sector=sector, **context)
    rendered = Presentation(output)
    assert [slide.slide_id for slide in rendered.slides] == [
        source.slides[number - 1].slide_id for number in kept
    ]
    original = _geometry(template)
    assert _geometry(output) == {
        (new_number, *key[1:]): value
        for new_number, old_number in enumerate(kept, 1)
        for key, value in original.items()
        if key[0] == old_number
    }
    assert template.read_bytes() == before
    assert rendered.slides[0].shapes[0].text == "｜ショート動画を成果につなげる5つのポイント"


@pytest.mark.parametrize("value", [0, False, [], {}, "", "未定", "2026-99-99"])
def test_worker_rejects_invalid_posting_date(tmp_path: Path, value: Any) -> None:
    _template(tmp_path / "template.pptx")
    with pytest.raises(MediaOperationError, match="composer JSON is invalid"):
        _render(tmp_path, "worker", posting_start_date=value)
    assert not (tmp_path / "output.pptx").exists()


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("missing_page", [34, 46, 47, 48])
def test_incomplete_pr_chapter_is_retained(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, backend: str, missing_page: int
) -> None:
    template = tmp_path / "template.pptx"
    _template(template)
    prs = Presentation(template)
    prs.slides[missing_page - 1].shapes[0].text = "別構成の資料"
    prs.save(template)
    _pin_template(monkeypatch, template)
    output = _render(tmp_path, backend)
    assert len(Presentation(output).slides) == 84


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_input_inventory_is_checked_before_bu2_removal(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, backend: str
) -> None:
    template = tmp_path / "template.pptx"
    _template(template)
    prs = Presentation(template)
    item = prs.slides._sldIdLst[34]
    prs.part.drop_rel(item.rId)
    prs.slides._sldIdLst.remove(item)
    prs.save(template)
    _pin_template(monkeypatch, template)
    with pytest.raises((UnfilledPlaceholderError, MediaOperationError)):
        _render(tmp_path, backend)
    assert not (tmp_path / "output.pptx").exists()


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("unit", ["BU3", None, [], {}])
def test_invalid_business_unit_is_rejected(tmp_path: Path, backend: str, unit: Any) -> None:
    _template(tmp_path / "template.pptx")
    with pytest.raises((ValueError, MediaOperationError)):
        if backend == "local":
            render_pptx(
                tmp_path / "template.pptx", {}, tmp_path / "output.pptx", business_unit=unit
            )
        else:
            _render(tmp_path, backend, business_unit=unit)
    assert not (tmp_path / "output.pptx").exists()
