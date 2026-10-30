"""最終バイトのL1・補助欄・画像を完成判定に使う。外部接続なし。"""

import hashlib
import json
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from teamagent.media.render_child import _proposal
from teamagent.skills.proposal_deck.contract import ComposerOutput
from teamagent.skills.proposal_deck.renderer import render_deck
from tests.skills.proposal_deck.test_bu2_shape import _pin_template, _template
from tests.skills.proposal_deck.test_proposal_deck_skill import _full_composer_json
from tests.skills.proposal_deck.test_renderer_images import _TINY_JPEG


def _render(root: Path, backend: str, **updates: Any) -> tuple[Path, dict[str, Any]]:
    from teamagent.media.proposal_template import BASE_AUXILIARY_KEYS

    data = json.loads(_full_composer_json())
    data.update(
        template_profile="proposal-builder-v1",
        auxiliary_placeholders={key: "設定値" for key in BASE_AUXILIARY_KEYS},
        product_sector="美容",
    )
    data.update(updates)
    evidence = data.pop("worker_evidence", [])
    output = root / "output.pptx"
    metadata: dict[str, Any] = {}
    if backend == "local":
        render_deck(
            ComposerOutput.model_validate(data),
            root / "template.pptx",
            output,
            enable_images=True,
            metadata=metadata,
        )
    else:
        (root / "composer.json").write_text(json.dumps(data), encoding="utf-8")
        metadata = _proposal(
            root,
            {
                "kind": "proposal_pptx",
                "template": "template.pptx",
                "composer": "composer.json",
                "output": "output.pptx",
                "evidence": evidence,
                "fail_if_missing": True,
            },
        )
    return output, metadata


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize(
    "unit,sector,expected",
    [
        ("BU2", "美容", 72),
        ("BU2", "食品", 68),
        ("BU1", "美容", 83),
        ("BU1", "食品", 79),
    ],
)
@pytest.mark.parametrize("damage", [False, True])
def test_l1_checks_spec_count_and_final_bytes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    backend: str,
    unit: str,
    sector: str,
    expected: int,
    damage: bool,
) -> None:
    template = tmp_path / "template.pptx"
    _template(template)
    if damage:
        prs = Presentation(str(template))
        prs.slides[33].shapes[0].text = "異なる章見出し"
        prs.save(str(template))
    _pin_template(monkeypatch, template)
    output, metadata = _render(tmp_path, backend, business_unit=unit, product_sector=sector)
    assert metadata["expected_slides"] == expected
    assert metadata["output_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert metadata["deck_qa"]["error_count"] == int(damage and unit == "BU2")


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize("image_state", ["missing", "broken", "valid"])
def test_required_image_slot_counts_actual_insertions(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    backend: str,
    image_state: str,
) -> None:
    template = tmp_path / "template.pptx"
    _template(template)
    prs = Presentation(str(template))
    slot = prs.slides[5].shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 5454720, 609905, 609905)
    slot.text = "ここに画像"
    shape_id = slot.shape_id
    prs.save(str(template))
    _pin_template(monkeypatch, template)
    updates: dict[str, Any] = {}
    if image_state != "missing":
        image = tmp_path / "image.jpg"
        image.write_bytes(_TINY_JPEG if image_state == "valid" else b"broken")
        updates["evidence_images"] = {
            58: [
                {
                    "placeholder_id": 58,
                    "rank": 1,
                    "keyword": "合成",
                    "image_path": str(image),
                }
            ]
        }
        updates["worker_evidence"] = [{"placeholder_id": 58, "rank": 1, "path": image.name}]
    _, metadata = _render(tmp_path, backend, **updates)
    slot_result = next(item for item in metadata["image_slots"] if item["shape_id"] == shape_id)
    assert slot_result["slide"] == 6
    assert slot_result["expected"] == 1
    assert slot_result["inserted"] == int(image_state == "valid")
    assert any("image_insertion_failed" in issue for issue in metadata["field_issues"]) is (
        image_state == "broken"
    )
    if image_state == "missing":
        assert any("画像未挿入 1 枠（6枚目）" in w for w in metadata["field_warnings"])


@pytest.mark.parametrize("backend", ["local", "worker"])
@pytest.mark.parametrize(
    "unit,sector,removed_page", [("BU2", "美容", 38), ("BU2", "食品", 67), ("BU1", "美容", 84)]
)
def test_image_skips_bu_and_sector_excluded_pages(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    backend: str,
    unit: str,
    sector: str,
    removed_page: int,
) -> None:
    template = tmp_path / "template.pptx"
    _template(template)
    prs = Presentation(str(template))
    for slide in (prs.slides[removed_page - 1], prs.slides[81]):
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = "ここに画像"
    prs.save(str(template))
    _pin_template(monkeypatch, template)
    image = tmp_path / "synthetic.jpg"
    image.write_bytes(_TINY_JPEG)
    _, metadata = _render(
        tmp_path,
        backend,
        business_unit=unit,
        product_sector=sector,
        evidence_images={
            58: [{"placeholder_id": 58, "rank": 1, "keyword": "合成", "image_path": str(image)}]
        },
        worker_evidence=[{"placeholder_id": 58, "rank": 1, "path": image.name}],
    )
    assert metadata["evidence_images"] == 1
    assert len(metadata["image_slots"]) == 1
    assert metadata["image_slots"][0]["inserted"] == 1
    assert not any(issue.startswith("image_") for issue in metadata["field_issues"])


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_missing_auxiliary_is_visible_and_unused_hearing_is_exempt(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    backend: str,
) -> None:
    template = tmp_path / "template.pptx"
    _template(template)
    _pin_template(monkeypatch, template)
    output, metadata = _render(tmp_path, backend, hearing_items=[f"確認事項{i}" for i in range(4)])
    prs = Presentation(str(output))
    assert "条件未指定・要相談" in "\n".join(s.text for s in prs.slides[-3].shapes)
    assert len(prs.slides[-1].shapes) == 4
    assert not any("PB-BU2-" in issue for issue in metadata["field_issues"])
    assert any("条件未指定・要相談" in w for w in metadata["field_warnings"])
    assert not any("PB-BU2-HEARING" in issue for issue in metadata["field_issues"])


@pytest.mark.parametrize("qa_fails", [False, True])
def test_l1_exception_is_explicit_without_losing_final_artifact(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    qa_fails: bool,
) -> None:
    from teamagent.media.proposal_quality import inspect_proposal, proposal_readiness_issues

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    output = tmp_path / "output.pptx"
    prs.save(str(output))
    original = output.read_bytes()
    if qa_fails:

        def fail(*args: Any, **kwargs: Any) -> Any:
            raise RuntimeError("private detail")

        monkeypatch.setattr("teamagent.media.deck_qa.inspect_pptx", fail)
    metadata = inspect_proposal(output, expected_slides=1)
    metadata["field_issues"] = []
    assert ("pptx_l1_unavailable" in proposal_readiness_issues(metadata)) is qa_fails
    assert output.read_bytes() == original
    assert "private detail" not in str(metadata)


@pytest.mark.parametrize(
    "identity", ["matches", "wrong-sha", "wrong-count", "image-dropped", "image-surplus"]
)
def test_deck_checks_worker_inspection_identity(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    identity: str,
) -> None:
    from unittest.mock import MagicMock

    from teamagent.skills.base import SkillContext
    from teamagent.skills.proposal_deck.schema import ProposalDeckInput
    from teamagent.skills.proposal_deck.skill import ProposalDeckSkill

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    template = tmp_path / "template.pptx"
    prs.save(str(template))
    body = template.read_bytes()
    skill = ProposalDeckSkill(bedrock=MagicMock())
    composer = ComposerOutput.model_validate_json(_full_composer_json())
    if identity in {"image-dropped", "image-surplus"}:
        from teamagent.skills.proposal_deck.contract import EvidenceImage

        composer.evidence_images = {
            58: [
                EvidenceImage(
                    placeholder_id=58,
                    rank=rank,
                    keyword="合成",
                    image_path="/missing/image.jpg",
                )
                for rank in range(1, 9 if identity == "image-surplus" else 2)
            ]
        }
    monkeypatch.setattr(skill, "_compose", lambda *args: (composer, 0.0, []))

    def render(_composer: Any, _template: Path, output: Path, **kwargs: Any) -> Path:
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(body)
        kwargs["metadata"].update(
            {
                "output_sha256": hashlib.sha256(body).hexdigest()
                if identity != "wrong-sha"
                else "bad",
                "expected_slides": 68 if identity != "wrong-count" else 86,
                "deck_qa": {"error_count": 0},
                "field_issues": [],
                "requested_evidence_images": int(identity == "image-surplus"),
                "evidence_images": int(identity == "image-surplus"),
                "image_slots": [
                    {"slide": 54, "expected": 1, "inserted": int(identity == "image-surplus")}
                ],
                "field_warnings": ["要対応: 画像未挿入 1 枠（54枚目）"],
            }
        )
        return output

    monkeypatch.setattr(skill, "_render_pptx", render)
    monkeypatch.setattr(skill, "_publish_if_enabled", lambda *args, **kwargs: None)
    request = ProposalDeckInput(
        product_name="合成",
        goal="認知",
        target_persona="生活者",
        template_path=str(template),
        out_dir=str(tmp_path / "out"),
        product_sector="食品",
    ).model_copy(update={"template_profile": "proposal-builder-v1"})
    output = skill.run(request, SkillContext())
    assert ("qa_error" in output.render_metadata) is (identity in {"wrong-sha", "wrong-count"})
    if identity != "image-surplus":
        assert "要対応: 画像未挿入 1 枠（54枚目）" in output.quality_warnings
    if identity == "image-dropped":
        assert "image_insertion_failed:1" in output.render_metadata["field_issues"]
    if identity == "image-surplus":
        assert output.render_metadata["field_issues"] == []
        assert "要対応: 画像 7 件は貼り先なし" in output.quality_warnings
    assert output.render_metadata["deck_qa"] == {"error_count": 0}
    assert Path(output.pptx_path).read_bytes() == body


def test_auxiliary_spec_recognizes_token_split_across_paragraphs() -> None:
    from teamagent.media.proposal_quality import proposal_field_results, proposal_field_spec

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_textbox(0, 0, 2000000, 500000).text_frame
    frame.text = "{{PB-"
    frame.add_paragraph().text = "CASES}}"
    spec = proposal_field_spec(prs)
    result = proposal_field_results(prs, spec, {"PB-CASES": ""}, {})
    assert "auxiliary_missing:PB-CASES" in result["field_issues"]


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_no_declared_image_fields_is_valid_zero_requirement(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    backend: str,
) -> None:
    template = tmp_path / "template.pptx"
    _template(template)
    _pin_template(monkeypatch, template)
    _, metadata = _render(tmp_path, backend, template_profile="base")
    assert metadata["image_slots"] == []
    assert not any("image" in issue for issue in metadata["field_issues"])


@pytest.mark.parametrize("backend", ["local", "worker"])
def test_non_condition_auxiliary_review_is_hidden_but_still_an_issue(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    backend: str,
) -> None:
    from teamagent.media.proposal_template import BASE_AUXILIARY_KEYS

    template = tmp_path / "template.pptx"
    _template(template)
    _pin_template(monkeypatch, template)
    values = {key: "設定値" for key in BASE_AUXILIARY_KEYS}
    values["PB-ACCOUNTS"] = "合成アカウント：説明途中に要確認（出典URL未取得）\n別のアカウント"
    output, metadata = _render(tmp_path, backend, auxiliary_placeholders=values)
    prs = Presentation(str(output))
    text = "\n".join(shape.text for shape in prs.slides[3].shapes)
    assert "条件未指定・要相談" not in text
    assert "合成アカウント" not in text
    assert "auxiliary_missing:PB-ACCOUNTS" in metadata["field_issues"]
    assert "出典URL未取得" in prs.slides[3].notes_slide.notes_text_frame.text


def test_coordinate_only_rectangle_is_not_an_image_slot() -> None:
    from teamagent.media.proposal_quality import iter_proposal_image_slots

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 5454720, 609905, 609905)
    labelled = slide.shapes.add_textbox(0, 0, 2000000, 1000000)
    labelled.text = "SNS投稿イメージ\nなど貼り付け"
    assert [shape.shape_id for _, shape in iter_proposal_image_slots(prs)] == [labelled.shape_id]


def test_unexpected_integrated_template_reports_unreadable_image_spec() -> None:
    from teamagent.media.proposal_quality import proposal_field_results, proposal_field_spec

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    spec = proposal_field_spec(prs, template_profile="proposal-builder-v1")
    metadata = proposal_field_results(prs, spec, {}, {})
    assert "required_image_spec_missing" in metadata["field_issues"]
