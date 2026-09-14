"""clip_proposal テストの共通フィクスチャ。

**テンプレ実物は repo に置かない**（他社名・第三者の顔写真・社員実名を含む）。代わりに
``golden_template_shapes.json``（shape_id / 段落数 / 枠 EMU / vert のダンプ）から
**同じ構造の合成テンプレ**を組み立てる。

⚠ 合成テンプレを ``CLIP_TEMPLATE_INVENTORY_V1`` から組むと、台帳が実在資産と食い違って
いても受入テストは永久に緑になる（自己言及）。情報源を 2 本に分け、**ダンプ側を
「実物の写し」・台帳側を「実装が期待する契約」**として突き合わせる。台帳だけを
書き換えると ``test_validate_template_accepts_the_golden_dump`` が赤くなる。

合成テンプレは段落ごとに異なる ``sz``（文字サイズ）を持たせてある。差し替え後も
段落ごとの書式が残ることを検査するため（``_replace_placeholders`` 相当の実装に
差し替えると全段落が段落 0 へ潰れて必ず赤くなる）。

``build_unsanitized_template`` は **消毒漏れのテンプレ**（docProps に実在社員名・
shape を消しても残る孤児 media part）を再現する。V6（``sanitize_output``）の回帰用。
"""

from __future__ import annotations

import json
import struct
import zlib
from pathlib import Path
from typing import Any

from teamagent.skills.clip_proposal.analysis import (
    AppealAxis,
    ClipPlan,
    ClipProposalAnalysis,
    ClipWindow,
    Community,
    CommunityTerm,
    Insight,
)
from teamagent.skills.clip_proposal.inventory import ImageSlotSpec, TextFrameSpec

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

#: 段落番号 → ``sz``。差し替え後も段落ごとに保たれることを検査する。
PARAGRAPH_SIZES = (500, 700, 750, 900, 1000, 1100, 1200)

GOLDEN_DUMP_PATH = Path(__file__).with_name("golden_template_shapes.json")


def load_golden_dump() -> dict[str, Any]:
    """消毒済みテンプレの受入ダンプ（台帳とは独立した情報源）。"""

    return json.loads(GOLDEN_DUMP_PATH.read_text(encoding="utf-8"))


def golden_text_frames() -> tuple[TextFrameSpec, ...]:
    """ダンプ側のテキスト枠。``max_chars`` は差し替え対象外なので 0 のまま。"""

    return tuple(
        TextFrameSpec(
            shape_id=int(entry["shape_id"]),
            role=str(entry["role"]),
            template_paragraphs=int(entry["paragraphs"]),
            width_emu=int(entry["width_emu"]),
            height_emu=int(entry["height_emu"]),
            vert=str(entry.get("vert") or ""),
        )
        for entry in load_golden_dump()["text_frames"]
    )


def golden_image_slots() -> tuple[ImageSlotSpec, ...]:
    return tuple(
        ImageSlotSpec(
            shape_id=int(entry["shape_id"]),
            role=str(entry["role"]),
            kind=str(entry["kind"]),  # type: ignore[arg-type]
            width_emu=int(entry["width_emu"]),
            height_emu=int(entry["height_emu"]),
        )
        for entry in load_golden_dump()["image_slots"]
    )


def _qn(tag: str) -> str:
    return f"{{{_A}}}{tag}"


def make_png_bytes(width: int = 4, height: int = 4) -> bytes:
    """依存を足さずに作る最小 PNG（画像スロット用）。"""

    raw = b"".join(b"\x00" + b"\x80\x80\x80" * width for _ in range(height))

    def chunk(kind: bytes, payload: bytes) -> bytes:
        body = kind + payload
        return struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def _force_shape_id(shape: Any, shape_id: int) -> None:
    element = shape._element
    for name in ("nvSpPr", "nvPicPr", "nvGrpSpPr"):
        holder = getattr(element, name, None)
        if holder is not None:
            holder.cNvPr.set("id", str(shape_id))
            holder.cNvPr.set("name", f"shape-{shape_id}")
            return
    raise AssertionError(f"cannot set shape id on {shape}")


def _add_text_frame(slide: Any, spec: TextFrameSpec) -> Any:
    from lxml import etree
    from pptx.util import Emu

    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(spec.width_emu), Emu(spec.height_emu))
    _force_shape_id(box, spec.shape_id)
    frame = box.text_frame
    body = frame._txBody

    if spec.vert:
        body.find(_qn("bodyPr")).set("vert", spec.vert)

    # 既定で 1 段落あるので、不足分を足す。段落ごとに sz を変えて書式保持を検査可能にする。
    while len(body.findall(_qn("p"))) < spec.template_paragraphs:
        frame.add_paragraph()
    for index, paragraph in enumerate(body.findall(_qn("p"))):
        run = etree.SubElement(paragraph, _qn("r"))
        props = etree.SubElement(run, _qn("rPr"))
        props.set("lang", "ja-JP")
        props.set("sz", str(PARAGRAPH_SIZES[index % len(PARAGRAPH_SIZES)]))
        text = etree.SubElement(run, _qn("t"))
        text.text = f"{spec.role}#{index}"
    return box


def _add_image_slot(slide: Any, spec: ImageSlotSpec) -> Any:
    import io

    from pptx.util import Emu

    if spec.kind == "picture":
        shape = slide.shapes.add_picture(
            io.BytesIO(make_png_bytes()),
            Emu(0),
            Emu(0),
            Emu(spec.width_emu),
            Emu(spec.height_emu),
        )
    else:
        shape = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(spec.width_emu), Emu(spec.height_emu))
    _force_shape_id(shape, spec.shape_id)
    return shape


def build_synthetic_template(
    path: str,
    *,
    omit_shape_id: int | None = None,
    paragraph_delta: dict[int, int] | None = None,
    author: str = "",
    orphan_media: bytes | None = None,
) -> str:
    """受入ダンプと同じ構造の合成テンプレを作る（**台帳は参照しない**）。

    ``omit_shape_id`` で 1 枠だけ落とす / ``paragraph_delta`` で段落数をずらすことで、
    テンプレ改竄検知（fail-closed）のテストが書ける。``author`` / ``orphan_media`` は
    消毒漏れ（docProps の実在社員名・shape を消しても残る media part）の再現用。
    """

    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    deltas = paragraph_delta or {}
    for spec in golden_text_frames():
        if spec.shape_id == omit_shape_id:
            continue
        delta = deltas.get(spec.shape_id, 0)
        if delta:
            spec = TextFrameSpec(
                shape_id=spec.shape_id,
                role=spec.role,
                template_paragraphs=max(1, spec.template_paragraphs + delta),
                width_emu=spec.width_emu,
                height_emu=spec.height_emu,
                vert=spec.vert,
                max_chars=spec.max_chars,
            )
        _add_text_frame(slide, spec)
    for slot in golden_image_slots():
        if slot.shape_id == omit_shape_id:
            continue
        _add_image_slot(slide, slot)

    if author:
        presentation.core_properties.author = author
        presentation.core_properties.last_modified_by = author
        presentation.core_properties.comments = f"{author} が編集"
    if orphan_media is not None:
        # 画像を足してから shape だけ消す。part（ppt/media/imageN.png）は残る
        # ＝ 消毒スクリプトが shape を消しただけで終わったときの実際の形。
        import io

        from pptx.util import Emu

        picture = slide.shapes.add_picture(
            io.BytesIO(orphan_media), Emu(0), Emu(0), Emu(914400), Emu(914400)
        )
        picture._element.getparent().remove(picture._element)

    presentation.save(path)
    return path


def build_unsanitized_template(
    path: str, *, author: str = "高林 拓也", orphan_long_edge: int = 8
) -> str:
    """消毒漏れテンプレ（docProps に実在社員名 ＋ 孤児 media part）。"""

    return build_synthetic_template(
        path, author=author, orphan_media=make_png_bytes(orphan_long_edge, orphan_long_edge)
    )


def sample_community(index: int) -> Community:
    return Community(
        name=f"界隈{index}",
        terms=(
            CommunityTerm(term=f"観測語{index}", verified=True),
            CommunityTerm(term=f"未確認語{index}", verified=False),
        ),
        scale_note=f"約{index}万人",
        description=f"界隈{index}の説明",
        is_primary=index == 1,
    )


def sample_insight(index: int) -> Insight:
    return Insight(
        target=f"ターゲット{index}",
        insight=f"本音{index}",
        evidence_hint=f"判断材料{index}",
    )


def sample_clip(cell_index: int) -> ClipPlan:
    kind = "community" if cell_index < 5 else "insight"
    if kind == "community":
        source = sample_community(cell_index + 1)
        label, detail = source.render_name(), source.render_detail_lines()
    else:
        insight = sample_insight(cell_index - 4)
        label, detail = insight.target, insight.render_detail_lines()
    return ClipPlan(
        cell_index=cell_index,
        kind=kind,  # type: ignore[arg-type]
        label=label,
        detail_lines=tuple(detail),
        window=ClipWindow(
            start_sec=float(cell_index * 5),
            end_sec=float(cell_index * 5 + 18),
            hook_start_sec=float(cell_index * 5),
            hook_end_sec=float(cell_index * 5 + 2),
        ),
        hook_copy=f"フック{cell_index}",
        band_top=f"上帯{cell_index}",
        band_bottom=f"下帯{cell_index}",
        search_word="🔍テスト商材 採用",
        quote_evidence=f"引用{cell_index}",
        takeaway=f"伝わること{cell_index}",
    )


def sample_analysis(
    *, clip_count: int = 10, client_name: str = "テスト商事"
) -> ClipProposalAnalysis:
    return ClipProposalAnalysis(
        client_name=client_name,
        axes=tuple(
            AppealAxis(index=i + 1, text=f"訴求軸{i + 1}", quote_evidence=f"引用{i + 1}")
            for i in range(5)
        ),
        clips=tuple(sample_clip(i) for i in range(clip_count)),
        cost_usd=0.12,
        gemini_calls=2,
    )


__all__ = [
    "GOLDEN_DUMP_PATH",
    "PARAGRAPH_SIZES",
    "build_synthetic_template",
    "build_unsanitized_template",
    "golden_image_slots",
    "golden_text_frames",
    "load_golden_dump",
    "make_png_bytes",
    "sample_analysis",
    "sample_clip",
    "sample_community",
    "sample_insight",
]
