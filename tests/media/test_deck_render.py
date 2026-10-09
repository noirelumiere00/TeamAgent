"""DeckSpec → 編集できる PowerPoint（media worker 側の python-pptx 描画）。

各スライドの部品の種類（入力欄・表・グラフ・図）・ノート・書体・字の大きさを、保存した .pptx を
python-pptx で開き直して確かめる（描いた側のオブジェクトではなく、ファイルの中身を見る）。
"""

from __future__ import annotations

import io
from typing import Any

import pytest

pptx = pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER  # noqa: E402

from teamagent.media.deck_contracts import DeckSpec  # noqa: E402
from teamagent.media.deck_render import (  # noqa: E402
    DeckRenderError,
    render_deck_csv,
    render_deck_files,
    render_deck_pptx,
)
from teamagent.skills.search_surface_check.deck import build_surface_deck  # noqa: E402
from tests.skills.search_surface_check.deck_fixtures import (  # noqa: E402
    CLIENT,
    NOW,
    REPORT_ID,
    ai_conclusion,
    covers,
    surface,
)

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@pytest.fixture(scope="module")
def built() -> tuple[DeckSpec, Any]:
    s = surface(15, conclusion=ai_conclusion())
    deck = build_surface_deck(
        [s],
        client_name=CLIENT,
        client_accounts=["kurashiru.com"],
        measured_epoch=NOW,
        report_id=REPORT_ID,
        include_appendix=True,
        addressee="テスト株式会社 御中",
        covers=covers(s.posts),
    )
    body = render_deck_pptx(deck.spec, deck.images)
    return deck.spec, Presentation(io.BytesIO(body))


def _slide(spec: DeckSpec, prs: Any, slide_id: str) -> Any:
    index = [s.slide_id for s in spec.slides].index(slide_id)
    return prs.slides[index]


def _kinds(slide: Any) -> list[str]:
    """部品の種類: 入力欄 / 表 / グラフ / 図 / 図の枠（画像なし）/ それ以外。"""
    out = []
    for shape in slide.shapes:
        name = shape.__class__.__name__
        if getattr(shape, "has_chart", False):
            out.append("chart")
        elif getattr(shape, "has_table", False):
            out.append("table")
        elif name == "PlaceholderPicture" or shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            out.append("picture")
        elif name == "PicturePlaceholder":
            out.append("picture_frame")
        elif shape.is_placeholder:
            out.append("placeholder")
        else:
            out.append("other")
    return out


def _layout_size_pt(slide: Any, shape: Any) -> float:
    """入力欄の字の大きさ（run に無ければレイアウトの lstStyle から）。"""
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                return float(run.font.size.pt)
    idx = shape.placeholder_format.idx
    layout_ph = next(p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == idx)
    rpr = layout_ph._element.find(f".//{_A}lstStyle/{_A}lvl1pPr/{_A}defRPr")
    return int(rpr.get("sz")) / 100


def test_layouts_are_named_and_slides_pick_them_by_name(built: tuple[DeckSpec, Any]) -> None:
    spec, prs = built
    names = [layout.name for layout in prs.slide_layouts]
    assert names == ["R_表紙", "R_結論", "R_数字", "R_グラフ", "R_表", "R_動画カード", "R_付録の表"]
    for slide_spec, slide in zip(spec.slides, prs.slides, strict=True):
        assert slide.slide_layout.name == slide_spec.layout
    assert prs.slide_width == 12_192_000 and prs.slide_height == 6_858_000


def test_part_kinds_per_slide(built: tuple[DeckSpec, Any]) -> None:
    spec, prs = built
    cover = _kinds(_slide(spec, prs, "SS-01"))
    assert cover.count("picture") == 1 and "placeholder" in cover
    assert "table" in _kinds(_slide(spec, prs, "SS-05"))
    assert "chart" in _kinds(_slide(spec, prs, "SS-07"))
    cards = _kinds(_slide(spec, prs, "SS-08"))
    assert cards.count("picture") == 5
    assert "table" in _kinds(_slide(spec, prs, "SS-09"))
    for slide in prs.slides:
        # 置いた部品はすべて入力欄・表・グラフ・図（手で置いた図形やテキストボックスは無い）
        assert "other" not in _kinds(slide)


def test_every_slide_has_title_condition_footer_page_number_and_notes(
    built: tuple[DeckSpec, Any],
) -> None:
    _, prs = built
    for number, slide in enumerate(prs.slides, start=1):
        types = [s.placeholder_format.type for s in slide.placeholders]
        assert PP_PLACEHOLDER.TITLE in types
        assert PP_PLACEHOLDER.FOOTER in types and PP_PLACEHOLDER.SLIDE_NUMBER in types
        texts = [s.text_frame.text for s in slide.placeholders if s.has_text_frame]
        assert any(t.startswith("「スパイスカレー 作り方」で検索した上位 15 本") for t in texts)
        number_ph = next(
            s
            for s in slide.placeholders
            if s.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
        )
        assert number_ph._element.find(f".//{_A}fld").get("type") == "slidenum"
        assert number_ph.text_frame.text == str(number)
        notes = slide.notes_slide.notes_text_frame.text
        assert notes.startswith("【このページは何を見て何を出しているか】")


def test_no_empty_prompt_placeholders_are_left(built: tuple[DeckSpec, Any]) -> None:
    _, prs = built
    for slide in prs.slides:
        for shape in slide.placeholders:
            if shape.has_text_frame:
                assert shape.text_frame.text.strip(), shape.name


def test_theme_fonts_colors_and_japanese_language(built: tuple[DeckSpec, Any]) -> None:
    _, prs = built
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    theme = prs.slide_master.part.part_related_by(RT.THEME).blob.decode("utf-8")
    assert theme.count('typeface="游ゴシック"') == 4  # 見出し・本文 × latin・ea
    assert '<a:accent1><a:srgbClr val="2F4B7C"/>' in theme
    assert '<a:dk2><a:srgbClr val="2A3351"/>' in theme
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    assert run.font._rPr.get("lang") == "ja-JP"
                    assert run.font.name is None  # 書体はテーマを参照（run に書かない）


def test_font_sizes_follow_the_template(built: tuple[DeckSpec, Any]) -> None:
    spec, prs = built
    for slide in prs.slides:
        title = slide.shapes.title
        assert _layout_size_pt(slide, title) == 28
    conclusion = _slide(spec, prs, "SS-04")
    body = next(s for s in conclusion.placeholders if s.placeholder_format.idx == 1)
    assert _layout_size_pt(conclusion, body) == 16

    def table_pt(slide: Any) -> float:
        table = next(s for s in slide.shapes if getattr(s, "has_table", False)).table
        return float(table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size.pt)

    assert table_pt(_slide(spec, prs, "SS-05")) == 14  # 6 行以下
    assert table_pt(_slide(spec, prs, "SS-09")) == 12  # 上位 10 本
    assert table_pt(_slide(spec, prs, "SS-12-1")) == 11  # 付録


def test_table_is_native_with_header_links_and_right_aligned_numbers(
    built: tuple[DeckSpec, Any],
) -> None:
    spec, prs = built
    slide = _slide(spec, prs, "SS-09")
    frame = next(s for s in slide.shapes if getattr(s, "has_table", False))
    assert frame.name == "SS-09｜上位 10 本"
    table = frame.table
    header = [table.cell(0, c).text for c in range(len(table.columns))]
    assert header[:2] == ["順位", "アカウント（元投稿）"] and "保存率" in header
    link = table.cell(1, 1).text_frame.paragraphs[0].runs[0].hyperlink.address
    assert link and link.startswith("https://www.tiktok.com/@gonosara/")
    plays_col = header.index("再生")
    from pptx.enum.text import PP_ALIGN

    assert table.cell(1, plays_col).text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
    tc_pr = table.cell(0, 0)._tc.tcPr
    assert tc_pr.find(f"{_A}solidFill/{_A}schemeClr").get("val") == "accent1"
    # 罫線は横線だけ（accent4）・縦線なし
    body_pr = table.cell(1, 0)._tc.tcPr
    assert body_pr.find(f"{_A}lnB/{_A}solidFill/{_A}schemeClr").get("val") == "accent4"
    assert body_pr.find(f"{_A}lnL/{_A}noFill") is not None


def test_chart_is_native_with_embedded_data_top_first(built: tuple[DeckSpec, Any]) -> None:
    spec, prs = built
    frame = next(s for s in _slide(spec, prs, "SS-07").shapes if getattr(s, "has_chart", False))
    assert frame.name == "SS-07｜投稿者の構成"
    chart = frame.chart
    assert chart.chart_type == XL_CHART_TYPE.BAR_STACKED_100
    assert chart.part.chart_workbook.xlsx_part is not None  # 「データの編集」で開ける
    assert chart.category_axis.reverse_order is True
    assert list(chart.plots[0].categories) == ["本数", "再生"]
    assert chart.plots[0].data_labels.number_format == "0%"
    assert chart.has_legend


def test_pictures_are_not_cropped_and_have_alt_text_and_links(
    built: tuple[DeckSpec, Any],
) -> None:
    spec, prs = built
    slide = _slide(spec, prs, "SS-08")
    pictures = [s for s in slide.shapes if s.__class__.__name__ == "PlaceholderPicture"]
    assert len(pictures) == 5
    for pic in pictures:
        assert (pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom) == (0, 0, 0, 0)
        descr = pic._element.nvPicPr.cNvPr.get("descr")
        assert descr.endswith("表紙") and " 位 @" in descr
        assert pic.click_action.hyperlink.address.startswith("https://www.tiktok.com/")
    second = next(p for p in pictures if p._element.nvPicPr.cNvPr.get("descr").startswith("2 位"))
    assert second.width / second.height == pytest.approx(3 / 4, rel=0.01)  # 3:4 のまま収める


def test_sections_and_document_properties(built: tuple[DeckSpec, Any]) -> None:
    _, prs = built
    xml = prs.part._element.xml
    for name in ("表紙", "本編", "付録"):
        assert f'name="{name}"' in xml
    props = prs.core_properties
    assert props.identifier == REPORT_ID
    assert REPORT_ID in props.comments and "取得時点" in props.comments
    assert props.title == "「スパイスカレー 作り方」TikTok 検索 上位 15 本の顔ぶれ"


def test_csv_has_bom_and_all_rows(built: tuple[DeckSpec, Any]) -> None:
    spec, _ = built
    body = render_deck_csv(spec)
    assert body is not None and body.startswith("﻿".encode())
    lines = body.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("検索語,媒体,順位")
    assert len(lines) == 16


def test_files_entry_and_missing_image_is_rejected(built: tuple[DeckSpec, Any]) -> None:
    spec, _ = built
    with pytest.raises(DeckRenderError, match="missing"):
        render_deck_files(spec.model_dump_json(), {})
    s = surface(3)
    deck = build_surface_deck([s], client_name=None, measured_epoch=NOW, report_id=REPORT_ID)
    files = render_deck_files(deck.spec.model_dump_json(), deck.images)
    assert set(files) == {"deck.pptx", "deck.csv"}
    prs = Presentation(io.BytesIO(files["deck.pptx"]))
    cover = prs.slides[0]
    frames = [s for s in cover.placeholders if s.placeholder_format.type == PP_PLACEHOLDER.PICTURE]
    assert len(frames) == 1  # 画像が無くても「枠だけ」は残す


def _picture_package_errors(body: bytes) -> list[str]:
    """PowerPoint が表示できる画像の関係・MIME・バイト列・形状を直接調べる。"""
    import posixpath
    import zipfile
    from xml.etree import ElementTree as ET

    from PIL import Image

    ns = {
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    errors = []
    count = 0
    with zipfile.ZipFile(io.BytesIO(body)) as package:
        ct = ET.fromstring(package.read("[Content_Types].xml"))
        defaults = {
            c.attrib["Extension"]: c.attrib["ContentType"] for c in ct if "Extension" in c.attrib
        }
        overrides = {
            c.attrib["PartName"]: c.attrib["ContentType"] for c in ct if "PartName" in c.attrib
        }
        for name in package.namelist():
            if not name.startswith("ppt/slides/slide") or not name.endswith(".xml"):
                continue
            root = ET.fromstring(package.read(name))
            relname = posixpath.join(
                posixpath.dirname(name), "_rels", posixpath.basename(name) + ".rels"
            )
            rels = {r.attrib["Id"]: r.attrib for r in ET.fromstring(package.read(relname))}
            for pic in root.findall(".//p:pic", ns):
                count += 1
                geom = pic.find("p:spPr/a:prstGeom", ns)
                if geom is None or geom.attrib.get("prst") != "rect":
                    errors.append("picture geometry")
                blip = pic.find("p:blipFill/a:blip", ns)
                rid = blip.attrib.get("{" + ns["r"] + "}embed") if blip is not None else None
                rel = rels.get(rid, {})
                if rel.get("TargetMode") == "External" or not rel.get("Type", "").endswith(
                    "/image"
                ):
                    errors.append("image relationship")
                    continue
                target = posixpath.normpath(posixpath.join(posixpath.dirname(name), rel["Target"]))
                data = package.read(target)
                ext = target.rsplit(".", 1)[1]
                mime = overrides.get("/" + target, defaults.get(ext))
                expected = (
                    "image/png"
                    if data.startswith(b"\x89PNG\r\n\x1a\n")
                    else "image/jpeg"
                    if data.startswith(b"\xff\xd8")
                    else None
                )
                if not expected or mime != expected:
                    errors.append("image content type or signature")
                with Image.open(io.BytesIO(data)) as im:
                    im.verify()
        if count != 6:
            errors.append(f"picture count: {count}")
    return errors


def _image_deck_bytes() -> bytes:
    from PIL import Image

    from tests.skills.search_surface_check.deck_fixtures import png

    s = surface(5)
    images = covers(s.posts)
    jpeg = io.BytesIO()
    Image.open(io.BytesIO(png(9, 16))).save(jpeg, format="JPEG")
    images[s.posts[0].url] = jpeg.getvalue()
    deck = build_surface_deck(
        [s], client_name=None, measured_epoch=NOW, report_id=REPORT_ID, covers=images
    )
    return render_deck_pptx(deck.spec, deck.images)


def test_embedded_images_have_geometry_relationships_content_types_and_real_formats() -> None:
    body = _image_deck_bytes()
    assert _picture_package_errors(body) == []
    assert len(Presentation(io.BytesIO(body)).slides) > 1


@pytest.mark.parametrize("mutation", ["geometry", "content_type", "relationship"])
def test_image_package_checker_rejects_mutations(mutation: str) -> None:
    import zipfile
    from xml.etree import ElementTree as ET

    body = _image_deck_bytes()
    output = io.BytesIO()
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(body)) as source, zipfile.ZipFile(output, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if mutation == "geometry" and name.endswith(".xml") and name.startswith("ppt/slide"):
                root = ET.fromstring(data)
                for parent in root.iter():
                    for child in list(parent):
                        if child.tag == a + "prstGeom":
                            parent.remove(child)
                data = ET.tostring(root)
            elif mutation == "content_type" and name == "[Content_Types].xml":
                data = data.replace(b"image/jpeg", b"image/png")
            elif mutation == "relationship" and name.startswith("ppt/slides/_rels/"):
                data = data.replace(b"relationships/image", b"relationships/hyperlink")
            target.writestr(name, data)
    assert _picture_package_errors(output.getvalue())
    assert not _picture_package_errors(body)


def test_empty_table_cells_have_font_size_and_safe_numeric_column_spacing(
    built: tuple[DeckSpec, Any],
) -> None:
    spec, prs = built
    empty = 0
    for ss, slide in zip(spec.slides, prs.slides, strict=True):
        fill = next((f for f in ss.fills if f.kind == "table"), None)
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            assert fill is not None
            for row in shape.table.rows:
                for col, cell in enumerate(row.cells):
                    for para in cell.text_frame.paragraphs:
                        end = para._p.find(_A + "endParaRPr")
                        assert end is not None and end.get("sz") == str(int(fill.font_pt * 100))
                        assert end.get("lang") == "ja-JP"
                        if not cell.text:
                            empty += 1
                            assert not para.runs
                    if col > 0 and col - 1 in fill.numeric_cols and col not in fill.numeric_cols:
                        assert cell.margin_left >= 182880
    assert empty > 0


def test_notes_theme_and_dynamic_chart_labels(built: tuple[DeckSpec, Any]) -> None:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    _, prs = built
    theme = prs.notes_master.part.part_related_by(RT.THEME).blob.decode()
    assert theme.count('typeface="游ゴシック"') == 4
    assert not any(font in theme for font in ("Calibri", "ＭＳ", "Arial"))
    charts = [sh.chart for sl in prs.slides for sh in sl.shapes if sh.has_chart]
    assert charts
    c = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    for chart in charts:
        assert chart._chartSpace.find(".//" + c + "dLbl") is None
        assert chart._chartSpace.find(".//" + c + "dLbls/" + c + "delete") is None
        formats = chart._chartSpace.findall(".//" + c + "ser/" + c + "dLbls/" + c + "numFmt")
        assert formats and all(f.get("formatCode").startswith('[<0.06]"";') for f in formats)


@pytest.mark.parametrize(
    "value",
    [
        '=HYPERLINK("x","y")',
        "+α",
        "-cmd",
        "@SUM(1+1)",
        "\t=1+1",
        "\r=1+1",
        " =1+1",
        "@普通のメンション",
    ],
)
def test_csv_neutralizes_formula_cells(value: str) -> None:
    import csv

    from teamagent.media.deck_contracts import CsvSpec

    s = surface(3)
    spec = build_surface_deck([s], client_name=None, measured_epoch=NOW, report_id=REPORT_ID).spec
    # header、全列を同じ入口で無害化する。
    spec = spec.model_copy(
        update={"csv": CsvSpec(columns=(value, "本文", "再生"), rows=((value, value, "123"),))}
    )
    data = render_deck_csv(spec)
    assert data is not None
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0][0] == "\u200b" + value
    assert rows[1] == ["\u200b" + value, "\u200b" + value, "123"]


def test_table_has_one_row_of_editing_space_and_conclusion_fits() -> None:
    from teamagent.skills._deck.layouts import report_template

    template = report_template()
    caption = template.layout("R_動画カード").box("caption_1")
    assert caption.h >= int(1.8 * 914400)
    assert caption.y + caption.h < template.layout("R_動画カード").box("footer").y
    table = template.layout("R_表").box("table")
    note = template.layout("R_表").box("note")
    assert table.y + table.h + table.h // 11 <= note.y
    body = template.layout("R_結論").box("body")
    assert body.font_pt == 16 and body.space_before_pt == 26
    # 全角100字/点、游ゴシック16pt、約44字/行、3行/点を保守的に見積もる。
    height_in = (9 * 16 * 1.4 + 3 * 26) / 72
    assert height_in <= body.h / 914400
    assert (3 * 16 * 1.4 + 3 * 26) / 72 >= body.h / 914400 * 0.4
