"""LibreOffice 不要の合成座標・合成デッキによる描画 QA 契約。"""

from __future__ import annotations

import json
import subprocess
import sys
import time

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from teamagent.media.render_qa import (
    Glyph,
    ShapeBox,
    evaluate,
    load_profile,
    prepare_probe,
    probe_color,
)
from teamagent.media.render_qa_child import execute, run, soffice_command


def box(key="a", rect=(0, 0, 100, 100), **kwargs):
    return ShapeBox(
        1, key, key, rect, rect, probe_color(1 if key == "a" else 2), has_text=True, **kwargs
    )


def glyph(key="a", x=10, y=20, size=10, text="合成", font="Arial"):
    return Glyph(
        1, (x, y - size, x + 8, y + 2), y, size, probe_color(1 if key == "a" else 2), font, text
    )


def kinds(report):
    return {f["kind"] for f in report["findings"]}


def test_fits_and_baseline_descender_overflow():
    shape = box(rect=(0, 0, 100, 100))
    assert evaluate([shape], [glyph(y=90)], 200, 200)["status"] == "ok"
    # baseline 自体は枠内。+0.25em がなければ必ず見逃す境界ケース。
    report = evaluate([shape], [glyph(y=59, size=40), glyph(y=99, size=40)], 200, 200)
    assert "text_overflow_bottom" in kinds(report)
    assert report["findings"][0]["severity"] == "warn"


@pytest.mark.parametrize("y,severity", [(102, "warn"), (108, "error")])
def test_bottom_thresholds(y, severity):
    report = evaluate([box()], [glyph(y=y)], 200, 200)
    assert report["findings"][0]["kind"] == "text_overflow_bottom"
    assert report["findings"][0]["severity"] == severity


def test_right_and_outside_slide():
    report = evaluate([box()], [glyph(x=110)], 115, 100)
    assert {"text_overflow_right", "text_outside_slide"} <= kinds(report)


def test_collisions_and_exact_template_baseline():
    shapes = [box(), box("b")]
    chars = [glyph(), glyph("b")]
    report = evaluate(shapes, chars, 200, 200)
    finding = next(f for f in report["findings"] if f["kind"] == "text_collision")
    baseline = [
        {
            "slide": 1,
            "shape_key": "a",
            "kind": "text_collision",
            "details": {k: v for k, v in finding["details"].items() if k != "shape_key"},
        }
    ]
    excluded = evaluate(shapes, chars, 200, 200, baseline=baseline)
    assert excluded["status"] == "ok"
    assert excluded["counts"]["baseline_excluded"] == 1
    assert "text_collision" in kinds(
        evaluate(shapes, [glyph(x=11), glyph("b")], 200, 200, baseline=baseline)
    )


def test_autofit_picture_and_missing():
    growing = box(rect=(0, 0, 100, 20), autofit="spAutoFit")
    neighbour = ShapeBox(1, "pic", "image", (0, 20, 100, 80), (0, 20, 100, 80), picture=True)
    report = evaluate([growing, neighbour, box("b")], [glyph(y=25)], 200, 200)
    assert {"autofit_grow_collision", "text_over_picture", "text_missing"} <= kinds(report)
    assert "text_overflow_bottom" not in kinds(report)


def test_rotation_counts_only():
    report = evaluate([box(rotation=45)], [glyph(x=500)], 200, 200)
    assert report["findings"] == []
    assert report["counts"]["unsupported_rotation"] == 1


@pytest.mark.parametrize(
    "text,font,kind",
    [
        ("□", "Arial", "notdef"),
        ("(cid:0)", "Arial", "notdef"),
        # 見本の字は、結果の欄の名前（expected 等）や書体名と重ならない字にする（1 文字の x は誤検知する）。
        ("本", ".notdef", "notdef"),
        ("本", "Unexpected", "glyph_fallback"),
        ("😀", "Arial", "glyph_fallback"),
    ],
)
def test_missing_glyphs_and_fallback(text, font, kind):
    report = evaluate([box(expected_fonts=("Arial",))], [glyph(text=text, font=font)], 200, 200)
    assert kind in kinds(report)
    assert text not in json.dumps(report, ensure_ascii=False)


def test_line_clustering_and_font_substitute_warning_width():
    shape = box()
    report = evaluate([shape], [glyph(y=20), glyph(x=20, y=20.5), glyph(y=102)], 200, 200)
    finding = (
        next(f for f in report["findings"] if f["kind"] == "text_overflow_bottom")
        if "text_overflow_bottom" in kinds(report)
        else None
    )
    # 2 baseline clusters, spacing is observed (81.5pt), not glyph bbox heights.
    assert finding is None
    assert (
        evaluate([shape], [glyph(y=102)], 200, 200, font_substitute="plain_noto")["findings"] == []
    )
    assert (
        evaluate([shape], [glyph(y=110)], 200, 200, font_substitute="plain_noto")["counts"]["error"]
        == 1
    )


def test_unmatched_color_is_review_and_shape_names_are_bounded():
    char = glyph()
    char.color = (255, 255, 255)
    shape = box()
    shape.name = "名前" * 100
    report = evaluate([shape], [char], 200, 200)
    assert report["status"] == "needs_review"
    assert report["counts"]["unmatched_glyphs"] == 1
    assert len(report["findings"][0]["shape"]) == 40


def test_probe_colors_are_unique_and_separated():
    colors = [probe_color(i) for i in range(2000)]
    assert len(set(colors)) == len(colors)
    # 全ペアの差は、少なくとも1チャネルで8以上（単なる定数検査にしない）。
    for i, color in enumerate(colors):
        assert all(
            max(abs(a - b) for a, b in zip(color, other, strict=True)) >= 8 for other in colors[:i]
        )


def test_probe_copy_preserves_geometry_and_table_cells(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "合成サンプル"
    run.font.size, run.font.name, run.font.bold = Pt(23), "Yu Gothic UI", True
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "左", "右"
    source, output = tmp_path / "source.pptx", tmp_path / "probe.pptx"
    deck.save(source)
    boxes, width, height = prepare_probe(source, output)
    probe = Presentation(output)
    changed = probe.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert (changed.font.size, changed.font.name, changed.font.bold) == (
        run.font.size,
        run.font.name,
        run.font.bold,
    )
    assert width > 0 and height > 0
    assert len({s.color for s in boxes}) == 3
    assert source.read_bytes() != output.read_bytes()
    assert (
        Presentation(source).slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.color.type
        is None
    )


def test_nested_group_geometry(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    shape.text = "合成"
    group.left, group.top, group.width, group.height = Inches(2), Inches(2), Inches(4), Inches(2)
    source = tmp_path / "group.pptx"
    deck.save(source)
    boxes, _, _ = prepare_probe(source, tmp_path / "probe.pptx")
    assert boxes[0].rect == pytest.approx((144, 144, 432, 288))


def test_profile_and_soffice_security(tmp_path):
    thresholds, baseline = load_profile("proposal_builder_v1")
    assert thresholds["error_lines"] == 0.5 and baseline == []
    with pytest.raises(ValueError, match="unknown_profile"):
        load_profile("../../secrets")
    command = soffice_command(tmp_path)
    assert "--infilter=Impress MS PowerPoint 2007 XML" in command
    assert not any("--accept" in item for item in command)
    assert any('"ExportHiddenSlides"' in item and '"true"' in item for item in command)
    assert any(tmp_path.as_uri() in item for item in command)


def test_no_renderer_is_unavailable(tmp_path, monkeypatch):
    monkeypatch.setattr("teamagent.media.render_qa_child.shutil.which", lambda _: None)
    result = execute(tmp_path / "missing.pptx", tmp_path / "output", None)
    assert result["status"] == "unavailable"
    assert (tmp_path / "output/render_qa.json").is_file()


def test_timeout_kills_process_group(tmp_path, monkeypatch):
    # macOS は RLIMIT_AS の finite 設定を拒否する。実イメージは Linux。
    import resource

    original = resource.setrlimit

    def portable_limit(kind, values):
        if sys.platform == "darwin" and kind in {resource.RLIMIT_AS, resource.RLIMIT_NPROC}:
            return
        original(kind, values)

    monkeypatch.setattr(resource, "setrlimit", portable_limit)
    marker = tmp_path / "orphan"
    # 孫プロセスが生き残ればあとで marker を作る。ネットワーク・外部ツール不要。
    script = "import subprocess,sys,time; subprocess.Popen([sys.executable,'-c',sys.argv[1]]); time.sleep(30)"
    child = f"import time,pathlib; time.sleep(0.8); pathlib.Path({str(marker)!r}).touch()"
    with pytest.raises(subprocess.TimeoutExpired):
        run([sys.executable, "-c", script, child], time.monotonic() + 0.25)
    time.sleep(1)
    assert not marker.exists()


def test_linux_resource_limits_are_all_applied(monkeypatch):
    import resource

    from teamagent.media.render_qa_child import limits

    applied = {}
    monkeypatch.setattr(resource, "setrlimit", lambda kind, value: applied.update({kind: value}))
    limits()
    assert applied == {
        resource.RLIMIT_CPU: (300, 300),
        resource.RLIMIT_AS: (3 << 30, 3 << 30),
        resource.RLIMIT_FSIZE: (1 << 30, 1 << 30),
        resource.RLIMIT_NOFILE: (1024, 1024),
        resource.RLIMIT_NPROC: (256, 256),
        resource.RLIMIT_CORE: (0, 0),
    }


def test_pdfminer_extracts_color_and_matrix_baseline(tmp_path):
    pytest.importorskip("pdfminer.high_level", reason="render-qa-runtime 未導入の CI")
    from teamagent.media.render_qa import extract_glyphs

    # 小さな合成 PDF。実デッキ・LibreOffice・追加 PDF ライブラリ不要。
    stream = b"BT /F1 10 Tf 0 0 0.031372549 rg 1 0 0 1 20 70 Tm (A) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 100] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(pdf)
    pdf += b"xref\n0 6\n0000000000 65535 f \n"
    for offset in offsets[1:]:
        pdf += f"{offset:010d} 00000 n \n".encode()
    pdf += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path = tmp_path / "synthetic.pdf"
    path.write_bytes(pdf)
    glyphs, pages = extract_glyphs(path)
    assert pages == [(200, 100)]
    assert len(glyphs) == 1
    assert glyphs[0].color == probe_color(1)
    assert glyphs[0].baseline == 30
    assert glyphs[0].rect[0] == 20 and glyphs[0].size == 10


def test_contact_sheet_and_flagged_detail_limit(tmp_path):
    from PIL import Image

    from teamagent.media.render_qa_child import make_thumbnails

    findings = [{"slide": n, "details": {"rect": [10, 10, 30, 30]}} for n in range(1, 13)]
    (tmp_path / "result.json").write_text(json.dumps({"findings": findings}))
    (tmp_path / "geometry.json").write_text(json.dumps({"width": 200, "height": 100}))
    for n in range(1, 13):
        Image.new("RGB", (320, 160), "white").save(tmp_path / f"thumb-{n}.png")
        if n <= 10:
            Image.new("RGB", (1280, 640), "white").save(tmp_path / f"detail-{n}.png")
    make_thumbnails(tmp_path)
    assert len(list(tmp_path.glob("slide-*.png"))) == 10
    with Image.open(tmp_path / "contact_sheet.png") as image:
        assert image.getpixel((6, 6)) == (255, 0, 0)
    with Image.open(tmp_path / "slide-001.png") as image:
        assert image.getpixel((64, 64)) == (255, 0, 0)


def test_empty_slides_still_have_page_identity(tmp_path):
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.slides.add_slide(deck.slide_layouts[6])
    source = tmp_path / "empty.pptx"
    deck.save(source)
    boxes, width, height = prepare_probe(source, tmp_path / "probe.pptx")
    result = evaluate(boxes, [], width, height)
    assert result["counts"]["slides"] == 2
    assert result["status"] == "ok"


@pytest.mark.parametrize("family", ["Yu Gothic UI", "Yu Gothic", "游ゴシック", "游ゴシック体"])
def test_palt_substitute_is_allowed_and_plain_noto_requires_review(family):
    shape = box(expected_fonts=(family,))
    palt = evaluate([shape], [glyph(font="TAPaltGothic-Regular")], 200, 200)
    assert palt["font_substitute"] == "palt"
    assert palt["status"] == "ok"
    plain = evaluate(
        [shape], [glyph(font="NotoSansCJKJP-Regular")], 200, 200, font_substitute="plain_noto"
    )
    assert plain["font_substitute"] == "plain_noto"
    assert plain["findings"] == []
    assert plain["status"] == "needs_review"
    unexpected = evaluate([shape], [glyph(font="NotoSansCJKJP-Regular")], 200, 200)
    assert "glyph_fallback" in kinds(unexpected)
