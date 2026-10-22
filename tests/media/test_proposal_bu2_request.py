"""BU2 quantities must survive an exact numeric-token check against the user's text."""

from decimal import Decimal

import pytest

from teamagent.media.proposal_bu2 import BU2Content, bu2_auxiliary_values, bu2_source_notes
from tests.skills.proposal_deck.test_bu2_content import bu2_payload

REQUEST_TEXT = "2本、100,001回、2.5円/回、4週間、4回接触、理解率60%、管理費10%"


@pytest.mark.parametrize(
    "text",
    [REQUEST_TEXT, "２本、１０．０００１万回、２．５円／回、４週間、４回接触、６０％、１０％"],
)
def test_request_numbers_and_notation_variants_are_kept(text: str) -> None:
    original = BU2Content.model_validate(bu2_payload())
    grounded = original.grounded_in_request(text)
    assert grounded == original
    assert BU2Content.model_validate_json(grounded.model_dump_json()) == grounded


@pytest.mark.parametrize("text", ["", "提案資料を作成してください", "2024年、20本、0.4円、60万回"])
def test_missing_or_unmatched_request_removes_all_numbers(text: str) -> None:
    original = BU2Content.model_validate(bu2_payload())
    grounded = original.grounded_in_request(text)
    assert grounded.plans[0].count is None
    assert grounded.plans[0].views is None
    assert grounded.plans[0].unit_price_yen is None
    assert grounded.plans[0].period == ""
    assert grounded.estimate_terms.management_fee_rate is None
    assert grounded.simulation is not None
    assert grounded.simulation.views is None
    assert grounded.simulation.contacts is None
    assert grounded.simulation.understanding_rate is None
    values = bu2_auxiliary_values(grounded)
    for key in (
        "PB-BU2-PLAN-TRIAL-COUNT",
        "PB-BU2-PLAN-TRIAL-VIEWS",
        "PB-BU2-PLAN-TRIAL-UNIT-PRICE",
        "PB-BU2-PLAN-TRIAL-PERIOD",
        "PB-BU2-PLAN-TRIAL-AMOUNT",
        "PB-BU2-ESTIMATE-NOTES",
        "PB-BU2-SIM-VIEWS",
        "PB-BU2-SIM-CONTACTS",
        "PB-BU2-SIM-RATE",
        "PB-BU2-SIM-REACH",
        "PB-BU2-SIM-RESULT",
        "PB-BU2-SIM-ASSUMPTION",
    ):
        assert "依頼に記載が無いため未記入" in values[key]
    assert grounded.hearing_items == original.hearing_items
    assert original == BU2Content.model_validate(bu2_payload())


def test_matching_simulation_values_survive_but_dependent_results_do_not() -> None:
    grounded = BU2Content.model_validate(bu2_payload()).grounded_in_request("100001回、理解率60%")
    assert grounded.simulation is not None
    assert grounded.simulation.views == 100001
    assert grounded.simulation.contacts is None
    assert grounded.simulation.understanding_rate == Decimal("0.6")
    values = bu2_auxiliary_values(grounded)
    assert values["PB-BU2-SIM-VIEWS"] == "100,001回（想定）"
    assert values["PB-BU2-SIM-RATE"] == "60%（仮定）"
    assert "要確認" in values["PB-BU2-SIM-REACH"]
    assert "要確認" in values["PB-BU2-SIM-RESULT"]


@pytest.mark.parametrize("text", ["40本、2024年、0.4円", "4万回", "4%", "-4回", "−4回"])
def test_number_components_do_not_authorize_other_values(text: str) -> None:
    payload = {"plans": [{"name": "Trial", "count": 4, "assumption": "依頼条件"}]}
    assert BU2Content.model_validate(payload).grounded_in_request(text).plans[0].count is None


def test_ungrounded_numbers_in_assumptions_do_not_leak_into_notes() -> None:
    payload = bu2_payload()
    payload["plans"][0]["assumption"] = "独自に仮定した費用は777円"
    payload["simulation"]["assumption"] = "独自に仮定した接触は888回"
    grounded = BU2Content.model_validate(payload).grounded_in_request(REQUEST_TEXT)
    assert "777" not in str(bu2_source_notes(grounded))
    assert "888" not in bu2_auxiliary_values(grounded)["PB-BU2-SIM-ASSUMPTION"]


def test_removed_plan_assumption_is_listed_for_review_even_when_amount_matches() -> None:
    from pptx import Presentation

    from teamagent.media.proposal_cleanup import prepare_proposal_content

    payload = bu2_payload()
    payload["plans"][0]["assumption"] = "777円を仮定"
    grounded = BU2Content.model_validate(payload).grounded_in_request(REQUEST_TEXT)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slot = "PB-BU2-PLAN-TRIAL-AMOUNT"
    slide.shapes.add_textbox(0, 0, 100, 100).text = "{{" + slot + "}}"
    _, values = prepare_proposal_content(
        presentation, {}, bu2_auxiliary_values(grounded), bu2_source_notes(grounded)
    )
    assert values[slot] == "250,003 円"
    notes = slide.notes_slide.notes_text_frame.text
    assert "要確認一覧" in notes
    assert "依頼に記載が無いため未記入" in notes
    assert "777" not in notes


def test_every_period_number_must_be_present() -> None:
    payload = bu2_payload()
    payload["plans"][0]["period"] = "4〜8週間"
    assert BU2Content.model_validate(payload).grounded_in_request("4週間").plans[0].period == ""


def test_period_cannot_match_inside_a_larger_number() -> None:
    assert (
        BU2Content.model_validate(bu2_payload()).grounded_in_request("14週間").plans[0].period == ""
    )


def test_request_numbers_are_not_rounded_into_a_match() -> None:
    payload = bu2_payload()
    payload["plans"][0]["count"] = 1
    source = "0.999999999999999999999999999999本、理解率59.99999999999999999999999999999%"
    grounded = BU2Content.model_validate(payload).grounded_in_request(source)
    assert grounded.plans[0].count is None
    assert grounded.simulation.understanding_rate is None


@pytest.mark.parametrize("period", ["四週間", "短期", "4万5千日", "4.2.3日", "4e2日"])
def test_unrecognized_period_is_not_treated_as_having_no_numbers(period: str) -> None:
    payload = bu2_payload()
    payload["plans"][0]["period"] = period
    assert BU2Content.model_validate(payload).grounded_in_request("4週間").plans[0].period == ""


def test_direct_deck_call_without_request_text_drops_bu2_numbers(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from teamagent.skills.base import SkillContext
    from teamagent.skills.proposal_deck.schema import ProposalDeckInput
    from teamagent.skills.proposal_deck.skill import ProposalDeckSkill

    request = ProposalDeckInput(
        product_name="合成食品", goal="認知", target_persona="家庭", **bu2_payload()
    )
    captured = []

    def capture(input, _ctx):
        captured.append(input)
        raise RuntimeError("stop before rendering")

    deck = ProposalDeckSkill(bedrock=object())
    monkeypatch.setattr(deck, "_compose", capture)
    with pytest.raises(RuntimeError, match="stop before rendering"):
        deck.run(request, SkillContext())
    assert captured[0].plans[0].unit_price_yen is None
    assert captured[0].simulation.contacts is None
