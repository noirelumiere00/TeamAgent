"""ZIP/OOXML 消毒を合成パッケージだけで検証する。"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from teamagent.media import render_qa
from teamagent.media.render_qa import sanitize_pptx

CT = "http://schemas.openxmlformats.org/package/2006/content-types"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def package(tmp_path: Path, additions=None):
    source = tmp_path / "source.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(1), Inches(1)
    ).text = "合成"
    deck.save(source)
    with zipfile.ZipFile(source) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    parts.update(additions or {})
    return source, parts


def write(source, parts, compression=zipfile.ZIP_DEFLATED):
    with zipfile.ZipFile(source, "w", compression) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def test_external_ole_activex_vba_fonts_media_are_removed(tmp_path):
    source, parts = package(tmp_path)
    additions = {
        "ppt/embeddings/ole.bin": b"ole",
        "ppt/activeX/activeX.bin": b"activex",
        "ppt/vbaProject.bin": b"vba",
        "ppt/fonts/font.bin": b"font",
        "ppt/media/video.bin": b"video",
    }
    parts.update(additions)
    tree = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(
        tree, f"{{{CT}}}Default", Extension="bin", ContentType="application/octet-stream"
    )
    parts["[Content_Types].xml"] = etree.tostring(tree)
    relname = "ppt/slides/_rels/slide1.xml.rels"
    rels = etree.fromstring(parts[relname])
    etree.SubElement(
        rels,
        f"{{{REL}}}Relationship",
        Id="external",
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        Target="https://example.invalid",
        TargetMode="External",
    )
    for index, name in enumerate(additions):
        etree.SubElement(
            rels,
            f"{{{REL}}}Relationship",
            Id=f"danger{index}",
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject",
            Target="../" + name.removeprefix("ppt/"),
        )
    parts[relname] = etree.tostring(rels)
    write(source, parts)
    output = tmp_path / "safe.pptx"
    counts = sanitize_pptx(source, output)
    assert counts["external"] == 1
    assert all(counts[k] == 1 for k in ("embeddings", "activex", "vba", "fonts", "media"))
    with zipfile.ZipFile(output) as archive:
        assert not set(additions) & set(archive.namelist())
        assert b'TargetMode="External"' not in archive.read(relname)
    assert len(Presentation(output).slides) == 1


@pytest.mark.parametrize(
    "name", ["../evil.xml", "/evil.xml", "ppt/../../evil.xml", "ppt\\evil.xml", "C:evil.xml"]
)
def test_zip_traversal_rejected(tmp_path, name):
    source, parts = package(tmp_path, {name: b"evil"})
    write(source, parts)
    with pytest.raises(ValueError, match="zip_path"):
        sanitize_pptx(source, tmp_path / "safe.pptx")


def test_zip_ratio_bomb_rejected(tmp_path):
    source, parts = package(tmp_path, {"bomb.xml": b"0" * 100_000})
    write(source, parts)
    with pytest.raises(ValueError, match="zip_ratio"):
        sanitize_pptx(source, tmp_path / "safe.pptx")


@pytest.mark.parametrize("limit,value", [("MAX_ENTRIES", 1), ("MAX_UNCOMPRESSED", 10)])
def test_zip_limits(tmp_path, monkeypatch, limit, value):
    source, parts = package(tmp_path)
    write(source, parts)
    monkeypatch.setattr(render_qa, limit, value)
    with pytest.raises(ValueError, match="zip_limit"):
        sanitize_pptx(source, tmp_path / "safe.pptx")


def test_unknown_content_type_rejected(tmp_path):
    source, parts = package(tmp_path)
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"application/xml", b"application/x-evil"
    )
    write(source, parts)
    with pytest.raises(ValueError, match="content_type"):
        sanitize_pptx(source, tmp_path / "safe.pptx")


def test_doctype_rejected(tmp_path):
    source, parts = package(tmp_path)
    parts["[Content_Types].xml"] = (
        b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///secret">]>'
        + parts["[Content_Types].xml"].split(b"?>", 1)[1]
    )
    write(source, parts)
    with pytest.raises(ValueError, match="xml_doctype"):
        sanitize_pptx(source, tmp_path / "safe.pptx")


def test_bitmap_is_reencoded_and_svg_removed(tmp_path):
    image = Image.new("RGB", (4, 4), "blue")
    stream = io.BytesIO()
    image.save(stream, "PNG")
    source, parts = package(
        tmp_path,
        {
            "ppt/media/image.png": stream.getvalue() + b"private trailing payload",
            "ppt/media/image.svg": b"<svg/>",
        },
    )
    tree = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(tree, f"{{{CT}}}Default", Extension="svg", ContentType="image/svg+xml")
    etree.SubElement(tree, f"{{{CT}}}Default", Extension="png", ContentType="image/png")
    parts["[Content_Types].xml"] = etree.tostring(tree)
    write(source, parts)
    output = tmp_path / "safe.pptx"
    counts = sanitize_pptx(source, output)
    assert counts["images_reencoded"] >= 1 and counts["media"] == 1
    with zipfile.ZipFile(output) as archive:
        assert "ppt/media/image.svg" not in archive.namelist()
        assert b"private trailing payload" not in archive.read("ppt/media/image.png")


def test_duplicate_entry_rejected(tmp_path):
    source, parts = package(tmp_path)
    write(source, parts)
    with pytest.warns(UserWarning), zipfile.ZipFile(source, "a") as archive:
        archive.writestr("ppt/presentation.xml", b"duplicate")
    with pytest.raises(ValueError, match="zip_path"):
        sanitize_pptx(source, tmp_path / "safe.pptx")
