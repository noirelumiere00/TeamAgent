from __future__ import annotations

import io
import json
import subprocess
import zipfile
from pathlib import Path
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from teamagent.media.deck_qa import inspect_pptx
from teamagent.media.render_child import _deck_qa
from teamagent.skills.proposal_deck.renderer import _add_picture_fit


def image(size: tuple[int, int] = (200, 100)) -> io.BytesIO:
    stream = io.BytesIO()
    Image.new("RGB", size, "white").save(stream, "PNG")
    stream.seek(0)
    return stream


def deck() -> tuple[Any, Any]:
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def kinds(prs: Any, tmp_path: Path, **kwargs: Any) -> set[str]:
    path = tmp_path / "test.pptx"
    prs.save(path)
    return {f.kind for f in inspect_pptx(path, **kwargs).findings}


@pytest.mark.parametrize("size", [(400, 100), (100, 400), (200, 200), (1000, 1), (1, 1000)])
def test_fit_centers_and_contains(size: tuple[int, int]) -> None:
    _prs, slide = deck()
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(2))
    _add_picture_fit(slide, image(size).getvalue(), box)
    pic = slide.shapes[-1]
    assert box.left <= pic.left <= pic.left + pic.width <= box.left + box.width
    assert box.top <= pic.top <= pic.top + pic.height <= box.top + box.height
    assert abs((pic.left + pic.width / 2) - (box.left + box.width / 2)) <= 1
    assert abs((pic.top + pic.height / 2) - (box.top + box.height / 2)) <= 1
    assert pic.width / pic.height == pytest.approx(size[0] / size[1], rel=0.001)


@pytest.mark.parametrize("outside", [False, True])
def test_bounds(tmp_path: Path, outside: bool) -> None:
    prs, slide = deck()
    slide.shapes.add_picture(
        image(), prs.slide_width - Inches(1) + (Inches(0.1) if outside else 0), 0, width=Inches(1)
    )
    assert ("out_of_bounds" in kinds(prs, tmp_path)) is outside


@pytest.mark.parametrize("grouped", [False, True])
def test_nested_group_bounds(tmp_path: Path, grouped: bool) -> None:
    prs, slide = deck()
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_picture(image(), 0, 0, width=Inches(2))
    outer.left = prs.slide_width - Inches(1) if grouped else Inches(1)
    outer.width = Inches(2)
    assert ("out_of_bounds" in kinds(prs, tmp_path)) is grouped


@pytest.mark.parametrize("distorted", [False, True])
def test_aspect_crop(tmp_path: Path, distorted: bool) -> None:
    prs, slide = deck()
    pic = slide.shapes.add_picture(image(), 0, 0, width=Inches(2), height=Inches(2))
    if not distorted:
        pic.crop_left = pic.crop_right = 0.25
    assert ("image_aspect" in kinds(prs, tmp_path)) is distorted


@pytest.mark.parametrize("overflow", [False, True])
def test_frame(tmp_path: Path, overflow: bool) -> None:
    prs, slide = deck()
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    slide.shapes.add_picture(image(), Inches(1.1 if overflow else 1), Inches(1), width=Inches(2))
    assert ("image_frame_overflow" in kinds(prs, tmp_path)) is overflow


@pytest.mark.parametrize("long", [False, True])
@pytest.mark.parametrize("auto", [MSO_AUTO_SIZE.NONE, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE])
def test_text(tmp_path: Path, long: bool, auto: Any) -> None:
    prs, slide = deck()
    frame = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame
    frame.auto_size = auto
    frame.word_wrap = True
    frame.text = "あ" * (100 if long else 5)
    frame.paragraphs[0].runs[0].font.size = Pt(20)
    path = tmp_path / "test.pptx"
    prs.save(path)
    found = [f for f in inspect_pptx(path).findings if f.kind == "text_overflow"]
    assert bool(found) is long
    if found:
        assert found[0].severity == ("warn" if auto == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE else "error")


@pytest.mark.parametrize("collision", [False, True])
def test_shape_autofit(tmp_path: Path, collision: bool) -> None:
    prs, slide = deck()
    frame = slide.shapes.add_textbox(0, 0, Inches(2), Inches(0.3)).text_frame
    frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    frame.word_wrap = True
    frame.text = "あ" * 40
    if collision:
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.5), Inches(2), Inches(1))
    assert ("text_autofit_collision" in kinds(prs, tmp_path)) is collision


@pytest.mark.parametrize("marker", ["{{x}}", "<<x", "TODO", "xxx", "○○", "要確認（データ未検出）"])
def test_markers(tmp_path: Path, marker: str) -> None:
    prs, slide = deck()
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(2)).text = marker
    assert "placeholder" in kinds(prs, tmp_path)


def test_marker_false_positive_and_review_count(tmp_path: Path) -> None:
    prs, slide = deck()
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(2)).text = "TODOLIST axxxb 要確認 要確認"
    assert "placeholder" not in kinds(prs, tmp_path)
    report = inspect_pptx(tmp_path / "test.pptx")
    assert next(f for f in report.findings if f.kind == "review_required").details["count"] == 2


@pytest.mark.parametrize("empty", [False, True])
def test_empty_text(tmp_path: Path, empty: bool) -> None:
    prs, slide = deck()
    slide.shapes.add_textbox(0, 0, Inches(3), Inches(2)).text = "" if empty else "済"
    assert ("empty_text" in kinds(prs, tmp_path)) is empty


@pytest.mark.parametrize("overlap", [False, True])
def test_overlap(tmp_path: Path, overlap: bool) -> None:
    prs, slide = deck()
    for x in (0, 0.5 if overlap else 4):
        slide.shapes.add_textbox(Inches(x), 0, Inches(3), Inches(1)).text = "説明"
    assert ("text_overlap" in kinds(prs, tmp_path)) is overlap


def test_background_edge_and_decorative_band(tmp_path: Path) -> None:
    prs, slide = deck()
    slide.shapes.add_picture(
        image((400, 300)), 0, 0, width=prs.slide_width, height=prs.slide_height
    )
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.5))
    slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "説明"
    assert not {"out_of_bounds", "image_frame_overflow", "text_overlap"} & kinds(prs, tmp_path)


@pytest.mark.parametrize("expected", [1, 2])
def test_count(tmp_path: Path, expected: int) -> None:
    prs, _ = deck()
    assert ("slide_count" in kinds(prs, tmp_path, expected_slides=expected)) is (expected != 1)


@pytest.mark.parametrize("font", ["Arial", "Nonstandard Test Font"])
def test_font(tmp_path: Path, font: str) -> None:
    prs, slide = deck()
    frame = slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text_frame
    frame.text = "説明"
    frame.paragraphs[0].runs[0].font.name = font
    assert ("font_unavailable" in kinds(prs, tmp_path)) is (font != "Arial")
    assert font in inspect_pptx(tmp_path / "test.pptx").fonts


def mutate(path: Path, change: str) -> None:
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    key = next(n for n in files if n.startswith("ppt/media/"))
    if change == "missing":
        del files[key]
    elif change == "empty":
        files[key] = b""
    elif change == "webp":
        files[key] = b"RIFF\x10\0\0\0WEBP" + b"\0" * 30
    elif change == "avif":
        files[key] = b"\0\0\0\x20ftypavif" + b"\0" * 30
    elif change == "types":
        files["[Content_Types].xml"] = files["[Content_Types].xml"].replace(
            b'<Default Extension="png" ContentType="image/png"/>', b""
        )
    elif change == "mismatch":
        files["[Content_Types].xml"] = files["[Content_Types].xml"].replace(
            b"image/png", b"image/jpeg"
        )
    elif change == "broken":
        files[key] = b"broken"
    with zipfile.ZipFile(path, "w") as z:
        for n, data in files.items():
            z.writestr(n, data)


@pytest.mark.parametrize(
    "change", ["missing", "empty", "webp", "avif", "types", "mismatch", "broken", "valid"]
)
def test_package_images(tmp_path: Path, change: str) -> None:
    prs, slide = deck()
    slide.shapes.add_picture(image(), 0, 0, width=Inches(2))
    path = tmp_path / "test.pptx"
    prs.save(path)
    if change != "valid":
        mutate(path, change)
    findings = inspect_pptx(path).findings
    assert any(f.kind.startswith("image_") or f.kind == "package_unreadable" for f in findings) is (
        change != "valid"
    )
    if change == "webp":
        assert any(f.details.get("problem") == "signature_mismatch" for f in findings)
    if change == "missing":
        assert any(
            f.kind == "image_missing" and f.slide == 1 and f.shape.startswith("Picture")
            for f in findings
        )


def test_summary_and_failure_isolation(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    prs, _ = deck()
    path = tmp_path / "test.pptx"
    prs.save(path)
    assert _deck_qa(path, expected_slides=1)["deck_qa"]["error_count"] == 0

    def fail(*args: Any, **kwargs: Any) -> Any:
        raise ValueError("QA failure")

    monkeypatch.setattr("teamagent.media.deck_qa.inspect_pptx", fail)
    assert _deck_qa(path, expected_slides=1)["qa_error"] == "ValueError"


@pytest.mark.parametrize("expected", [1, 2])
def test_cli(tmp_path: Path, expected: int) -> None:
    import sys

    prs, _ = deck()
    path = tmp_path / "test.pptx"
    prs.save(path)
    run = subprocess.run(
        [
            sys.executable,
            "scripts/deck_qa.py",
            str(path),
            "--json",
            "--expected-slides",
            str(expected),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert run.returncode == (expected != 1)
    assert json.loads(run.stdout)["slide_count"] == 1


@pytest.mark.parametrize("qa_fails", [False, True])
def test_proposal_render_records_qa(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, qa_fails: bool
) -> None:
    from teamagent.media import render_child

    prs, slide = deck()
    slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "{{1}}"
    prs.save(tmp_path / "template.pptx")
    (tmp_path / "composer.json").write_text(json.dumps({"placeholders": {"1": "説明"}}))
    if qa_fails:

        def fail(*args: Any, **kwargs: Any) -> Any:
            raise RuntimeError("検査例外")

        monkeypatch.setattr("teamagent.media.deck_qa.inspect_pptx", fail)
    stats = render_child._proposal(
        tmp_path.resolve(),
        {
            "kind": "proposal_pptx",
            "template": "template.pptx",
            "composer": "composer.json",
            "evidence": [],
            "output": "output.pptx",
            "fail_if_missing": False,
        },
    )
    assert (tmp_path / "output.pptx").is_file()
    assert stats["template_slides"] == 1
    assert ("qa_error" in stats) is qa_fails
    assert ("deck_qa" in stats) is not qa_fails


@pytest.mark.parametrize("outside", [False, True])
def test_rotated_shape(tmp_path: Path, outside: bool) -> None:
    prs, slide = deck()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(2))
    shape.rotation = 45
    if outside:
        shape.left = 0
    assert ("out_of_bounds" in kinds(prs, tmp_path)) is outside


@pytest.mark.parametrize("overflow", [False, True])
def test_explicit_line_breaks(tmp_path: Path, overflow: bool) -> None:
    prs, slide = deck()
    frame = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.text = "説明\v" * (10 if overflow else 1)
    assert ("text_overflow" in kinds(prs, tmp_path)) is overflow


def test_inherited_theme_font(tmp_path: Path) -> None:
    prs, slide = deck()
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(2)).text = "説明"
    master = slide.slide_layout.slide_master
    style = master._element.xpath("./p:txStyles/p:otherStyle/a:lvl1pPr/a:defRPr")[0]
    style.set("sz", "2000")
    from pptx.oxml.xmlchemy import OxmlElement

    face = OxmlElement("a:latin")
    face.set("typeface", "Missing Inherited Font")
    for node in style.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}latin"):
        style.remove(node)
    style.append(face)
    # 日本語の本文では east Asian の既定フォントも対象にする。
    for node in style.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}ea"):
        node.set("typeface", "Missing Inherited Font")
    assert "font_unavailable" in kinds(prs, tmp_path)
    assert "Missing Inherited Font" in inspect_pptx(tmp_path / "test.pptx").fonts


@pytest.mark.parametrize("outside", [False, True])
def test_table_bounds(tmp_path: Path, outside: bool) -> None:
    prs, slide = deck()
    table = slide.shapes.add_table(
        1, 1, prs.slide_width - Inches(1), 0, Inches(2 if outside else 1), Inches(1)
    )
    table.table.cell(0, 0).text = "済"
    assert ("out_of_bounds" in kinds(prs, tmp_path)) is outside


def test_autofit_background_is_not_collision(tmp_path: Path) -> None:
    prs, slide = deck()
    slide.shapes.add_picture(
        image((400, 300)), 0, 0, width=prs.slide_width, height=prs.slide_height
    )
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.3)).text_frame
    frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    frame.word_wrap = True
    frame.text = "あ" * 40
    assert "text_autofit_collision" not in kinds(prs, tmp_path)


@pytest.mark.parametrize("outside", [False, True])
def test_chart_bounds(tmp_path: Path, outside: bool) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    prs, slide = deck()
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("系列", [1, 2])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        prs.slide_width - Inches(1),
        0,
        Inches(2 if outside else 1),
        Inches(1),
        data,
    )
    assert ("out_of_bounds" in kinds(prs, tmp_path)) is outside


@pytest.mark.parametrize("strict", [False, True])
def test_aspect_severity_and_no_duplicates(tmp_path: Path, strict: bool) -> None:
    prs, slide = deck()
    slide.shapes.add_picture(image(), 0, 0, width=Inches(2), height=Inches(2))
    path = tmp_path / "test.pptx"
    prs.save(path)
    found = [f for f in inspect_pptx(path, strict=strict).findings if f.kind == "image_aspect"]
    assert len(found) == 1
    assert found[0].severity == ("error" if strict else "warn")


def rewrite_package(path: Path, replacements: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    files.update(replacements)
    with zipfile.ZipFile(path, "w") as archive:
        for name, blob in files.items():
            archive.writestr(name, blob)


@pytest.mark.parametrize("extension", ["jpg", "jpeg"])
@pytest.mark.parametrize("content_type", ["image/jpg", "image/jpeg"])
def test_jpeg_aliases(tmp_path: Path, extension: str, content_type: str) -> None:
    prs, slide = deck()
    stream = io.BytesIO()
    Image.new("RGB", (200, 100)).save(stream, "JPEG")
    stream.seek(0)
    slide.shapes.add_picture(stream, 0, 0, width=Inches(2))
    path = tmp_path / "test.pptx"
    prs.save(path)
    with zipfile.ZipFile(path) as archive:
        types = archive.read("[Content_Types].xml").replace(b"image/jpeg", content_type.encode())
        rels = archive.read("ppt/slides/_rels/slide1.xml.rels")
        key = next(n for n in archive.namelist() if n.startswith("ppt/media/"))
        target = key.rsplit(".", 1)[0] + "." + extension
        blob = archive.read(key)
        if extension == "jpeg":
            entry = f'<Default Extension="jpg" ContentType="{content_type}"/>'.encode()
            types = types.replace(
                entry, entry + entry.replace(b'Extension="jpg"', b'Extension="jpeg"')
            )
        rels = rels.replace(key.rsplit("/", 1)[-1].encode(), target.rsplit("/", 1)[-1].encode())
    rewrite_package(
        path, {"[Content_Types].xml": types, "ppt/slides/_rels/slide1.xml.rels": rels, target: blob}
    )
    assert not [f for f in inspect_pptx(path).findings if f.kind.startswith("image_")]


@pytest.mark.parametrize("change", ["mismatch", "broken", "empty", "types", "webp", "avif"])
def test_registration_severity(tmp_path: Path, change: str) -> None:
    prs, slide = deck()
    slide.shapes.add_picture(image(), 0, 0, width=Inches(2))
    path = tmp_path / "test.pptx"
    prs.save(path)
    mutate(path, change)
    found = [f for f in inspect_pptx(path).findings if f.kind == "image_registration"]
    assert found
    assert any(f.severity == "error" for f in found) is (change in {"types", "webp", "avif"})
    assert not [f for f in inspect_pptx(path).findings if f.kind == "image_broken"]


@pytest.mark.parametrize("geometry", ["none", "picture", "layout", "master", "custom"])
def test_picture_geometry_inheritance(tmp_path: Path, geometry: str) -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[8])
    picture = slide.placeholders[1].insert_picture(image())
    layout = slide.slide_layout.placeholders[1]
    master = slide.slide_layout.slide_master.placeholders[1]
    for shape in (picture, layout, master):
        for node in shape._element.xpath("./p:spPr/a:prstGeom | ./p:spPr/a:custGeom"):
            node.getparent().remove(node)
    if geometry != "none":
        target = {"picture": picture, "layout": layout, "master": master, "custom": picture}[
            geometry
        ]
        node = OxmlElement("a:custGeom" if geometry == "custom" else "a:prstGeom")
        if geometry != "custom":
            node.set("prst", "rect")
        node.append(OxmlElement("a:avLst"))
        target._element.spPr.append(node)
    path = tmp_path / "test.pptx"
    prs.save(path)
    found = [f for f in inspect_pptx(path).findings if f.kind == "picture_no_geometry"]
    assert len(found) == (1 if geometry == "none" else 0)
    if found:
        assert found[0].slide == 1 and found[0].severity == "error"


@pytest.mark.parametrize("source", ["layout", "master"])
def test_placeholder_effective_properties(tmp_path: Path, source: str) -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    shape = slide.placeholders[1]
    layout = slide.slide_layout.placeholders[1]
    master = slide.slide_layout.slide_master.placeholders[1]
    # master の idx と layout の idx を違えて type による対応を検証する。
    master._element.xpath("./p:nvSpPr/p:nvPr/p:ph")[0].set("idx", "99")
    target = layout if source == "layout" else master
    for candidate in (shape, layout, master):
        for node in candidate._element.xpath("./p:txBody/a:p/a:pPr"):
            node.getparent().remove(node)
    for candidate in (shape, layout):
        for node in candidate._element.xpath("./p:spPr/a:xfrm"):
            node.getparent().remove(node)
        body = candidate.text_frame._txBody.bodyPr
        body.attrib.clear()
        for node in list(body):
            body.remove(node)
        for node in list(candidate.text_frame._txBody.xpath("./a:lstStyle")[0]):
            node.getparent().remove(node)
    target.left, target.top = Inches(1), Inches(1)
    target.width, target.height = Inches(4), Inches(1)
    body = target.text_frame._txBody.bodyPr
    body.attrib.update({"wrap": "none", "lIns": "0", "rIns": "0", "tIns": "0", "bIns": "0"})
    for child in list(body):
        body.remove(child)
    body.append(OxmlElement("a:normAutofit"))
    props = OxmlElement("a:lvl1pPr")
    run = OxmlElement("a:defRPr")
    run.set("sz", "1000")
    props.append(run)
    target.text_frame._txBody.xpath("./a:lstStyle")[0].append(props)
    shape.text = "あ" * 40
    path = tmp_path / "test.pptx"
    prs.save(path)
    assert not [
        f
        for f in inspect_pptx(path).findings
        if f.shape == shape.name and f.kind in {"text_overflow", "out_of_bounds"}
    ]
    # 同じ継承 autoFit を残して長くしても error にはしない。
    shape.text = "あ\v" * 40
    prs.save(path)
    found = [
        f
        for f in inspect_pptx(path).findings
        if f.shape == shape.name and f.kind == "text_overflow"
    ]
    assert len(found) == 1 and found[0].severity == "warn"


@pytest.mark.parametrize("auto", [MSO_AUTO_SIZE.NONE, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE])
def test_modest_text_overflow_is_warning(tmp_path: Path, auto: Any) -> None:
    prs, slide = deck()
    frame = slide.shapes.add_textbox(0, 0, Inches(4), Inches(0.4)).text_frame
    frame.auto_size = auto
    frame.text = "説明"
    frame.paragraphs[0].runs[0].font.size = Pt(22)
    path = tmp_path / "test.pptx"
    prs.save(path)
    found = [f for f in inspect_pptx(path).findings if f.kind == "text_overflow"]
    assert len(found) == 1 and found[0].severity == "warn"


def test_oversized_insets_do_not_prove_text_clipping(tmp_path: Path) -> None:
    prs, slide = deck()
    frame = slide.shapes.add_textbox(0, 0, Inches(4), Inches(0.1)).text_frame
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_top = frame.margin_bottom = Inches(0.3)
    frame.text = "説明"
    path = tmp_path / "test.pptx"
    prs.save(path)
    found = [f for f in inspect_pptx(path).findings if f.kind == "text_overflow"]
    assert len(found) == 1 and found[0].severity == "warn"


@pytest.mark.parametrize("kind", ["svg", "emf", "linked", "no_blip"])
def test_non_bitmap_pictures_are_not_broken(tmp_path: Path, kind: str) -> None:
    prs, slide = deck()
    picture = slide.shapes.add_picture(image(), 0, 0, width=Inches(2))
    blip = picture._element.xpath("./p:blipFill/a:blip")[0]
    if kind == "linked":
        rid = blip.attrib.pop(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
        blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link", rid)
    elif kind == "no_blip":
        blip.getparent().remove(blip)
    path = tmp_path / "test.pptx"
    prs.save(path)
    if kind in {"svg", "emf"}:
        with zipfile.ZipFile(path) as archive:
            types = archive.read("[Content_Types].xml").replace(
                b"image/png", b"image/svg+xml" if kind == "svg" else b"image/x-emf"
            )
            media = next(n for n in archive.namelist() if n.startswith("ppt/media/"))
        rewrite_package(
            path,
            {
                "[Content_Types].xml": types,
                media: b'<svg xmlns="http://www.w3.org/2000/svg"/>' if kind == "svg" else b"EMF",
            },
        )
    assert not [f for f in inspect_pptx(path).findings if f.kind == "image_broken"]


def test_repeated_aspect_finding_is_coalesced(tmp_path: Path) -> None:
    prs, slide = deck()
    for x in (0, 3, 6):
        picture = slide.shapes.add_picture(image(), Inches(x), 0, width=Inches(2), height=Inches(2))
        picture.name = "repeated image"
    path = tmp_path / "test.pptx"
    prs.save(path)
    assert len([f for f in inspect_pptx(path).findings if f.kind == "image_aspect"]) == 1


def test_missing_picture_relationship(tmp_path: Path) -> None:
    from xml.etree import ElementTree as ET

    prs, slide = deck()
    slide.shapes.add_picture(image(), 0, 0, width=Inches(2))
    path = tmp_path / "test.pptx"
    prs.save(path)
    key = "ppt/slides/_rels/slide1.xml.rels"
    with zipfile.ZipFile(path) as archive:
        tree = ET.fromstring(archive.read(key))
    for rel in list(tree):
        if rel.get("Type", "").endswith("/image"):
            tree.remove(rel)
    rewrite_package(path, {key: ET.tostring(tree)})
    found = [f for f in inspect_pptx(path).findings if f.kind == "image_missing"]
    assert len(found) == 1 and found[0].severity == "error"


def test_redaction_marker_is_warning(tmp_path: Path) -> None:
    prs, slide = deck()
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "○○"
    path = tmp_path / "test.pptx"
    prs.save(path)
    found = [f for f in inspect_pptx(path).findings if f.kind == "placeholder"]
    assert len(found) == 1 and found[0].severity == "warn"


def test_explicit_small_runs_override_large_default(tmp_path: Path) -> None:
    prs, slide = deck()
    frame = slide.shapes.add_textbox(0, 0, Inches(4), Inches(0.6)).text_frame
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.text = "説明"
    frame.paragraphs[0].font.size = Pt(60)
    frame.paragraphs[0].runs[0].font.size = Pt(10)
    assert "text_overflow" not in kinds(prs, tmp_path)


def test_inherited_paragraph_level_size(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    shape = slide.placeholders[1]
    shape.height = Inches(0.4)
    shape.text = "説明"
    shape.text_frame.paragraphs[0].level = 1
    shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
    master = slide.slide_layout.slide_master
    style = master._element.xpath("./p:txStyles/p:bodyStyle/a:lvl2pPr/a:defRPr")[0]
    style.set("sz", "1000")
    path = tmp_path / "test.pptx"
    prs.save(path)
    assert not [
        f
        for f in inspect_pptx(path).findings
        if f.shape == shape.name and f.kind == "text_overflow"
    ]


@pytest.mark.parametrize("registered", [True, False])
def test_xml_comments_do_not_hide_registration(tmp_path: Path, registered: bool) -> None:
    # lxml はコメントも子として返す。コメント入りの部品でも落ちず、登録漏れを見逃さない。
    prs, slide = deck()
    slide.shapes.add_picture(image(), 0, 0, width=Inches(2))
    path = tmp_path / "test.pptx"
    prs.save(path)
    with zipfile.ZipFile(path) as archive:
        types = archive.read("[Content_Types].xml")
        rels = archive.read("ppt/slides/_rels/slide1.xml.rels")
    if not registered:
        types = types.replace(b'<Default Extension="png" ContentType="image/png"/>', b"")
    types = types.replace(b"<Default ", b"<!-- note --><Default ", 1)
    rels = rels.replace(b"<Relationship ", b"<!-- note --><Relationship ", 1)
    rewrite_package(path, {"[Content_Types].xml": types, "ppt/slides/_rels/slide1.xml.rels": rels})
    found = [f for f in inspect_pptx(path).findings if f.kind == "image_registration"]
    assert bool(found) is not registered


class _Summary:
    def __init__(self, payload: dict[str, Any]) -> None:
        self.payload = payload

    def summary(self) -> dict[str, Any]:
        return self.payload


@pytest.mark.parametrize(
    "case",
    ["nan", "huge_details", "huge_exception"],
)
def test_qa_metadata_never_breaks_the_media_result(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, case: str
) -> None:
    # 検査の結果が 32KB 超・NaN だと MediaJobResult の検証（_canonical_json・allow_nan=False・
    # 32KB）で依頼全体が失敗し、資料が届かない。
    from teamagent.media.contracts import _canonical_json

    prs, _ = deck()
    path = tmp_path / "test.pptx"
    prs.save(path)
    finding = {"slide": 1, "shape": "s", "kind": "text_overflow", "severity": "warn"}
    if case == "nan":
        payload = {
            "error_count": 0,
            "warn_count": 1,
            "top_findings": [{**finding, "details": {"x": float("nan")}}],
        }
    else:
        payload = {
            "error_count": 1,
            "warn_count": 0,
            "top_findings": [{**finding, "details": {"collisions": ["x" * 40] * 2000}}],
        }

    def fake(*args: Any, **kwargs: Any) -> Any:
        if case == "huge_exception":
            raise ValueError("x" * 100_000)
        return _Summary(payload)

    monkeypatch.setattr("teamagent.media.deck_qa.inspect_pptx", fake)
    metadata = {"slides": 1, **_deck_qa(path, expected_slides=1)}
    # 子プロセスの出力と同じ往復を通す（NaN はここで JSON に載る）。
    metadata = json.loads(json.dumps(metadata))
    assert len(_canonical_json(metadata)) <= 32 * 1024
    if case != "huge_exception":
        assert metadata["deck_qa"]["top_findings_dropped"] is True
        assert metadata["deck_qa"]["error_count"] == payload["error_count"]


@pytest.mark.parametrize("placement", ["fits", "outside", "collision"])
def test_table_grows_as_whole(tmp_path: Path, placement: str) -> None:
    prs, slide = deck()
    top = prs.slide_height - Inches(0.5) if placement == "outside" else Inches(1)
    shape = slide.shapes.add_table(2, 2, Inches(1), top, Inches(4), Inches(0.4))
    for row in shape.table.rows:
        for cell in row.cells:
            cell.text = "説明\v" * 3 + "説明"
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    if placement == "collision":
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(4), Inches(1))
    path = tmp_path / "table.pptx"
    prs.save(path)
    findings = inspect_pptx(path).findings
    assert not [f for f in findings if f.kind == "text_overflow"]
    found = [f for f in findings if f.kind == "table_overflow"]
    assert len(found) == (0 if placement == "fits" else 1)
    if found:
        assert found[0].severity == ("error" if placement == "outside" else "warn")
        assert bool(found[0].details["collisions"]) is (placement == "collision")


@pytest.mark.parametrize("merged", [False, True])
@pytest.mark.parametrize("wide", [False, True])
def test_table_grid_width_and_span(tmp_path: Path, merged: bool, wide: bool) -> None:
    prs, slide = deck()
    shape = slide.shapes.add_table(1, 3, 0, prs.slide_height - Inches(0.8), Inches(6), Inches(0.2))
    table = shape.table
    table.columns[0].width = Inches(3 if wide else 0.5)
    table.columns[1].width = Inches(2 if wide else 0.5)
    table.columns[2].width = Inches(1 if wide else 5)
    cell = table.cell(0, 0)
    if merged:
        cell.merge(table.cell(0, 1))
    cell.text = "あ" * (45 if merged else 25)
    cell.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    assert ("table_overflow" in kinds(prs, tmp_path)) is (not wide)


@pytest.mark.parametrize("source", ["layout", "master"])
@pytest.mark.parametrize("overflow", [False, True])
@pytest.mark.parametrize("property_name", ["size", "insets"])
def test_inherited_text_capacity(
    tmp_path: Path, source: str, overflow: bool, property_name: str
) -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    shape = slide.placeholders[1]
    layout = slide.slide_layout.placeholders[1]
    master = slide.slide_layout.slide_master.placeholders[1]
    target = layout if source == "layout" else master
    shape.width, shape.height = Inches(2), Inches(0.5)
    for candidate in (shape, layout, master):
        body = candidate.text_frame._txBody.bodyPr
        body.attrib.clear()
        for child in list(body):
            body.remove(child)
        for child in candidate._element.xpath("./p:txBody/a:p/a:pPr | ./p:txBody/a:lstStyle/*"):
            child.getparent().remove(child)
    target.text_frame._txBody.bodyPr.attrib.update(
        {"lIns": "0", "rIns": "0", "tIns": "0", "bIns": "0"}
    )
    props = OxmlElement("a:lvl1pPr")
    run = OxmlElement("a:defRPr")
    run.set("sz", "3000" if overflow and property_name == "size" else "1000")
    props.append(run)
    target.text_frame._txBody.xpath("./a:lstStyle")[0].append(props)
    if property_name == "insets" and overflow:
        target.text_frame._txBody.bodyPr.set("tIns", str(Inches(0.5)))
    shape.text = "あ" * 14
    path = tmp_path / "inherited.pptx"
    prs.save(path)
    found = [
        f
        for f in inspect_pptx(path).findings
        if f.shape == shape.name and f.kind == "text_overflow"
    ]
    assert bool(found) is overflow


@pytest.mark.parametrize("include_edges", [False, True])
def test_explicit_percentage_spacing_and_paragraph_edges(
    tmp_path: Path, include_edges: bool
) -> None:
    prs, slide = deck()
    # 行間 100%（spcPct）は指定なしと同じ 1 行＝字の 1.2 倍（20pt → 24pt）。0.36in（25.9pt）に収まる。
    frame = slide.shapes.add_textbox(0, 0, Inches(2), Inches(0.36)).text_frame
    frame.margin_top = frame.margin_bottom = 0
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.text = "説明"
    p = frame.paragraphs[0]
    p.font.size = Pt(20)
    p.line_spacing = 1.0
    p.space_before = p.space_after = Pt(20)
    frame._txBody.bodyPr.set("spcFirstLastPara", "1" if include_edges else "0")
    assert ("text_overflow" in kinds(prs, tmp_path)) is include_edges


@pytest.mark.parametrize("large_margin", [False, True])
def test_table_cell_margins(tmp_path: Path, large_margin: bool) -> None:
    prs, slide = deck()
    shape = slide.shapes.add_table(1, 1, 0, prs.slide_height - Inches(0.5), Inches(3), Inches(0.2))
    cell = shape.table.cell(0, 0)
    cell.text = "説明"
    cell.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    cell.margin_top = Inches(0.7) if large_margin else 0
    cell.margin_bottom = 0
    assert ("table_overflow" in kinds(prs, tmp_path)) is large_margin


@pytest.mark.parametrize("long", [False, True])
def test_table_vertical_merge_capacity(tmp_path: Path, long: bool) -> None:
    prs, slide = deck()
    shape = slide.shapes.add_table(2, 1, 0, prs.slide_height - Inches(1), Inches(3), Inches(0.8))
    cell = shape.table.cell(0, 0)
    cell.merge(shape.table.cell(1, 0))
    cell.text = "説明\v" * (12 if long else 2) + "説明"
    for run in cell.text_frame.paragraphs[0].runs:
        run.font.size = Pt(10)
    assert ("table_overflow" in kinds(prs, tmp_path)) is long


@pytest.mark.parametrize("overflow", [False, True])
def test_footer_uses_other_master_style(tmp_path: Path, overflow: bool) -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    prs, slide = deck()
    shape = slide.shapes.add_textbox(0, 0, Inches(3), Inches(0.3))
    ph = OxmlElement("p:ph")
    ph.set("type", "ftr")
    ph.set("idx", "42")
    shape._element.xpath("./p:nvSpPr/p:nvPr")[0].append(ph)
    shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
    shape.text = "説明"
    master = slide.slide_layout.slide_master
    for candidate in master.placeholders:
        if int(candidate.placeholder_format.type) == 15:
            for child in candidate._element.xpath("./p:txBody/a:lstStyle/* | ./p:txBody/a:p/a:pPr"):
                child.getparent().remove(child)
    style = master._element.xpath("./p:txStyles/p:otherStyle/a:lvl1pPr/a:defRPr")[0]
    style.set("sz", "3000" if overflow else "1000")
    master._element.xpath("./p:txStyles/p:bodyStyle/a:lvl1pPr/a:defRPr")[0].set("sz", "4000")
    assert ("text_overflow" in kinds(prs, tmp_path)) is overflow


@pytest.mark.parametrize("explicit", [False, True])
def test_explicit_single_spacing_equals_default(tmp_path: Path, explicit: bool) -> None:
    # PowerPoint では行間の指定なしと「100%」は同じ見た目。判定も同じでなければならない。
    prs, slide = deck()
    frame = slide.shapes.add_textbox(0, 0, Inches(2), Inches(0.3)).text_frame
    frame.margin_top = frame.margin_bottom = 0
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.text = "説明"
    p = frame.paragraphs[0]
    p.font.size = Pt(20)
    if explicit:
        p.line_spacing = 1.0
    # 20pt × 1.2 = 24pt は 0.3in（21.6pt）× 1.1 を超える。
    assert "text_overflow" in kinds(prs, tmp_path)


@pytest.mark.parametrize("break_size", [None, 20])
def test_line_break_size_follows_its_own_or_previous_run(
    tmp_path: Path, break_size: int | None
) -> None:
    # python-pptx の改行（a:br）は rPr を持たない。段落の既定（18pt）で数えると 10pt の 3 行が
    # 高く見積もられ、収まる枠を「あふれ」と誤る。改行に字の大きさがあればそれを使う。
    prs, slide = deck()
    frame = slide.shapes.add_textbox(0, 0, Inches(3), Pt(10 * 1.2 * 3)).text_frame
    frame.margin_top = frame.margin_bottom = 0
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.text = "説明\v説明\v説明"
    for run in frame.paragraphs[0].runs:
        run.font.size = Pt(10)
    if break_size is not None:
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for br in frame.paragraphs[0]._p.findall(f"{ns}br"):
            props = br.makeelement(f"{ns}rPr", {"sz": str(break_size * 100)})
            br.insert(0, props)
    assert ("text_overflow" in kinds(prs, tmp_path)) is (break_size is not None)
