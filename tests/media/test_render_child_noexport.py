"""media worker の撮影（render_child._slides）が、撮る前に [data-noexport] を隠すこと（T21）。

編集ヒント（position:fixed）を隠さないと、PPTX の全スライドの左上に焼き込まれる（D-01・本番と
同じ条件で 7 枚すべてに写ることを再現済み）。pptx_export.shoot_sections と同じ CSS を入れる。
"""

from __future__ import annotations

from pathlib import Path

import pytest

from teamagent.media import render_child
from tests.skills.video_algorithm.chromium import (
    FakePlaywright,
    assert_hidden_before_shots,
    chromium_path,
    edit_tip_brightness,
    launched_browser,
)


def _slides_html() -> str:
    """本物のスライド HTML（編集ヒント＝position:fixed の data-noexport を含む・外部参照なし）。"""
    from teamagent.skills.video_algorithm.schema import VideoAlgorithmOutput
    from teamagent.skills.video_algorithm.slides import render_slides

    html = render_slides(VideoAlgorithmOutput(query="新宿 ランチ"))
    # 2 枚にする（表紙を複製）。2 枚目以降にも焼き込まれていたことを確かめるため。
    first = html.index('<section class="slide')
    end = html.index("</section>", first) + len("</section>")
    return html[:end] + html[first:end] + html[end:]


def _manifest(root: Path) -> dict[str, object]:
    (root / "slides.html").write_text(_slides_html(), encoding="utf-8")
    return {
        "kind": "slides",
        "html": "slides.html",
        "output": "slides.pptx",
        "selector": ".slide",
        "width": 1280,
        "height": 720,
        "scale": 2,
    }


def test_render_child_hides_noexport_before_screenshots(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """壊し方: render_child._slides の add_style_tag を外す → 赤。"""
    import playwright.sync_api

    calls: list[tuple[str, object]] = []
    monkeypatch.setattr(playwright.sync_api, "sync_playwright", FakePlaywright(calls))
    meta = render_child._slides(tmp_path.resolve(), _manifest(tmp_path))
    assert meta == {"slides": 2, "network_requests_allowed": 0}
    assert_hidden_before_shots(calls)
    assert render_child._NOEXPORT_CSS == "[data-noexport]{display:none!important}"


def test_render_child_real_browser_pptx_has_no_edit_tip(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """実描画（chromium が無ければ skip）: PPTX の画像の左上が白い。"""
    from pptx import Presentation

    with launched_browser():
        pass
    exe = chromium_path()
    if exe is None:
        pytest.skip("CHROMIUM_PATH が無い")
    monkeypatch.setenv("CHROMIUM_PATH", exe)
    root = tmp_path.resolve()
    meta = render_child._slides(root, _manifest(root))
    assert meta["slides"] == 2
    prs = Presentation(str(root / "slides.pptx"))
    blobs = [sh.image.blob for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13]
    assert len(blobs) == 2
    for blob in blobs:
        assert edit_tip_brightness(blob) > 245
