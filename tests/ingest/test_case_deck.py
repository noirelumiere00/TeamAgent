"""事例集 PPTX を事例単位に切って取り込む（case_corpus_decks・2026-10-06）。

社名・本文はすべて **架空**（実 deck の顧客名はフィクスチャにもコミットにも入れない）。
ただし型は実 deck（Drive コネクタで全文確認・11 事例）と同じにしてある:
  - 区切りスライド（ブランド名 1 語＋共通フッタ）→ 表題スライド（①②③＋全体結果ハイライト）
  - 続きスライドに表題「◯◯ 様」を繰り返す事例（実 deck では 4 枚繰り返しが 2 事例）
  - 区切りの無い事例・見出し（①）の無い事例・本文中の「◯◯様からは…」
  - 表題 shape より前に本文 shape が来るスライド（shape の並び順は保証されない）

フェイクは本番の失敗モードを再現する:
  - 抽出 0 件（印の無い deck）/ 上限超過 / 破損 / 表題の型崩れ / 同名の会社が 2 事例 /
    ACL 空（permissions が取れない）/ ファイルがフォルダに無い
変異（どれも該当テストが赤になることを手で確認済み・PR 本文に記載）:
  - case_corpus 付与を消す → test_pipeline_writes_one_document_per_case
  - 区切りの正規表現（全体結果ハイライト）を壊す → 件数系のテスト
  - 表題で切る（続きスライドの「◯◯ 様」でも切る）→ test_repeated_title_slides_stay_in_one_case
"""

from __future__ import annotations

import hashlib
from io import BytesIO
from typing import Any

import pytest

from teamagent.adapters.gdrive_client import DriveFile, DrivePermission
from teamagent.ingest.case_deck import (
    assign_external_ids,
    case_external_id,
    format_case_pages,
    parse_case_title,
    split_case_deck,
)
from teamagent.ingest.loader import CaseDeckSpec
from teamagent.ingest.office_extract import (
    OfficePayloadError,
    extract_pptx_slide_shapes,
)

FOOTER = "Kakuu Garden 18th Floor 1-2-3 Kakuu, Minato-ku, Tokyo 100-0000 Japan"
TABLE = "媒体 | 投稿本数 | 再生数 | オーガニック割合\nTikTok | 27本 | 3,791,739回 | 55.73%"
MARKER = "本施策の全体結果ハイライト"

# 実 deck と同じ型の 16 枚（1 枚目は表紙）。
DECK_SHAPES: list[tuple[str, ...]] = [
    ("2026 ショート動画事例集", "社外秘"),  # 1 表紙（どの事例にも入らない）
    ("北斗", FOOTER),  # 2 区切り
    (  # 3 事例1の始まり
        FOOTER,
        "青葉レコード様：北斗シスターズ様",
        "2",
        "① UGC風動画で話題化",
        "380%達成",
        "投稿本数27本で約380万回再生。舞台裏動画で話題化。",
        "② 過去半年で№1",
        "最も高い\n指名検索を獲得",
        "文脈設計を起点に、指名検索数を獲得。",
        "③ メディア露出の獲得",
        "経済紙への掲載",
        "本施策の成果により、経済紙に掲載。",
        MARKER,
        "北斗シスターズの新曲 TikTokショート動画施策",
        "本施策は、新曲の発売に合わせて実施。目標対比380％の再生回数を獲得した。",
    ),
    (FOOTER, "施策結果：再生数", "3", TABLE),  # 4
    ("本施策の構造設計", "4", "フェーズごとに訴求を変えた。"),  # 5
    ("羊の箱", FOOTER),  # 6 区切り
    (  # 7 事例2の始まり（見出しと成果が 1 shape・印と小見出しが 1 shape）
        FOOTER,
        "銀河フィルム株式会社 様「羊の箱」",
        "7",
        "① TikTok上で大規模な認知獲得\n約375万回再生を獲得",
        "投稿本数47本で、総再生数3,748,533回を記録。",
        "② 勝ちクリエイティブの創出",
        "単一の動画が\n260万回再生を記録",
        "最も再生された動画は権威性訴求の構成。",
        f"{MARKER}\n映画「羊の箱」 TikTokショート動画施策",
        "本施策は、映画の公開に合わせて実施。総再生数約374万回を獲得した。",
    ),
    ("施策結果：TikTok内での占有状況", "8"),  # 8 短いが区切りではない
    ("みなと銀行", FOOTER),  # 9 区切り
    (  # 10 事例3の始まり（ブランドなし）
        FOOTER,
        "みなと銀行 様",
        "22",
        "① 目標再生数を超過達成",
        "119%達成",
        "・133本投稿で総再生118.7万回を獲得",
        MARKER,
        "みなと銀行 ショート動画施策（3月・4月）",
        "総再生119万回で目標を119%達成した。",
    ),
    (FOOTER, "みなと銀行 様", "23", "短期間・多本数投稿により、目標再生数を超過達成。"),  # 11
    (FOOTER, "みなと銀行 様", "24", "投稿開始後、再生数と連動して指名検索が上昇。"),  # 12
    ("白樺スリープ", FOOTER),  # 13 区切り
    (  # 14 事例4の始まり（本文の「◯◯様からは」が表題より前にある）
        FOOTER,
        "依頼先の東西グループ様からは、その後も継続でご発注いただいた。",
        "白樺様 白樺スリープ「深呼吸まくら」",
        "17",
        "① 1千万円規模の売上創出",
        "751件のCVを獲得",
        f"{MARKER}（4つの成功）",
        "深呼吸まくら ショート動画施策",
        "高価格帯商材でありながら751件のCVを獲得。",
    ),
    (  # 15 事例5（区切りなし・注記つき表題）
        "North Field Japan株式会社様（整備済み端末の再販事業）",
        "32",
        "① 認知拡大×来場促進",
        "ポップアップ連動\n145%達成",
        MARKER,
        "「再生品」認知拡大 ポップアップ施策",
        "目標95万回に対し138.1万回再生。",
    ),
    (  # 16 事例6（見出しなし＝概要文の数字入りの文に落ちる）
        "株式会社みらい製菓様",
        "43",
        MARKER,
        "グミ 新商品発売×ショート動画 成功事例",
        "計100本の動画を集中投下。KPI比136.0%（約204万再生）を達成した。日常の悩みをフックにした。",
    ),
]


def _slides(shapes: list[tuple[str, ...]] | None = None) -> list[tuple[int, tuple[str, ...]]]:
    return [(i, s) for i, s in enumerate(shapes or DECK_SHAPES, start=1)]


# ── 表題の型 ───────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("青葉レコード様：北斗シスターズ様", ("青葉レコード", "北斗シスターズ")),
        ("みなと食品様：フルーツ酢", ("みなと食品", "フルーツ酢")),
        ("青葉ストア様：スポンジくん様", ("青葉ストア", "スポンジくん")),
        ("銀河フィルム株式会社 様「羊の箱」", ("銀河フィルム株式会社", "羊の箱")),
        ("SEVEN STAR 様『星の身分』", ("SEVEN STAR", "星の身分")),
        ("白樺様 白樺スリープ「深呼吸まくら」", ("白樺", "白樺スリープ「深呼吸まくら」")),
        ("みなと銀行 様", ("みなと銀行", "")),
        ("Skyline 様", ("Skyline", "")),
        (
            "North Field Japan株式会社様（整備済み端末の再販事業）",
            ("North Field Japan株式会社", ""),
        ),
        ("SEA LINE振興会様", ("SEA LINE振興会", "")),
        ("株式会社みらい製菓様", ("株式会社みらい製菓", "")),
        # 段落・run に割れた表題（shape 内で改行）も 1 本として読む
        ("青葉レコード様：\n北斗シスターズ様", ("青葉レコード", "北斗シスターズ")),
    ],
)
def test_title_variants_seen_in_the_real_deck(title: str, expected: tuple[str, str]) -> None:
    assert parse_case_title(title) == expected


@pytest.mark.parametrize(
    "text",
    [
        "依頼先の東西グループ様からは、その後も継続でご発注いただいた。",
        "① 目標再生数を超過達成",
        "本施策の全体結果ハイライト",
        "施策結果：再生数",
        "",
        "様",
    ],
)
def test_non_title_shapes_are_not_titles(text: str) -> None:
    assert parse_case_title(text) is None


# ── 切り出し ───────────────────────────────────────────────────────────


def test_split_finds_every_case_with_real_slide_ranges() -> None:
    result = split_case_deck(_slides())
    assert result.slide_count == 16
    assert result.start_slides == 6
    assert result.unmatched_titles == 0
    got = [(c.client_name, c.brand_key, c.slide_from, c.slide_to) for c in result.cases]
    assert got == [
        ("青葉レコード", "北斗シスターズ", 2, 5),
        ("銀河フィルム", "羊の箱", 6, 8),
        ("みなと銀行", "", 9, 12),
        ("白樺", "白樺スリープ", 13, 14),
        ("North Field Japan", "", 15, 15),
        ("みらい製菓", "", 16, 16),
    ]


def test_repeated_title_slides_stay_in_one_case() -> None:
    """続きスライドに「みなと銀行 様」を繰り返しても 1 事例のまま（表題では切らない）。

    変異: 始まりの判定を「表題が読めるスライド」にすると 3 件に割れて赤。
    """
    cases = split_case_deck(_slides()).cases
    minato = [c for c in cases if c.client_name == "みなと銀行"]
    assert len(minato) == 1
    assert [num for num, _ in minato[0].pages] == [9, 10, 11, 12]


def test_divider_is_attached_to_its_own_case_not_the_previous_one() -> None:
    """区切り（「羊の箱」）は次の事例に入り、短い結果スライドは区切りと誤認しない。"""
    cases = split_case_deck(_slides()).cases
    first, second = cases[0], cases[1]
    assert all("羊の箱" not in text for _, text in first.pages)
    assert second.pages[0] == (6, "羊の箱")
    assert second.slide_to == 8  # 「施策結果：TikTok内での占有状況」は事例2の最後
    assert all(num != 1 for c in cases for num, _ in c.pages)  # 表紙はどこにも入らない


def test_footer_and_page_numbers_are_removed_from_bodies() -> None:
    cases = split_case_deck(_slides()).cases
    for case in cases:
        for _, text in case.pages:
            assert "Kakuu Garden" not in text
            assert not any(line.strip().isdigit() for line in text.split("\n"))
    assert TABLE in cases[0].pages[2][1]  # 表は残す


def test_body_honorific_before_the_title_is_not_taken_as_title() -> None:
    """表題 shape より前に「依頼先の東西グループ様からは…」があっても表題を取り違えない。"""
    case = split_case_deck(_slides()).cases[3]
    assert (case.company, case.brand) == ("白樺", "白樺スリープ「深呼吸まくら」")
    assert case.product == "白樺スリープ「深呼吸まくら」"


def test_effect_is_one_line_built_from_numbered_highlights() -> None:
    cases = split_case_deck(_slides()).cases
    assert cases[0].effect == (
        "UGC風動画で話題化：380%達成／過去半年で№1：最も高い指名検索を獲得／"
        "メディア露出の獲得：経済紙への掲載"
    )
    assert cases[1].effect == (
        "TikTok上で大規模な認知獲得：約375万回再生を獲得／"
        "勝ちクリエイティブの創出：単一の動画が260万回再生を記録"
    )
    assert cases[4].effect == "認知拡大×来場促進：ポップアップ連動145%達成"
    # 見出しが無い型は概要文の「数字を含む文」だけ
    assert cases[5].effect == "計100本の動画を集中投下。KPI比136.0%（約204万再生）を達成した。"
    assert all(len(c.effect) <= 160 for c in cases)


def test_effect_is_capped_at_160_chars() -> None:
    long_heading = "① " + "あ" * 200
    shapes = [("青葉レコード様：北斗シスターズ様", long_heading, MARKER)]
    case = split_case_deck(_slides(shapes)).cases[0]
    assert len(case.effect) == 160


def test_product_falls_back_to_subheading_without_brand() -> None:
    cases = split_case_deck(_slides()).cases
    assert cases[2].product == "みなと銀行 ショート動画施策（3月・4月）"
    assert cases[4].product == "「再生品」認知拡大 ポップアップ施策"
    assert cases[1].product == "羊の箱"  # ブランドがあればブランド


def test_broken_title_drops_the_case_and_is_counted_not_guessed() -> None:
    """表題の型崩れ: 区切り（「羊の箱」）から社名を推測して埋めない。"""
    shapes = list(DECK_SHAPES)
    start = list(shapes[6])
    start[1] = "銀河フィルムの事例"  # 「様」が無い
    shapes[6] = tuple(start)
    result = split_case_deck(_slides(shapes))
    assert result.unmatched_titles == 1
    assert len(result.cases) == 5
    assert "銀河フィルム" not in [c.client_name for c in result.cases]


def test_deck_without_start_marker_yields_zero_cases() -> None:
    """印を壊した deck（テンプレ差し替え等）は 0 件。表題だけでは切らない。"""
    shapes = [tuple(s.replace("全体結果ハイライト", "まとめ") for s in sl) for sl in DECK_SHAPES]
    result = split_case_deck(_slides(shapes))
    assert result.cases == ()
    assert result.start_slides == 0


# ── external_id ────────────────────────────────────────────────────────


def test_external_id_is_stable_when_cases_are_reordered() -> None:
    cases = list(split_case_deck(_slides()).cases)
    ids, _ = assign_external_ids("FILE", cases)
    ids_rev, _ = assign_external_ids("FILE", list(reversed(cases)))
    assert sorted(ids) == sorted(ids_rev)
    assert all(i.startswith("FILE:case:") for i in ids)
    assert len(set(ids)) == len(ids)


def test_same_company_twice_gets_distinct_ids_and_is_counted() -> None:
    """同名の会社が 2 事例（ブランドなし同士）でも 1 文書に潰れない。"""
    shapes = [
        *DECK_SHAPES,
        ("みなと銀行 様", "① 秋の施策", "120%達成", MARKER, "みなと銀行 秋の施策"),
    ]
    cases = split_case_deck(_slides(shapes)).cases
    ids, duplicates = assign_external_ids("FILE", cases)
    assert duplicates == 1
    assert len(set(ids)) == len(ids) == 7
    assert ids[-1] == case_external_id("FILE", "みなと銀行", "", 2)


def test_external_id_ignores_corporate_form_and_width() -> None:
    a = case_external_id("F", "みらい製菓", "")
    assert a == case_external_id("F", "みらい製菓 ", "")
    assert a != case_external_id("F", "みらい製菓", "グミ")


def test_case_pages_carry_case_name_in_every_chunk_source() -> None:
    case = split_case_deck(_slides()).cases[2]
    pages = format_case_pages(case, deck_name="事例集", external_use_note="クライアント展開NG")
    assert pages[0][1].startswith("事例: みなと銀行\n出典: 事例集 スライド 9〜12\n")
    assert "対外利用: クライアント展開NG" in pages[0][1]
    assert all("みなと銀行の事例（スライド" in text for _, text in pages)


# ── 抽出（bounded OOXML parser）────────────────────────────────────────


def _build_pptx(slides: list[tuple[str, ...]], *, table_on: int | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[6]  # Blank
    for idx, shapes in enumerate(slides, start=1):
        slide = prs.slides.add_slide(layout)
        for n, text in enumerate(shapes):
            if table_on == idx and text == TABLE:
                rows = text.split("\n")
                cells = [r.split(" | ") for r in rows]
                frame = slide.shapes.add_table(
                    len(cells), len(cells[0]), Inches(1), Inches(4), Inches(6), Inches(1)
                )
                for r, row in enumerate(cells):
                    for c, value in enumerate(row):
                        frame.table.cell(r, c).text = value
                continue
            box = slide.shapes.add_textbox(
                Inches(0.2), Inches(0.1 + n * 0.3), Inches(8), Inches(0.3)
            )
            lines = text.split("\n")
            box.text_frame.text = lines[0]
            for line in lines[1:]:
                box.text_frame.add_paragraph().text = line
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_extract_keeps_shape_boundaries_runs_and_empty_slides() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(layout)
    box = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    para = box.text_frame.paragraphs[0]
    para.add_run().text = "青葉レ"
    para.add_run().text = "コード様：北斗シスターズ様"
    group = s1.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text_frame.text = "群の中"
    prs.slides.add_slide(layout)  # 空スライド
    buf = BytesIO()
    prs.save(buf)
    data = buf.getvalue()

    slides = extract_pptx_slide_shapes(data, expected_size=len(data))
    assert slides == [(1, ("青葉レコード様：北斗シスターズ様", "群の中")), (2, ())]


def test_end_to_end_pptx_to_cases() -> None:
    """python-pptx で組んだ実ファイル → bounded parser → 切り出し（表も 1 shape）。"""
    data = _build_pptx(DECK_SHAPES, table_on=4)
    slides = extract_pptx_slide_shapes(
        data,
        expected_size=len(data),
        expected_md5=hashlib.md5(data, usedforsecurity=False).hexdigest(),
    )
    assert len(slides) == 16
    assert TABLE in slides[3][1]
    result = split_case_deck(slides)
    assert [c.client_name for c in result.cases] == [
        "青葉レコード",
        "銀河フィルム",
        "みなと銀行",
        "白樺",
        "North Field Japan",
        "みらい製菓",
    ]


def test_extract_over_text_limit_is_classified_not_empty() -> None:
    data = _build_pptx(DECK_SHAPES)
    with pytest.raises(OfficePayloadError) as raised:
        extract_pptx_slide_shapes(data, max_chars=50)
    assert raised.value.category == "unsafe_content_volume"


def test_extract_over_compressed_limit_is_classified(monkeypatch: pytest.MonkeyPatch) -> None:
    import teamagent.ingest.office_extract as office_extract

    data = _build_pptx(DECK_SHAPES[:3])
    monkeypatch.setattr(office_extract, "MAX_OFFICE_COMPRESSED_BYTES", len(data) - 1)
    with pytest.raises(OfficePayloadError) as raised:
        extract_pptx_slide_shapes(data)
    assert raised.value.category == "unsafe_archive"


def test_real_deck_size_fits_the_office_limits() -> None:
    """実 deck は 44,855,013 bytes（Drive metadata・2026-10-06）。上限に収まることを固定。

    圧縮後サイズの上限（256MB）と Drive download の上限（256MB）の両方。上限を下げる変更で
    実 deck が黙って 0 件にならないように。
    """
    from teamagent.adapters.gdrive_client import DEFAULT_GDRIVE_DOWNLOAD_MAX_BYTES
    from teamagent.ingest.office_extract import MAX_OFFICE_COMPRESSED_BYTES

    real_size = 44_855_013
    assert real_size < MAX_OFFICE_COMPRESSED_BYTES
    assert real_size < DEFAULT_GDRIVE_DOWNLOAD_MAX_BYTES


# ── pipeline（_ingest_case_deck）────────────────────────────────────────

FILE_ID = "1FakeDeckFileId_abcdefghijklmnop"
FOLDER_ID = "1FakeFolderId"
NG_FOLDER = "20990101_各社事例集★クライアント展開NG"


class _FakeEmbedder:
    def embed_passage(self, text: str) -> list[float]:
        return [0.0, 0.0, 0.0]


class _FakeDrive:
    def __init__(
        self,
        data: bytes,
        *,
        listed: bool = True,
        size: int | None = None,
        perms: list[DrivePermission] | None = None,
        perms_error: bool = False,
    ) -> None:
        self.data = data
        self.listed = listed
        self.size = len(data) if size is None else size
        self.perms = perms
        self.perms_error = perms_error
        self.downloads = 0

    def list_files(self, folder_id: str | None, request_id: str, **_: Any) -> Any:
        assert folder_id == FOLDER_ID
        files = []
        if self.listed:
            files.append(
                DriveFile(
                    id=FILE_ID,
                    name="架空の事例集.pptx",
                    mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    modified_time="2026-07-29T04:13:50Z",
                    size=self.size,
                    web_view_link=f"https://drive.google.com/file/d/{FILE_ID}/view",
                    md5_checksum=hashlib.md5(self.data, usedforsecurity=False).hexdigest(),
                )
            )
        files.append(
            DriveFile(id="other", name="別資料.pptx", mime_type="x", modified_time=None, size=1)
        )
        return files, None

    def download_file_bytes(self, file_id: str, request_id: str, **_: Any) -> bytes:
        assert file_id == FILE_ID
        self.downloads += 1
        return self.data

    def list_permissions(self, file_id: str, request_id: str, **_: Any) -> list[DrivePermission]:
        if self.perms_error:
            raise RuntimeError("permissions unavailable")
        if self.perms is not None:
            return self.perms
        return [
            DrivePermission(id="1", type="user", role="organizer", email_address="a@example.co.jp"),
            DrivePermission(
                id="2", type="group", role="writer", email_address="team@example.co.jp"
            ),
        ]


class _FakeRepo:
    def __init__(
        self,
        *,
        stored: dict[str, dict[str, str]] | None = None,
        existing: list[str] | None = None,
        sticky_error: bool = False,
    ) -> None:
        self.upserts: list[tuple[Any, list[Any]]] = []
        self.stored = stored or {}
        self.existing = existing or []
        self.sticky_error = sticky_error
        self.retired: list[str] = []
        self.metadata_lookups: list[str] = []

    def get_document_metadata_values(
        self, source_type: str, external_ids: Any, keys: Any
    ) -> dict[str, dict[str, str]]:
        self.metadata_lookups.append(source_type)
        if self.sticky_error:
            raise RuntimeError("db down")
        return {k: v for k, v in self.stored.items() if k in set(external_ids)}

    def upsert_document_with_chunks(
        self, doc: Any, chunks: list[Any], request_id: str, **_: Any
    ) -> str:
        self.upserts.append((doc, chunks))
        return "id"

    def list_case_deck_external_ids(self, source_type: str, file_id: str) -> list[str]:
        assert source_type == "other"
        return list(self.existing)

    def retire_case_documents(
        self, source_type: str, external_ids: list[str], *, retired_at_iso: str
    ) -> int:
        self.retired.extend(external_ids)
        return len(external_ids)


def _spec(**kw: Any) -> CaseDeckSpec:
    base: dict[str, Any] = {
        "file_id": FILE_ID,
        "folder_id": FOLDER_ID,
        "folder_name": NG_FOLDER,
        "name": "架空の事例集（PPTX）",
        "min_cases": 3,
        "extra_metadata": {
            "topic": "ショート動画事例集",
            "case_external_use": "ng",
            "case_external_use_note": "クライアント展開NG（社内のみ）",
        },
    }
    base.update(kw)
    return CaseDeckSpec(**base)


def _ingest(
    drive: _FakeDrive,
    repo: _FakeRepo,
    *,
    spec: CaseDeckSpec | None = None,
    dry_run: bool = False,
    collector: Any = None,
) -> tuple[int, int]:
    from teamagent.ingest.pipeline import _ingest_case_deck

    return _ingest_case_deck(
        spec or _spec(),
        embedder=_FakeEmbedder(),
        repository=repo,  # type: ignore[arg-type]
        owner_email="ingest@example.co.jp",
        dry_run=dry_run,
        request_id="r-deck",
        warning_collector=collector,
        client=drive,
    )


@pytest.fixture(scope="module")
def deck_bytes() -> bytes:
    return _build_pptx(DECK_SHAPES, table_on=4)


def test_pipeline_writes_one_document_per_case(deck_bytes: bytes) -> None:
    """1 事例 = 1 document。母集団の印・照合キー・ACL・NG が載る。

    変異: metadata の ``case_corpus: "true"`` を消すと赤（母集団に入らない）。
    """
    repo = _FakeRepo()
    docs, chunks = _ingest(_FakeDrive(deck_bytes), repo)
    assert docs == 6 and chunks >= 6
    by_client = {doc.metadata["client_name"]: doc for doc, _ in repo.upserts}
    assert set(by_client) == {
        "青葉レコード",
        "銀河フィルム",
        "みなと銀行",
        "白樺",
        "North Field Japan",
        "みらい製菓",
    }
    doc = by_client["青葉レコード"]
    md = doc.metadata
    assert doc.source_type == "other"  # gdrive にしない（ACL 同期・stale の誤爆を避ける）
    assert md["case_corpus"] == "true"
    assert md["case_source"] == "deck"
    assert md["case_brand"] == "北斗シスターズ"
    assert md["case_product"] == "北斗シスターズ"
    assert md["case_effect"].startswith("UGC風動画で話題化：380%達成")
    assert md["case_external_use"] == "ng"
    assert md["case_external_use_note"] == "クライアント展開NG（社内のみ）"
    assert "case_owner" not in md  # 担当者は deck に無い＝付けない
    assert md.get("case_industry") is None  # 業種は推測しない
    assert md["topic"] == "ショート動画事例集"
    assert (md["case_slide_from"], md["case_slide_to"]) == ("2", "5")
    assert doc.source_uri == f"https://drive.google.com/file/d/{FILE_ID}/view"
    assert doc.acl_emails == ["a@example.co.jp", "ingest@example.co.jp"]
    assert doc.acl_groups == ["team@example.co.jp"]
    assert doc.external_id.startswith(f"{FILE_ID}:case:")
    assert by_client["みらい製菓"].metadata["case_company"] == "株式会社みらい製菓"
    assert "case_brand" not in by_client["みなと銀行"].metadata
    # 本文（検索で当たる側）には事例名と NG が入る
    first_chunk = repo.upserts[0][1][0].content
    assert "事例: 青葉レコード（北斗シスターズ）" in first_chunk
    assert "対外利用: クライアント展開NG（社内のみ）" in first_chunk
    assert repo.metadata_lookups == ["other"]  # sticky は source_type=other で読む


def test_ng_comes_from_folder_name_even_without_yaml_flag(deck_bytes: bytes) -> None:
    """yaml に case_external_use を書き忘れても、フォルダ名の「展開NG」で ng に倒れる。"""
    repo = _FakeRepo()
    _ingest(_FakeDrive(deck_bytes), repo, spec=_spec(extra_metadata={"topic": "x"}))
    uses = {doc.metadata["case_external_use"] for doc, _ in repo.upserts}
    assert uses == {"ng"}


def test_sticky_ng_is_not_downgraded(deck_bytes: bytes) -> None:
    """前回 ng の事例は、フォルダ名も yaml も NG でなくなっても ng のまま。"""
    ids, _ = assign_external_ids(
        FILE_ID, split_case_deck(extract_pptx_slide_shapes(deck_bytes)).cases
    )
    repo = _FakeRepo(
        stored={ids[0]: {"case_external_use": "ng", "case_external_use_note": "前回NG"}}
    )
    _ingest(
        _FakeDrive(deck_bytes),
        repo,
        spec=_spec(folder_name="20990101_事例集", extra_metadata={}),
    )
    uses = {doc.external_id: doc.metadata["case_external_use"] for doc, _ in repo.upserts}
    assert uses[ids[0]] == "ng"
    assert uses[ids[1]] == "unknown"


def test_sticky_lookup_failure_writes_nothing(deck_bytes: bytes) -> None:
    from teamagent.ingest.pipeline import CaseCorpusStickyLookupError

    repo = _FakeRepo(sticky_error=True)
    with pytest.raises(CaseCorpusStickyLookupError):
        _ingest(_FakeDrive(deck_bytes), repo)
    assert repo.upserts == []


def test_zero_cases_warns_and_keeps_existing_documents() -> None:
    """抽出 0 件（印の無い deck）: 何も書かず、退役もしない（既存の事例文書を残す）。"""
    from teamagent.ingest.pipeline import _IngestWarningCollector

    shapes = [tuple(s.replace("全体結果ハイライト", "まとめ") for s in sl) for sl in DECK_SHAPES]
    repo = _FakeRepo(existing=[f"{FILE_ID}:case:aaaa"])
    collector = _IngestWarningCollector()
    assert _ingest(_FakeDrive(_build_pptx(shapes)), repo, collector=collector) == (0, 0)
    assert repo.upserts == []
    assert repo.retired == []
    reasons = collector.snapshot("case_decks", FILE_ID).reasons
    assert reasons == {"case_deck_too_few_cases": 1}


def test_too_few_cases_below_min_cases_writes_nothing(deck_bytes: bytes) -> None:
    repo = _FakeRepo()
    assert _ingest(_FakeDrive(deck_bytes), repo, spec=_spec(min_cases=7)) == (0, 0)
    assert repo.upserts == []


def test_over_size_limit_fails_loudly_without_download(
    deck_bytes: bytes, monkeypatch: pytest.MonkeyPatch
) -> None:
    from teamagent.ingest.pipeline import CaseDeckIngestError

    drive = _FakeDrive(deck_bytes, size=300 * 1024 * 1024)
    repo = _FakeRepo()
    with pytest.raises(CaseDeckIngestError, match="MAX_OFFICE_COMPRESSED_BYTES"):
        _ingest(drive, repo)
    assert drive.downloads == 0
    assert repo.upserts == []


def test_corrupt_payload_fails_loudly(deck_bytes: bytes) -> None:
    from teamagent.ingest.pipeline import CaseDeckIngestError

    broken = deck_bytes[: len(deck_bytes) // 2]
    repo = _FakeRepo()
    with pytest.raises(CaseDeckIngestError, match="extract failed"):
        _ingest(_FakeDrive(broken, size=len(deck_bytes)), repo)
    assert repo.upserts == []


def test_file_missing_from_folder_fails_loudly(deck_bytes: bytes) -> None:
    from teamagent.ingest.pipeline import CaseDeckIngestError

    repo = _FakeRepo()
    with pytest.raises(CaseDeckIngestError, match="not found"):
        _ingest(_FakeDrive(deck_bytes, listed=False), repo)
    assert repo.upserts == []


def test_acl_failure_narrows_to_owner_and_warns(deck_bytes: bytes) -> None:
    """permissions が取れない（ACL 空）: 広げずに取込アカウントだけへ絞り、warning を出す。"""
    from teamagent.ingest.pipeline import _IngestWarningCollector

    repo = _FakeRepo()
    collector = _IngestWarningCollector()
    _ingest(_FakeDrive(deck_bytes, perms_error=True), repo, collector=collector)
    assert repo.upserts
    for doc, _ in repo.upserts:
        assert doc.acl_emails == ["ingest@example.co.jp"]
        assert doc.acl_groups == []
    assert collector.snapshot("case_decks", FILE_ID).reasons == {"case_deck_acl_owner_only": 1}


def test_dry_run_writes_nothing(deck_bytes: bytes) -> None:
    repo = _FakeRepo(existing=[f"{FILE_ID}:case:gone"])
    docs, _ = _ingest(_FakeDrive(deck_bytes), repo, dry_run=True)
    assert docs == 6
    assert repo.upserts == [] and repo.retired == []


def test_cases_removed_from_deck_are_retired(deck_bytes: bytes) -> None:
    ids, _ = assign_external_ids(
        FILE_ID, split_case_deck(extract_pptx_slide_shapes(deck_bytes)).cases
    )
    gone = f"{FILE_ID}:case:gone"
    repo = _FakeRepo(existing=[*ids, gone])
    _ingest(_FakeDrive(deck_bytes), repo)
    assert repo.retired == [gone]


def test_retire_is_skipped_when_a_title_could_not_be_read() -> None:
    """表題が読めず落とした事例がある run では退役しない（どれが消えたか分からない）。"""
    shapes = list(DECK_SHAPES)
    start = list(shapes[6])
    start[1] = "銀河フィルムの事例"
    shapes[6] = tuple(start)
    repo = _FakeRepo(existing=[f"{FILE_ID}:case:old"])
    _ingest(_FakeDrive(_build_pptx(shapes)), repo)
    assert len(repo.upserts) == 5
    assert repo.retired == []


def test_mass_retire_is_skipped(deck_bytes: bytes) -> None:
    existing = [f"{FILE_ID}:case:old{i}" for i in range(10)]
    repo = _FakeRepo(existing=existing)
    _ingest(_FakeDrive(deck_bytes), repo)
    assert repo.retired == []


# ── runner の配線 ─────────────────────────────────────────────────────


def test_runner_runs_case_decks_with_gdrive_kind_only(monkeypatch: pytest.MonkeyPatch) -> None:
    import teamagent.ingest.pipeline as pipeline
    from teamagent.ingest.loader import IngestSources

    calls: list[str] = []

    def _fake_handler(spec: Any, **_: Any) -> tuple[int, int]:
        calls.append(spec.file_id)
        return 1, 1

    monkeypatch.setattr(pipeline, "_ingest_case_deck", _fake_handler)
    monkeypatch.setattr(pipeline, "_ingest_gdrive_folder", lambda spec, **_: (0, 0))
    sources = IngestSources(
        version=1,
        slack_channels=(),
        gdrive_folders=(),
        gsheets=(),
        case_corpus_decks=(_spec(),),
    )
    runner = pipeline.IngestRunner(
        repository=_FakeRepo(),  # type: ignore[arg-type]
        embedder=_FakeEmbedder(),
        owner_email="ingest@example.co.jp",
        dry_run=True,
    )
    result = runner.run(sources, kinds=["slack"])
    assert calls == [] and "case_decks" not in result.by_kind
    result = runner.run(sources, kinds=["gdrive"])
    assert calls == [FILE_ID]
    assert result.by_kind["case_decks"].documents_upserted == 1
    assert "gdrive" in result.by_kind  # gdrive の集計とは別の kind


# ── yaml の読み込み ───────────────────────────────────────────────────


def _load(tmp_path: Any, body: str, *, strict: bool = False) -> Any:
    from teamagent.ingest.loader import load_ingest_sources

    path = tmp_path / "sources.yaml"
    path.write_text("version: 1\n" + body, encoding="utf-8")
    return load_ingest_sources(path, skip_placeholder=not strict)


def test_loader_parses_case_corpus_decks(tmp_path: Any) -> None:
    sources = _load(
        tmp_path,
        "case_corpus_decks:\n"
        "  - file_id: F1\n    folder_id: D1\n    folder_name: X\n    name: 事例集\n"
        "    min_cases: 4\n    extra_metadata:\n      topic: t\n",
    )
    (deck,) = sources.case_corpus_decks
    assert (deck.file_id, deck.folder_id, deck.name, deck.min_cases) == ("F1", "D1", "事例集", 4)
    assert deck.extra_metadata == {"topic": "t"}


def test_loader_skips_placeholder_deck_and_strict_raises(tmp_path: Any) -> None:
    body = (
        "case_corpus_decks:\n  - file_id: REPLACE_WITH_DECK\n    folder_id: D1\n    name: 事例集\n"
    )
    assert _load(tmp_path, body).case_corpus_decks == ()
    with pytest.raises(ValueError, match="placeholder"):
        _load(tmp_path, body, strict=True)


@pytest.mark.parametrize(
    ("body", "message"),
    [
        ("  - file_id: F1\n    folder_id: D1\n    name: a\n" * 2, "2 回"),
        ("  - file_id: F1\n    folder_id: D1\n    name: a\n    min_cases: 0\n", "min_cases"),
        ("  - file_id: F1\n    folder_id: D1\n", "name"),
        ("  - file_id: F1\n    name: a\n", "folder_id"),
    ],
)
def test_loader_rejects_misconfigured_decks(tmp_path: Any, body: str, message: str) -> None:
    with pytest.raises(ValueError, match=message):
        _load(tmp_path, "case_corpus_decks:\n" + body)
